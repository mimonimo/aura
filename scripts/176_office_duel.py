#!/usr/bin/env python3
"""오피스 파일 읽기 대결 — docling(지금) · 형식별 라이브러리(python-docx·openpyxl·python-pptx) · kordoc 을 실제 원본으로 견준다.

정답: 오피스 파일 안의 XML 에서 직접 모은 글자(docx w:t · pptx a:t · xlsx 공유 문자열·셀 값). 옛 xls 는 LibreOffice 로 xlsx 로
바꾼 뒤 같은 방식. 표 정답은 docx·pptx 의 표(행·열 수), xlsx 는 시트마다 쓰인 범위.
지표(문서마다 → 갈래별 평균): 낱말 재현율·정밀도(낱말 여러 벌 그대로 견줌), 표 개수 맞음, 표 모양(행·열) 맞음 비율, 실패, 걸린 시간.
결과: data/eval/office_duel_<tag>.json 과 요약 표(OFFICE_DUEL). 원본은 고치지 않는다.

사용(DGX): PYTHONPATH=src .venv-parse/bin/python scripts/176_office_duel.py --n 25 --seed 3
"""
from __future__ import annotations

import argparse
import json
import random
import re
import shutil
import statistics
import subprocess
import sys
import tempfile
import time
import zipfile
from collections import Counter, defaultdict
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "src"))

OUT_OF_SCOPE = re.compile(r"지출|증빙|스캔|영수|정산|집행")
_WORD = re.compile(r"[가-힣A-Za-z0-9]+")
_T = {"docx": re.compile(r"<w:t(?:\s[^>]*)?>([^<]*)</w:t>"), "pptx": re.compile(r"<a:t>([^<]*)</a:t>")}


def words(text: str) -> Counter:
    return Counter(w.lower() for w in _WORD.findall(text or ""))


def _unesc(s: str) -> str:
    import html
    return html.unescape(s)


def reference(path: Path, ext: str) -> tuple[Counter, list[tuple[int, int]]]:
    """(정답 낱말, 정답 표 모양 [(행, 열)])."""
    with zipfile.ZipFile(path) as zf:
        names = zf.namelist()
        if ext == "docx":
            xml = zf.read("word/document.xml").decode("utf-8", "replace")
            text = " ".join(_unesc(t) for t in _T["docx"].findall(xml))
            shapes = []
            for tb in re.findall(r"<w:tbl>.*?</w:tbl>", xml, flags=re.S):      # 겹친 표는 바깥 표만 셈(근사)
                rows = re.findall(r"<w:tr[ >].*?</w:tr>", tb, flags=re.S)
                shapes.append((len(rows), max((len(re.findall(r"<w:tc>|<w:tc ", r)) for r in rows), default=0)))
            return words(text), shapes
        if ext == "pptx":
            text, shapes = [], []
            for n in sorted(x for x in names if re.match(r"ppt/slides/slide\d+\.xml$", x)):
                xml = zf.read(n).decode("utf-8", "replace")
                text += [_unesc(t) for t in _T["pptx"].findall(xml)]
                for tb in re.findall(r"<a:tbl>.*?</a:tbl>", xml, flags=re.S):
                    shapes.append((len(re.findall(r"<a:tr[ >]", tb)), len(re.findall(r"<a:gridCol[ >/]", tb))))
            return words(" ".join(text)), shapes
        # xlsx
        shared = []
        if "xl/sharedStrings.xml" in names:
            sx = zf.read("xl/sharedStrings.xml").decode("utf-8", "replace")
            for si in re.findall(r"<si>.*?</si>", sx, flags=re.S):
                shared.append("".join(_unesc(t) for t in re.findall(r"<t(?:\s[^>]*)?>([^<]*)</t>", si)))
        text, shapes = [], []
        for n in sorted(x for x in names if re.match(r"xl/worksheets/sheet\d+\.xml$", x)):
            xml = zf.read(n).decode("utf-8", "replace")
            rows, cols = set(), set()
            for m in re.finditer(r'<c r="([A-Z]+)(\d+)"([^>]*?)(?:/>|>(.*?)</c>)', xml, flags=re.S):
                attrs, body = m.group(3), m.group(4) or ""
                v = re.search(r"<v>([^<]*)</v>", body)
                inline = re.findall(r"<t(?:\s[^>]*)?>([^<]*)</t>", body)
                val = ""
                if 't="s"' in attrs and v:
                    try:
                        val = shared[int(v.group(1))]
                    except (ValueError, IndexError):
                        val = ""
                elif inline:
                    val = "".join(_unesc(t) for t in inline)
                elif v:
                    val = v.group(1)
                if val.strip():
                    text.append(val)
                    rows.add(int(m.group(2)))
                    cols.add(m.group(1))
            if rows:
                shapes.append((len(rows), len(cols)))
        return words(" ".join(text)), shapes


def to_xlsx(path: Path, tmp: Path) -> Path | None:
    from zzaimy.app.office_pdf import soffice
    exe = soffice()
    if not exe:
        return None
    subprocess.run([exe, "--headless", "--norestore", f"-env:UserInstallation=file://{tmp}/profile", "--convert-to", "xlsx",
                    "--outdir", str(tmp), str(path)], capture_output=True, timeout=180, env={"HOME": str(tmp), "PATH": "/usr/bin:/bin"})
    out = tmp / (path.stem + ".xlsx")
    return out if out.is_file() else None


def by_docling(path: Path, tmp: Path):
    from zzaimy.ingest.parsers.docling import DoclingParser
    r = DoclingParser().parse(path)
    text = " ".join(p.text for p in r.pages) + " " + " ".join(c.text for t in r.tables for c in t.cells)
    return text, [(t.n_rows, t.n_cols) for t in r.tables]


def by_kordoc(path: Path, tmp: Path):
    from zzaimy.ingest.parsers import kordoc
    if not kordoc.available():
        raise RuntimeError("kordoc 없음")
    r = kordoc.KordocParser().parse(path, work_dir=tmp)
    pages = " ".join(p.text for p in r.pages)
    body = pages if pages.strip() else " ".join(e.text for e in r.entries if getattr(e, "text", None) and e.kind != "table")
    text = body + " " + " ".join(c.text for t in r.tables for c in t.cells)
    return text, [(t.n_rows, t.n_cols) for t in r.tables]


def by_libs(path: Path, tmp: Path):
    ext = path.suffix.lower().lstrip(".")
    if ext == "docx":
        import docx
        d = docx.Document(str(path))
        text = [p.text for p in d.paragraphs]
        shapes = []
        for t in d.tables:
            shapes.append((len(t.rows), max((len(r.cells) for r in t.rows), default=0)))
            seen = set()
            for r in t.rows:
                for c in r.cells:
                    if id(c._tc) not in seen:                 # 병합 칸은 한 번만
                        seen.add(id(c._tc))
                        text.append(c.text)
        return " ".join(text), shapes
    if ext == "pptx":
        from pptx import Presentation
        prs = Presentation(str(path))
        text, shapes = [], []

        def walk(shs):
            for sh in shs:
                if sh.shape_type == 6 and hasattr(sh, "shapes"):          # 그룹
                    walk(sh.shapes)
                if getattr(sh, "has_text_frame", False) and sh.has_text_frame:
                    text.append(sh.text_frame.text)
                if getattr(sh, "has_table", False) and sh.has_table:
                    tb = sh.table
                    shapes.append((len(tb.rows), len(tb.columns)))
                    for r in tb.rows:
                        for c in r.cells:
                            text.append(c.text)
        for s in prs.slides:
            walk(s.shapes)
        return " ".join(text), shapes
    if ext in ("xlsx", "xls"):
        src = path if ext == "xlsx" else to_xlsx(path, tmp)
        if src is None:
            raise RuntimeError("xls 변환 실패")
        import openpyxl
        wb = openpyxl.load_workbook(str(src), data_only=True, read_only=True)
        text, shapes = [], []
        for ws in wb.worksheets:
            rows, cols = set(), set()
            for row in ws.iter_rows():
                for c in row:
                    if c.value is not None and str(c.value).strip():
                        text.append(str(c.value))
                        rows.add(c.row)
                        cols.add(c.column)
            if rows:
                shapes.append((len(rows), len(cols)))
        return " ".join(text), shapes
    raise RuntimeError("형식 아님")


METHODS = {"docling": by_docling, "libs": by_libs, "kordoc": by_kordoc}
SUPPORTS = {"docling": {"docx", "xlsx", "pptx", "xls"}, "libs": {"docx", "xlsx", "pptx", "xls"}, "kordoc": {"docx", "xlsx", "xls"}}


def score(ref_w: Counter, ref_t: list, text: str, shapes: list) -> dict:
    cw = words(text)
    inter = sum((ref_w & cw).values())
    rec = inter / max(sum(ref_w.values()), 1)
    prec = inter / max(sum(cw.values()), 1)
    shape_hit = sum(1 for s in ref_t if s in shapes) / len(ref_t) if ref_t else (1.0 if not shapes else 0.0)
    return {"recall": round(rec, 4), "precision": round(prec, 4), "tables_ref": len(ref_t), "tables_got": len(shapes),
            "table_count_ok": len(ref_t) == len(shapes), "table_shape": round(shape_hit, 3)}


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--inventory", default=str(Path.home() / "archive_inventory.jsonl"))
    ap.add_argument("--root", default=str(Path.home() / "data"))
    ap.add_argument("--n", type=int, default=25, help="형식마다 표본 수")
    ap.add_argument("--seed", type=int, default=3)
    ap.add_argument("--tag", default=time.strftime("%Y%m%d%H%M"))
    ap.add_argument("--out", default=str(ROOT / "data" / "eval"))
    args = ap.parse_args()
    pool: dict[str, list[str]] = defaultdict(list)
    for line in open(args.inventory, encoding="utf-8", newline=""):
        try:
            it = json.loads(line)
        except ValueError:
            continue
        ext = (it.get("ext") or "").lower()
        if ext in ("docx", "xlsx", "pptx", "xls") and not it.get("dup_of") and not OUT_OF_SCOPE.search(it["rel"]) \
                and 2_000 < int(it.get("size") or 0) < 20_000_000:
            pool[ext].append(it["rel"])
    rng = random.Random(args.seed)
    rows = []
    for ext, rels in sorted(pool.items()):
        for rel in rng.sample(rels, min(args.n, len(rels))):
            path = Path(args.root) / rel
            with tempfile.TemporaryDirectory(prefix="zz-duel-") as td:
                tmp = Path(td)
                try:
                    ref_src = path if ext != "xls" else to_xlsx(path, tmp)
                    ref_w, ref_t = reference(ref_src, "xlsx" if ext == "xls" else ext) if ref_src else (Counter(), [])
                except Exception as e:
                    rows.append({"rel": rel, "ext": ext, "ref_error": f"{type(e).__name__}: {e}"[:120]})
                    continue
                if not ref_w:
                    continue                                   # 글이 없는 파일(그림만 든 발표 등)은 견줄 수 없다
                row = {"rel": rel, "ext": ext, "ref_words": sum(ref_w.values()), "ref_tables": len(ref_t)}
                for name, fn in METHODS.items():
                    if ext not in SUPPORTS[name]:
                        continue
                    t0 = time.time()
                    try:
                        text, shapes = fn(path, tmp)
                        row[name] = score(ref_w, ref_t, text, shapes) | {"sec": round(time.time() - t0, 2)}
                    except Exception as e:
                        row[name] = {"error": f"{type(e).__name__}: {e}"[:120], "sec": round(time.time() - t0, 2)}
                rows.append(row)
                print(ext, rel[-50:], {k: (v.get("recall") if isinstance(v, dict) else None) for k, v in row.items() if k in METHODS},
                      flush=True)
    summary: dict = {}
    for ext in sorted({r["ext"] for r in rows}):
        for name in METHODS:
            got = [r[name] for r in rows if r["ext"] == ext and name in r]
            if not got:
                continue
            ok = [g for g in got if "error" not in g]
            summary[f"{ext}|{name}"] = {
                "n": len(got), "fail": len(got) - len(ok),
                "recall": round(statistics.mean(g["recall"] for g in ok), 3) if ok else None,
                "precision": round(statistics.mean(g["precision"] for g in ok), 3) if ok else None,
                "table_count_ok": round(statistics.mean(g["table_count_ok"] for g in ok), 3) if ok else None,
                "table_shape": round(statistics.mean(g["table_shape"] for g in ok), 3) if ok else None,
                "sec_median": round(statistics.median(g["sec"] for g in got), 2)}
    out = Path(args.out) / f"office_duel_{args.tag}.json"
    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_text(json.dumps({"summary": summary, "rows": rows}, ensure_ascii=False, indent=1), encoding="utf-8")
    print("OFFICE_DUEL", json.dumps(summary, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
