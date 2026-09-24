"""옛 한글 구조 XML(pyhwp) → 변환기 입력 — 글자·문단·표·그림이 hwpx 와 같은 길로 옮겨진다."""

import base64
import io

from docx import Document
from docx.oxml.ns import qn

from zzaimy.ingest import hwp5_docx

def _png() -> str:
    from PIL import Image

    buf = io.BytesIO()
    Image.new("RGB", (8, 8), "white").save(buf, format="BMP")         # BMP 도 그림으로 들어가야 한다
    return base64.b64encode(buf.getvalue()).decode()


PNG = _png()

XML = f"""<HwpDoc version="5.0.3.0"><DocInfo><IdMappings ko-fonts="1" charshapes="2" parashapes="2" borderfills="2" styles="2"/>
<FaceName name="맑은 고딕"/>
<CharShape basesize="1000" bold="0" italic="0" underline="none" text-color="#000000"><FontFace ko="0"/></CharShape>
<CharShape basesize="1600" bold="1" italic="0" underline="line_through" underline-style="15" text-color="#2525f5"><FontFace ko="0"/></CharShape>
<ParaShape align="both" doubled-margin-left="0" indent="0" linespacing-type="ratio" linespacing="160" head-shape="none" level="0"/>
<ParaShape align="center" doubled-margin-left="0" indent="0" linespacing-type="ratio" linespacing="160" head-shape="outline" level="0"/>
<ParaShape align="both" doubled-margin-left="0" indent="0" linespacing-type="ratio" linespacing="160" head-shape="bullet" level="0" numbering-bullet-id="1"/>
<Bullet char="❍"/>
<BorderFill fillflags="00000000"><Border attribute-name="left" stroke-type="none" width="0.1mm"/><Border attribute-name="right" stroke-type="none" width="0.1mm"/><Border attribute-name="top" stroke-type="none" width="0.1mm"/><Border attribute-name="bottom" stroke-type="none" width="0.1mm"/></BorderFill>
<BorderFill fillflags="00000001"><Border attribute-name="left" stroke-type="solid" width="0.12mm"/><Border attribute-name="right" stroke-type="solid" width="0.12mm"/><Border attribute-name="top" stroke-type="solid" width="0.12mm"/><Border attribute-name="bottom" stroke-type="solid" width="0.12mm"/><FillColorPattern background-color="#d6d6d6"/></BorderFill>
<Style local-name="바탕글" name="Normal"/><Style local-name="개요 1" name="Outline 1"/>
<BinData><BinDataEmbedding storage-id="BIN0001" ext="bmp">{PNG}</BinDataEmbedding></BinData>
</DocInfo><BodyText><Section>
<Paragraph parashape-id="1" style-id="1" new-page="0"><LineSeg><SectionDef><PageDef width="59528" height="84188" left-offset="5669" right-offset="5669" top-offset="4251" bottom-offset="2834" header-offset="2834" footer-offset="2834" bookbinding-offset="0"/></SectionDef><Text charshape-id="1">AIDX 사업 개요</Text></LineSeg></Paragraph>
<Paragraph parashape-id="0" style-id="0" new-page="0"><LineSeg><Text charshape-id="0">첫 줄</Text><ControlChar name="LINE_BREAK"/><Text charshape-id="0">둘째 줄</Text></LineSeg></Paragraph>
<Paragraph parashape-id="2" style-id="0" new-page="0"><LineSeg><Text charshape-id="0">글머리표 항목</Text></LineSeg></Paragraph>
<Paragraph parashape-id="0" style-id="0" new-page="0"><LineSeg><TableControl inline="1" width="40000" height="2000"><TableBody rows="2" cols="2" borderfill-id="2"><TableRow>
<TableCell col="0" row="0" colspan="2" rowspan="1" width="40000" height="1000" borderfill-id="2" valign="middle"><Paragraph parashape-id="0" style-id="0"><LineSeg><Text charshape-id="1">제목 칸</Text></LineSeg></Paragraph></TableCell></TableRow><TableRow>
<TableCell col="0" row="1" colspan="1" rowspan="1" width="10000" height="1000" borderfill-id="2" valign="top"><Paragraph parashape-id="0" style-id="0"><LineSeg><Text charshape-id="0">왼쪽</Text></LineSeg></Paragraph></TableCell>
<TableCell col="1" row="1" colspan="1" rowspan="1" width="30000" height="1000" borderfill-id="2" valign="top"><Paragraph parashape-id="0" style-id="0"><LineSeg><Text charshape-id="0">오른쪽</Text></LineSeg></Paragraph></TableCell></TableRow></TableBody></TableControl></LineSeg></Paragraph>
<Paragraph parashape-id="0" style-id="0" new-page="0"><LineSeg><GShapeObjectControl inline="1" width="4000" height="4000"><ShapeComponent width="4000" height="4000" initial-width="4000" initial-height="4000"><ShapePicture><PictureInfo bindata-id="1"/></ShapePicture></ShapeComponent></GShapeObjectControl></LineSeg></Paragraph>
<Paragraph parashape-id="0" style-id="0" new-page="0"><LineSeg><GShapeObjectControl inline="0" width="40000" height="3000"><ShapeComponent width="40000" height="3000"><BorderLine stroke="solid"/><ShapeRectangle/><TextboxParagraphList valign="middle"><Paragraph parashape-id="1" style-id="0"><LineSeg><Text charshape-id="1">Ⅰ. 사업추진 목표</Text></LineSeg></Paragraph></TextboxParagraphList></ShapeComponent></GShapeObjectControl></LineSeg></Paragraph>
</Section></BodyText></HwpDoc>"""


def test_hwp5_xml_converts_like_hwpx(tmp_path):
    p = tmp_path / "plan.xml"
    p.write_text(XML, encoding="utf-8")
    data, stats = hwp5_docx.convert_xml(p)
    assert stats["tables"] == 1 and stats["images"] == 1 and stats["textboxes"] == 1
    d = Document(io.BytesIO(data))
    h = d.paragraphs[0]
    assert h.style.name.startswith("Heading") and h.runs[0].bold and h.runs[0].font.size.pt == 16
    assert h.text == "AI·DX 사업 개요"                          # 사설 영역 글리프는 가운뎃점으로
    assert not h.runs[0].font.underline                        # pyhwp 의 정의 안 된 밑줄 값은 밑줄이 아니다
    assert len(d.paragraphs[1]._p.findall(".//" + qn("w:br"))) == 1
    assert d.paragraphs[2].text == "❍ 글머리표 항목"                 # 자동 글머리표는 글로 붙는다
    assert d.paragraphs[0].runs[0].font.name == "Nanum Gothic"        # 맑은 고딕 → 독스에 있는 한글 글꼴
    t = d.tables[0]
    assert t.cell(0, 0).text == "제목 칸" and t.cell(0, 0)._tc is t.cell(0, 1)._tc
    assert t.cell(0, 0)._tc.tcPr.find(qn("w:shd")).get(qn("w:fill")) == "D6D6D6"
    assert t.cell(1, 1).text == "오른쪽"
    box = d.tables[1]                                          # 글상자는 테두리 있는 한 칸 표
    assert "Ⅰ. 사업추진 목표" in box.cell(0, 0).text
    assert round(d.sections[0].page_width.inches, 2) == 8.27


def test_text_box_table_gets_the_shape_width(tmp_path):
    p = tmp_path / "plan.xml"
    p.write_text(XML, encoding="utf-8")
    data, _ = hwp5_docx.convert_xml(p)
    d = Document(io.BytesIO(data))
    box = d.tables[1]                                          # 글상자(40000 HWPUNIT = 8000 twips)
    assert box._tbl.tblPr.find(qn("w:tblW")).get(qn("w:w")) == "8000"
    assert box._tbl.tblGrid.findall(qn("w:gridCol"))[0].get(qn("w:w")) == "8000"


def test_hwp5_header_and_page_number_controls(tmp_path):
    """pyhwp 의 Header(chid head)·PageNumberPosition(chid pgnp)·AutoNumbering(kind page)이 머리말·꼬리말·쪽 번호로 옮겨진다."""
    xml = XML.replace(
        '</SectionDef>',
        '<PageNumberPosition chid="pgnp" shape="0" position="bottom_center" dash="-"/>'
        '<Header chid="head" places="both_pages"><HeaderParagraphList paragraphs="1"><Paragraph parashape-id="0" style-id="0">'
        '<LineSeg><Text charshape-id="0">영남이공대학교</Text></LineSeg></Paragraph></HeaderParagraphList></Header>'
        '<Footer chid="foot" places="both_pages"><FooterParagraphList paragraphs="1"><Paragraph parashape-id="1" style-id="0">'
        '<LineSeg><AutoNumbering chid="atno" kind="page" footnoteshape="0" number="2"/></LineSeg></Paragraph></FooterParagraphList></Footer>'
        '</SectionDef>')
    p = tmp_path / "hf.xml"
    p.write_text(xml, encoding="utf-8")
    data, stats = hwp5_docx.convert_xml(p)
    assert stats["headers"] == 1 and stats["footers"] == 1 and "page_numbers" not in stats
    d = Document(io.BytesIO(data))
    assert d.sections[0].header.paragraphs[0].text == "영남이공대학교"
    assert [t.text for t in d.sections[0].footer._element.iter(qn("w:instrText"))] == ["PAGE"]
    assert d.paragraphs[0].text == "AI·DX 사업 개요"                       # 본문은 그대로


def test_paragraphs_under_a_column_set_inside_a_cell_are_kept(tmp_path):
    """셀이 다단(ColumnSet)이면 그 아래 문단·그림도 옮긴다."""
    xml = XML.replace(
        '<TableCell col="1" row="1" colspan="1" rowspan="1" width="30000" height="1000" borderfill-id="2" valign="top"><Paragraph parashape-id="0" style-id="0"><LineSeg><Text charshape-id="0">오른쪽</Text></LineSeg></Paragraph></TableCell>',
        '<TableCell col="1" row="1" colspan="1" rowspan="1" width="30000" height="1000" borderfill-id="2" valign="top"><ColumnSet>'
        '<Paragraph parashape-id="0" style-id="0"><LineSeg><Text charshape-id="0">다단 오른쪽</Text></LineSeg></Paragraph>'
        '<Paragraph parashape-id="0" style-id="0"><LineSeg><GShapeObjectControl inline="1" width="4000" height="4000"><ShapeComponent width="4000" height="4000">'
        '<ShapePicture><PictureInfo bindata-id="1"/></ShapePicture></ShapeComponent></GShapeObjectControl></LineSeg></Paragraph></ColumnSet></TableCell>')
    assert xml != XML
    p = tmp_path / "cs.xml"
    p.write_text(xml, encoding="utf-8")
    data, stats = hwp5_docx.convert_xml(p)
    d = Document(io.BytesIO(data))
    assert stats["images"] == 2
    cell = d.tables[0].cell(1, 1)
    assert "다단 오른쪽" in cell.text and cell._tc.findall(".//" + qn("w:drawing"))
