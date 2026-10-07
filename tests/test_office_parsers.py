"""오피스 직접 읽기(ADR-0051) — 문단·표가 순서대로, 엑셀은 시트마다 표 하나, 발표자료는 쪽마다 글·표."""
import pytest

from zzaimy.ingest.parsers import office

docx = pytest.importorskip("docx")
openpyxl = pytest.importorskip("openpyxl")
pptx = pytest.importorskip("pptx")


def test_docx_paragraphs_and_tables_in_order(tmp_path):
    d = docx.Document()
    d.add_heading("1. 사업 개요", level=1)
    d.add_paragraph("산학협력 기반 인재양성")
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text, t.cell(1, 0).text, t.cell(1, 1).text = "구분", "금액", "인건비", "100"
    p = tmp_path / "a.docx"
    d.save(p)
    r = office.parse_docx(p)
    assert [e.kind for e in r.entries] == ["heading", "text", "table"]
    assert r.tables[0].n_rows == 2 and {c.text for c in r.tables[0].cells} == {"구분", "금액", "인건비", "100"}
    assert "산학협력 기반 인재양성" in r.pages[0].text


def test_xlsx_sheet_is_one_table(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "예산"
    ws.append(["항목", "금액"])
    ws.append(["장비", 300])
    p = tmp_path / "b.xlsx"
    wb.save(p)
    r = office.parse_xlsx(p)
    assert len(r.tables) == 1 and (r.tables[0].n_rows, r.tables[0].n_cols) == (2, 2)
    assert "예산" in r.pages[0].text and "300" in r.pages[0].text


def test_pptx_text_and_table(tmp_path):
    prs = pptx.Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[5])
    s.shapes.title.text = "성과 확산"
    tb = s.shapes.add_table(2, 2, 0, 0, 100, 100).table
    tb.cell(0, 0).text, tb.cell(1, 1).text = "지표", "80%"
    p = tmp_path / "c.pptx"
    prs.save(p)
    r = office.parse_pptx(p)
    assert "성과 확산" in r.pages[0].text and len(r.tables) == 1


def test_pptx_slide_with_unknown_element_falls_back_to_xml_text(tmp_path, monkeypatch):
    """도형 해석이 실패한 슬라이드도 글자는 건진다(python-pptx 의 'has_ph_elm' 오류, 2026-10-06 13건)."""
    from pptx import Presentation
    from pptx.util import Inches

    from zzaimy.ingest.parsers import office
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[5])
    s.shapes.title.text = "사업 성과 요약"
    tb = s.shapes.add_textbox(Inches(1), Inches(2), Inches(4), Inches(1))
    tb.text_frame.text = "취업률 82%"
    path = tmp_path / "a.pptx"
    prs.save(path)

    import pptx.shapes.shapetree as st
    def boom(*a, **k):
        raise AttributeError("'lxml.etree._Element' object has no attribute 'has_ph_elm'")
    monkeypatch.setattr(st, "SlideShapeFactory", boom)
    r = office.parse_pptx(path)
    assert "사업 성과 요약" in r.pages[0].text and "취업률 82%" in r.pages[0].text
