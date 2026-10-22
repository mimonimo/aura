"""사업 관리 — 사업 현황 엑셀에서 연도별 값 채우기(program_sheet)와 전체·현행·지난 세 갈래, 연도별 표시. 엑셀은 가짜 값으로 만든다."""

from __future__ import annotations

import datetime as dt
import json

import openpyxl
import pytest

from tests.test_staff_grades import _add, _client, _no_google  # noqa: F401  (구글 끄기 고정물)
from tests.test_accounts import _app
from zzaimy.app import program_admin, program_sheet

LEDGER = {"programs": [
    {"name": "가나다 혁신사업", "terms": ["가나다"], "ministry": "교육부", "category": "일반재정지원",
     "school": {"years": [2024, 2026], "depts": ["기획처"]}},
    {"name": "라마바 지원사업", "terms": [], "ministry": "대구광역시", "category": "교내 장부",
     "school": {"years": [2025], "depts": ["산학협력단"]}},
    {"name": "가가가가 특화사업", "terms": ["가가가가"], "category": "교내 장부"},
    {"name": "나나나나 특화사업", "terms": ["나나나나"], "category": "교내 장부"},
    {"name": "사업 아닌 것 모음", "terms": ["사업 아닌 행사"], "category": "사업 아님"},
]}


def book_2024(path):
    """2024 꼴 — 순·사업명·…·사업 책임자·기타사항(업무진행부서)·국고·합계·상반기 합계, 합계 줄은 가로 병합."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "2024"
    ws["A1"] = "2024학년도 하반기 재정지원사업 수혜 실적"
    ws.merge_cells("A1:K1")
    head = ["순", "사업명", "세부 사업(과제)명", "지원기관", "총사업기간", "사업 책임자", "행정  담당자",
            "기타사항\n(업무진행부서)", "국고(A)", "합계", "상반기 합계"]
    for j, h in enumerate(head, 1):
        ws.cell(4, j, h)
    data = [
        [1, "2024년 가나다 혁신사업", "1단계", "교육부", "20240301-\n20270228", "BOSS", "김행정", "기획처", 700, 1000, 400],
        [2, "가나다 혁신사업", "2단계", "교육부", "20240301-\n20270228", "홍길순", "", "기획처", 100, 200, 50],
        [3, "라마바지원사업", "", "대구광역시", "20230101-20241231", "박책임", "", "산학협력단", 0, 300, 0],
        [None, None, None, None, None, None, None, None, None, None, None],
        [4, "새로운 교육 프로젝트 사업", "", "교육청", "2024.3.1~2024.12.31", "최새로", "", "교무처", 0, 50, 0],
        [5, "가가가가 나나나나 연계사업", "", "", "", "이모호", "", "", 0, 10, 0],
        [6, "사업 아닌 행사", "", "", "", "", "", "", 0, 5, 0],
    ]
    for i, row in enumerate(data, 5):
        for j, v in enumerate(row, 1):
            ws.cell(i, j, v)
    ws.cell(12, 1, "합계")
    ws.merge_cells("A12:H12")
    ws.cell(12, 10, 1565)
    wb.save(path)


def book_2025(path):
    """2025 꼴 — 앞에 구분 열, 열 순서가 다르고 총사업기간이 시작·끝 두 칸(가로 병합 머리줄, 날짜 값), 합계가 둘(그해·이월 포함)."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "260220"
    ws["A1"] = "2025학년도 재정지원사업 수혜실적"
    head = ["기존,신규,종료", "순", "사업명", "지원기관", "기타사항 (업무진행부서)", "총사업기간", None, "사업 책임자",
            "합계 (A+B)", "이월금(G)", "합계"]
    for j, h in enumerate(head, 1):
        ws.cell(4, j, h)
    ws.merge_cells("F4:G4")
    rows = [
        ["기존", 1, "가나다혁신사업", "교육부", "기획처", dt.datetime(2024, 3, 1), dt.datetime(2027, 2, 28), "BOSS", 1500, 100, 1600],
        ["신규", 2, "새로운 교육 프로젝트 사업", "교육청", "교무처", dt.datetime(2025, 3, 1), dt.datetime(2025, 12, 31), "최새로", 70, 0, 70],
        ["", "※ 금액 단위 원", None, None, None, None, None, None, None, None, None],
    ]
    for i, row in enumerate(rows, 5):
        for j, v in enumerate(row, 1):
            ws.cell(i, j, v)
    ws.cell(8, 1, "합계")
    ws.merge_cells("A8:E8")
    wb.save(path)


@pytest.fixture()
def world(tmp_path):
    app = _app(tmp_path)
    (tmp_path / "kg_external.json").write_text(json.dumps(LEDGER, ensure_ascii=False), encoding="utf-8")
    _add(app, "boss", dept="기획처", grade="head")
    _add(app, "lead1", dept="기획처", grade="lead")
    x24, x25 = tmp_path / "현황(2024).xlsx", tmp_path / "현황(2025).xlsx"
    book_2024(x24)
    book_2025(x25)
    return app, tmp_path, x24, x25


def test_read_book_matches_columns_by_header_words(world):
    _, _, x24, x25 = world
    a = program_sheet.read_book(x24)
    b = program_sheet.read_book(x25)
    assert {r["year"] for r in a} == {2024} and {r["year"] for r in b} == {2025}
    assert [r["name"] for r in a] == ["가나다 혁신사업", "가나다 혁신사업", "라마바지원사업", "새로운 교육 프로젝트 사업",
                                      "가가가가 나나나나 연계사업", "사업 아닌 행사"]       # 빈 줄·합계 줄 빠짐, 앞 연도 떼기
    first = a[0]
    assert first["manager"] == "BOSS" and first["dept"] == "기획처" and first["officer"] == "김행정"
    assert first["period"] == "2024.03.01~2027.02.28" and first["budget_won"] == 1000   # 「상반기 합계」가 아닌 「합계」
    assert first["national_won"] == 700 and first["sub"] == "1단계"
    g = b[0]                                                       # 열 순서가 달라도 같은 열쇠
    assert g["name"] == "가나다혁신사업" and g["manager"] == "BOSS" and g["dept"] == "기획처"
    assert g["period"] == "2024.03.01~2027.02.28"                  # 가로 병합 머리줄 = 시작·끝 두 칸
    assert g["budget_won"] == 1500                                 # 맨 왼쪽 합계(그해 합계)
    assert g["status"] == "기존" and g["extra"].get("이월금(G)") == "100"
    assert len(b) == 2                                             # ※ 메모 줄·합계 줄 빠짐


def test_plan_links_names_and_reports(world):
    app, tmp_path, x24, x25 = world
    items = program_sheet.read_book(x24) + program_sheet.read_book(x25)
    rows = program_admin.merged_rows(tmp_path)
    store = program_admin.load_store(tmp_path)
    sm, new = program_sheet.plan(items, rows, store, not_program=program_admin.not_program_keys(LEDGER))
    assert store == {"rows": {}, "added": {}}                      # 미리보기는 원래 값을 고치지 않는다
    matched = dict(sm["matched"])
    assert set(matched) == {"가나다 혁신사업", "라마바 지원사업"}
    assert [n["name"] for n in sm["new"]] == ["새로운 교육 프로젝트 사업"] and sm["new"][0]["years"] == [2024, 2025]
    assert sm["ambiguous"][0][0] == "가가가가 나나나나 연계사업"
    assert sm["not_program"] == ["사업 아닌 행사"]
    y = new["rows"]["program:가나다혁신사업"]["years"]
    assert y["2024"]["budget_won"] == 1200 and y["2024"]["lines"] == 2          # 세부 과제 두 줄은 더한다
    assert y["2024"]["manager"] == "BOSS, 홍길순" and y["2025"]["budget"] == "1,500원"
    key = sm["new"][0]["key"]
    assert new["added"][key]["source"] == "엑셀" and new["rows"][key]["years"]["2025"]["dept"] == "교무처"
    assert sm["cells"]["fill"] > 0 and sm["cells"]["change"] == 0 and sm["blank_top"] > 0
    # 한 번 더 채우면 새 줄은 이름으로 이어지고 바뀌는 칸이 없다
    rows2 = program_admin.merged_rows(tmp_path, ledger=LEDGER, store=new)
    sm2, _ = program_sheet.plan(items, rows2, new, not_program=program_admin.not_program_keys(LEDGER))
    assert not sm2["new"] and sm2["cells"]["fill"] == 0 and sm2["cells"]["change"] == 0


def _upload(c, path, **form):
    with open(path, "rb") as f:
        return c.post("/programs/admin/sheet/preview", data=form,
                      files={"files": (path.name, f.read(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})


def test_preview_apply_and_year_view(world):
    app, tmp_path, x24, x25 = world
    dev = _client(app, "zzdev", "devpass")
    doc = app.state.db.add_document("현황(2025)", str(x25), doc_type="grant")
    r = _upload(dev, x24, doc_ids=str(doc))
    assert r.status_code == 200 and "data-pa-preview" in r.text and "새로운 교육 프로젝트 사업" in r.text
    assert not (tmp_path / "program_admin.json").exists()          # 미리보기는 저장하지 않는다
    token = r.text.split('name="token" value="')[1].split('"')[0]
    r = dev.post("/programs/admin/sheet/apply", data={"doc": str(doc), "token": token}, follow_redirects=False)
    assert r.status_code == 303 and "ok=" in r.headers["location"]
    assert not list((tmp_path / "tmp" / "program_sheet").glob(f"{token}-*"))   # 올린 임시 파일은 지운다
    store = json.loads((tmp_path / "program_admin.json").read_text())
    assert set(store["rows"]["program:가나다혁신사업"]["years"]) == {"2024", "2025"}
    # 연도를 고르면 그해 진행한 사업만, 칸은 그해 값
    p24 = dev.get("/programs/admin?tab=all&year=2024").text
    assert "1,200원" in p24 and "1,500원" not in p24 and "라마바 지원사업" in p24
    p25 = dev.get("/programs/admin?tab=all&year=2025").text
    assert "1,500원" in p25 and "새로운 교육 프로젝트 사업" in p25 and "엑셀에서 온 줄" in p25
    assert "계정 있음" in p25                                        # 책임자 이름 BOSS = 계정 이름
    assert "lead1" not in json.dumps(store["rows"]["program:가나다혁신사업"].get("heads", []))   # 권한은 주지 않는다
    assert "BOSS" not in program_admin.load_store(tmp_path)["rows"]["program:가나다혁신사업"].get("heads", [])


def test_three_tabs_and_past_read_only(world):
    app, tmp_path, *_ = world
    dev = _client(app, "zzdev", "devpass")
    allp = dev.get("/programs/admin?tab=all").text
    assert "가나다 혁신사업" in allp and "라마바 지원사업" in allp and "지난 사업</span>" in allp
    body = allp.split("<tbody>")[1].split("</tbody>")[0]
    lamaba = [tr for tr in body.split("<tr") if "라마바 지원사업" in tr][0]
    assert "편집" not in lamaba and "숨기기" not in lamaba                 # 전체 갈래에서도 지난 사업 줄은 읽기만
    assert "현행 사업 <span>" in allp and "전체 <span>" in allp
    past = dev.get("/programs/admin?tab=past").text
    assert "편집" not in past.split("<tbody>")[1].split("</tbody>")[0] and "data-pa-sheet" not in past
    r = dev.post("/programs/admin/save", data={"key": "program:라마바지원사업", "name": "라마바 지원사업", "budget": "1원"},
                 follow_redirects=False)
    assert "err=" in r.headers["location"]                            # 지난 사업은 저장도 막는다


def test_hand_edit_not_overwritten_by_default(world):
    app, tmp_path, x24, _ = world
    dev = _client(app, "zzdev", "devpass")
    key = "program:가나다혁신사업"
    r = dev.post("/programs/admin/save", data={"key": key, "name": "가나다 혁신사업", "y_year": "2024", "y_budget": "9억 원",
                                                "y_manager": "손입력"}, follow_redirects=False)
    assert "ok=" in r.headers["location"]
    y = program_admin.load_store(tmp_path)["rows"][key]["years"]["2024"]
    assert y["budget"] == "9억 원" and set(y["by_hand"]) == {"budget", "manager"}
    token = _upload(dev, x24).text.split('name="token" value="')[1].split('"')[0]
    dev.post("/programs/admin/sheet/apply", data={"token": token})
    y = program_admin.load_store(tmp_path)["rows"][key]["years"]["2024"]
    assert y["budget"] == "9억 원" and y["manager"] == "손입력" and y["dept"] == "기획처"   # 손으로 고친 칸은 그대로, 빈칸만
    token = _upload(dev, x24).text.split('name="token" value="')[1].split('"')[0]
    dev.post("/programs/admin/sheet/apply", data={"token": token, "overwrite": "1"})
    y = program_admin.load_store(tmp_path)["rows"][key]["years"]["2024"]
    assert y["budget"] == "1,200원" and "by_hand" not in y                # 덮기를 고르면 덮는다
    # 손으로 더한 줄도 연도별 값을 가진다
    dev.post("/programs/admin/save", data={"key": "", "name": "손 사업", "y_year": "2026", "y_budget": "3억 원"})
    assert program_admin.load_store(tmp_path)["rows"]["program:손사업"]["years"]["2026"]["budget"] == "3억 원"
    assert "3억 원" in dev.get("/programs/admin?year=2026").text


def test_sheet_fill_admin_only(world):
    app, tmp_path, x24, _ = world
    for uid in ("boss", "lead1"):
        c = _client(app, uid)
        assert _upload(c, x24).status_code == 403
        assert c.post("/programs/admin/sheet/apply", data={"doc_ids": "1"}).status_code == 403
        assert "data-pa-sheet" not in c.get("/programs/admin").text
    assert not (tmp_path / "program_admin.json").exists()


def test_clerks_link_projects_and_detail(world):
    """행정 담당자 지정(clerks) · 프로젝트 잇기·끊기 · 새 프로젝트 · 줄 상세 · choices()."""
    app, tmp_path, *_ = world
    _add(app, "clerk1", dept="기획처", grade="clerk")
    _add(app, "lead9", dept="산학협력단", grade="lead")
    key = "program:가나다혁신사업"
    boss = _client(app, "boss")
    r = boss.post("/programs/admin/save", data={"key": key, "leads": ["lead1"], "clerks": ["clerk1"]}, follow_redirects=False)
    assert "ok=" in r.headers["location"]
    assert "err=" in boss.post("/programs/admin/save", data={"key": key, "clerks": ["lead9"]},
                               follow_redirects=False).headers["location"]        # 다른 부서
    roles = program_admin.roles_for_project({"program": key + "|2026"}, tmp_path)
    assert roles["clerks"] == ["clerk1"] and roles["leads"] == ["lead1"]
    assert "가나다 혁신사업" in _client(app, "clerk1").get("/programs/admin").text   # 지정된 사업을 본다
    # 잇기 — 관리자는 아무 프로젝트나, 사업 책임자는 자기가 만든 것만
    db = app.state.db
    p1 = db.create_project("grant", "남의 프로젝트", owner="zzdev")
    p2 = db.create_project("grant", "내 프로젝트", owner="lead1")
    lead = _client(app, "lead1")
    assert lead.post("/programs/admin/link", data={"key": key, "project_id": str(p1), "year": "2026"}).status_code == 403
    r = lead.post("/programs/admin/link", data={"key": key, "project_id": str(p2), "year": "2025"}, follow_redirects=False)
    assert "ok=" in r.headers["location"] and db.get_project(p2)["program"] == key + "|2025"
    page = lead.get("/programs/admin").text
    assert 'href="/project/%d"' % p2 in page and "프로젝트 연결" in page
    lead.post("/programs/admin/link", data={"key": key, "project_id": str(p2), "unlink": "1"})
    assert db.get_project(p2)["program"] == ""
    assert lead.post("/programs/admin/link", data={"key": "program:라마바지원사업", "project_id": str(p2)}).status_code == 404
    r = lead.post("/programs/admin/new-project", data={"key": key, "year": "2026"}, follow_redirects=False)
    assert r.status_code == 303
    pid = int(r.headers["location"].rsplit("/", 1)[1])
    assert db.get_project(pid)["program"] == key + "|2026" and db.get_project(pid)["owner"] == "lead1"
    _add(app, "plain", dept="기획처", grade="member")
    assert _client(app, "plain").post("/programs/admin/new-project", data={"key": key}).status_code == 403
    # 줄 상세
    d = lead.get("/programs/admin/row", params={"key": key}).json()
    assert d["name"] == "가나다 혁신사업" and d["roles"]["clerks"] == ["CLERK1"] and any(p["id"] == pid for p in d["projects"])
    assert lead.get("/programs/admin/row", params={"key": "program:라마바지원사업"}).status_code == 404
    assert boss.get("/programs/admin/row", params={"key": "program:라마바지원사업"}).status_code == 200   # 지난 사업은 읽기로
    # 프로젝트 쪽에서 쓸 사업 목록
    ch = program_admin.choices("all", data_dir=tmp_path)
    assert ch[0]["stage"] != "past" and {c["key"] for c in ch} >= {key, "program:라마바지원사업"}
    assert [c["key"] for c in program_admin.choices("own", user="clerk1", data_dir=tmp_path)] == [key]
    assert program_admin.program_value(key, 2026) == key + "|2026"
