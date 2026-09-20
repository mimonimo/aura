"""추출 문맥 고도화 회귀 잠금 — 표 평문·캡션·그림 글자(OCR)·한글 구조 (ADR-0016).

합성 입력만 쓴다. 외부 도구(tesseract·hwp5html)가 없는 환경에서는 조용히 생략돼야 한다.
"""

from __future__ import annotations

import pytest

import json
import zipfile

# ---- 표 평문 (검색·인용용) ------------------------------------------------------


def test_table_text_renders_caption_rows_and_fills_spans():
    from zzaimy.app.render import render_table_text, table_text

    payload = {
        "n_rows": 3, "n_cols": 3,
        "caption": "표 2. 연도별 예산 (단위: 천원)", "note": "※ 집행 기준",
        "cells": [
            [0, 0, 1, 1, 1, "구분"], [0, 1, 1, 2, 1, "예산"],
            [1, 0, 1, 1, 0, "국고"], [1, 1, 1, 1, 0, "1,200"], [1, 2, 1, 1, 0, "1,350"],
            [2, 0, 1, 1, 0, "대응"], [2, 1, 1, 1, 0, "300"], [2, 2, 1, 1, 0, "450"],
        ],
    }
    text = render_table_text(payload)
    lines = text.splitlines()
    assert lines[0] == "표 2. 연도별 예산 (단위: 천원)"
    assert lines[1] == "구분 | 예산 | 예산"          # colspan 채움 — 열마다 머리글 맥락
    assert lines[2] == "국고 | 1,200 | 1,350"
    assert lines[-1] == "※ 집행 기준"
    # 저장된 text가 있으면 그대로, 없으면 셀에서 만든다. JSON이 아니면 원문 그대로
    assert table_text(json.dumps({**payload, "text": "저장본"})) == "저장본"
    assert table_text(json.dumps(payload)) == text
    assert table_text("구분 | 값\n인건비 | 1,000") == "구분 | 값\n인건비 | 1,000"


def test_table_numbers_pass_numeric_verifier_via_material_text():
    """표 속 수치가 초안 재료·근거 허용목록에 들어간다 — 절대 규칙 1(수치는 인출)."""
    from zzaimy.app.drafter import material_text
    from zzaimy.verify.numbers import verify_numbers

    content = json.dumps({"n_rows": 2, "n_cols": 2, "cells": [
        [0, 0, 1, 1, 1, "연도"], [0, 1, 1, 1, 1, "취업률"],
        [1, 0, 1, 1, 0, "2025"], [1, 1, 1, 1, 0, "71.2%"],
    ]}, ensure_ascii=False)
    evidence = material_text({"kind": "table", "content": content})
    assert "2025 | 71.2%" in evidence and '"cells"' not in evidence
    assert verify_numbers("2025년 취업률 71.2%를 달성했다.", [evidence]).ok
    assert not verify_numbers("2025년 취업률 88.0%를 달성했다.", [evidence]).ok
    assert material_text({"kind": "image_text", "content": "그림 1. 체계도"}) == "그림 1. 체계도"


# ---- 구조 조각: 파서 캡션·그림 글자 ---------------------------------------------


def test_structured_chunks_attach_parser_captions_and_figure_text(tmp_path):
    from zzaimy.app.pipeline import DocumentProcessor
    from zzaimy.ingest.parsers.base import (
        ParsedEntry,
        ParsedImage,
        ParsedPage,
        ParsedTable,
        ParseResult,
        TableCell,
    )

    proc = DocumentProcessor()
    img = tmp_path / "fig_1.png"
    img.write_bytes(b"\x89PNG fake")
    proc._last_result = ParseResult(
        parser="fake", elapsed_s=0.0,
        pages=[ParsedPage(page_no=1, text="무시됨")],
        tables=[ParsedTable(page_no=1, n_rows=1, n_cols=2, cells=(
            TableCell(row=0, col=0, text="항목", is_header=True),
            TableCell(row=0, col=1, text="1,000"),
        ))],
        images=[ParsedImage(page_no=1, path=img)],
        entries=[
            ParsedEntry(page_no=1, kind="text", text="본문 문단입니다"),
            ParsedEntry(page_no=1, kind="table", ref=0, text="표 1. 사업비 내역\n※ 단위: 천원"),
            ParsedEntry(page_no=1, kind="image", ref=0, text="그림 1. 추진 체계도",
                        bbox=(1.0, 2.0, 30.0, 40.0)),
        ],
    )
    proc._last_image_text = {"fig_1.png": "기획처 → 산학협력단 → 학과"}
    chunks = proc._structured_chunks(do_mask=False)
    assert [c["kind"] for c in chunks] == ["text", "table", "image", "image_text"]
    data = json.loads(chunks[1]["content"])
    assert data["caption"] == "표 1. 사업비 내역" and data["note"] == "※ 단위: 천원"
    assert data["text"].splitlines() == ["표 1. 사업비 내역", "항목 | 1,000", "※ 단위: 천원"]
    assert chunks[2]["content"] == "fig_1.png"                     # 뷰어 계약(파일명) 유지
    assert chunks[3]["content"] == "그림 1. 추진 체계도\n기획처 → 산학협력단 → 학과"
    assert chunks[3]["bbox"] == chunks[2]["bbox"]                  # 그림 위치에 글자 레이어


def test_attach_captions_from_neighbouring_text():
    """파서가 캡션을 안 준 경우 — 앞줄 '표 N…'·단위 줄과 뒷줄 각주를 표에, '그림 N…'을 그림에."""
    from zzaimy.app.pipeline import DocumentProcessor

    table = json.dumps({"n_rows": 1, "n_cols": 2,
                        "cells": [[0, 0, 1, 1, 0, "국고"], [0, 1, 1, 1, 0, "1,200"]]})
    chunks = [
        {"kind": "text", "page_no": 1,
         "content": "사업비는 아래 표와 같다. 이 문장은 캡션이 아니다."},
        {"kind": "text", "page_no": 1, "content": "<표 3> 연도별 사업비"},
        {"kind": "text", "page_no": 1, "content": "(단위: 천원)"},
        {"kind": "table", "page_no": 1, "content": table},
        {"kind": "text", "page_no": 1, "content": "※ 2025년은 계획값"},
        {"kind": "text", "page_no": 1, "content": "다음 문단입니다."},
        {"kind": "image", "page_no": 2, "content": "fig_2.png"},
        {"kind": "text", "page_no": 2, "content": "[그림 2] 연차별 추진 일정"},
    ]
    out = DocumentProcessor._attach_captions(chunks)
    assert [c["kind"] for c in out] == ["text", "table", "text", "image", "image_text"]
    data = json.loads(out[1]["content"])
    assert data["caption"] == "<표 3> 연도별 사업비 (단위: 천원)"
    assert data["note"] == "※ 2025년은 계획값"
    assert data["text"].startswith("<표 3> 연도별 사업비 (단위: 천원)\n국고 | 1,200")
    assert out[4]["content"] == "[그림 2] 연차별 추진 일정" and out[4]["page_no"] == 2
    # 캡션 패턴이 아닌 줄은 건드리지 않는다
    keep = [{"kind": "text", "page_no": 1, "content": "그냥 문단"},
            {"kind": "table", "page_no": 1, "content": table}]
    assert DocumentProcessor._attach_captions(keep) == keep


def test_md_to_chunks_attaches_caption_to_table():
    from zzaimy.app.pipeline import DocumentProcessor

    md = ("## 예산\n\n<표 1> 연도별 예산\n\n| 구분 | 값 |\n|---|---|\n| 국고 | 1,200 |\n\n"
          "※ 단위: 천원\n\n본문 문장이다.")
    chunks = DocumentProcessor._md_to_chunks(md, lambda s: s)
    assert [c["kind"] for c in chunks] == ["heading", "table", "text"]
    data = json.loads(chunks[1]["content"])
    assert data["caption"] == "<표 1> 연도별 예산" and data["note"] == "※ 단위: 천원"
    assert "국고 | 1,200" in data["text"]


def test_compose_text_keeps_tables_in_reading_order():
    """표는 문서 끝이 아니라 제자리에 — 앞 문단·캡션과 함께 규정 조각·검토 입력이 된다."""
    from zzaimy.app.pipeline import DocumentProcessor
    from zzaimy.ingest.parsers.base import (
        ParsedEntry,
        ParsedPage,
        ParsedTable,
        ParseResult,
        TableCell,
    )

    t = ParsedTable(page_no=1, n_rows=1, n_cols=2,
                    cells=(TableCell(0, 0, "구분"), TableCell(0, 1, "값")))
    parsed = ParseResult(
        parser="mineru", elapsed_s=0.0, pages=[ParsedPage(1, "x")], tables=[t],
        entries=[
            ParsedEntry(page_no=1, kind="heading", text="2. 사업 예산"),
            ParsedEntry(page_no=1, kind="table", ref=0, text="표 1. 예산"),
            ParsedEntry(page_no=1, kind="text", text="위 표의 값은 계획값이다."),
        ],
    )
    assert DocumentProcessor._compose_text(parsed).split("\n\n") == [
        "2. 사업 예산", "표 1. 예산\n구분 | 값", "위 표의 값은 계획값이다.",
    ]
    # 구조 항목이 없는 파서(docling)는 페이지 본문 뒤에 그 페이지의 표
    parsed2 = ParseResult(
        parser="docling", elapsed_s=0.0,
        pages=[ParsedPage(1, "1쪽 본문"), ParsedPage(2, "2쪽 본문")], tables=[t],
    )
    assert DocumentProcessor._result_to_text(parsed2).split("\n\n") == [
        "1쪽 본문", "구분 | 값", "2쪽 본문",
    ]


def test_llm_correction_refreshes_table_text_and_covers_figure_text(monkeypatch):
    from zzaimy.app.pipeline import DocumentProcessor

    proc = DocumentProcessor.__new__(DocumentProcessor)
    monkeypatch.setattr(
        DocumentProcessor, "_correct_texts",
        lambda self, texts: [t.replace("사업멍", "사업명") for t in texts],
    )
    content = DocumentProcessor._finish_table_payload(
        {"n_rows": 1, "n_cols": 1, "cells": [[0, 0, 1, 1, 1, "사업멍"]]}
    )
    chunks = [
        {"kind": "table", "content": content},
        {"kind": "image_text", "content": "그림 속 사업멍 글자"},
    ]
    assert proc._llm_correct_chunks(chunks)
    data = json.loads(chunks[0]["content"])
    assert data["cells"][0][5] == "사업명" and data["text"] == "사업명"
    assert chunks[1]["content"] == "그림 속 사업명 글자"


# ---- 그림 속 글자 OCR (tesseract) ----------------------------------------------

_TSV_HEADER = ("level\tpage_num\tblock_num\tpar_num\tline_num\tword_num"
               "\tleft\ttop\twidth\theight\tconf\ttext")


def test_parse_tesseract_tsv_keeps_confident_words_by_line():
    from zzaimy.app.pipeline import DocumentProcessor

    rows = [
        _TSV_HEADER,
        "1\t1\t0\t0\t0\t0\t0\t0\t100\t100\t-1\t",
        "5\t1\t1\t1\t1\t1\t0\t0\t10\t10\t91.2\t취업률",
        "5\t1\t1\t1\t1\t2\t12\t0\t10\t10\t88.0\t71.2%",
        "5\t1\t1\t1\t2\t1\t0\t12\t10\t10\t30.5\t잡음",      # 저신뢰 — 버림
        "5\t1\t1\t1\t3\t1\t0\t24\t10\t10\t95.0\t|—~",       # 글자 아님 — 버림
        "5\t1\t2\t1\t1\t1\t0\t40\t10\t10\t80.0\t2025년",
    ]
    assert DocumentProcessor._parse_tesseract_tsv("\n".join(rows), 55.0) == "취업률 71.2%\n2025년"
    assert DocumentProcessor._parse_tesseract_tsv(_TSV_HEADER, 55.0) is None


def test_image_ocr_degrades_silently_without_tesseract(tmp_path, monkeypatch):
    import shutil

    from zzaimy.app.pipeline import DocumentProcessor

    monkeypatch.delenv("ZZAIMY_NO_IMAGE_OCR", raising=False)
    monkeypatch.setattr(DocumentProcessor, "_tess_cache", None)
    monkeypatch.setattr(shutil, "which", lambda name: None)
    proc = DocumentProcessor()
    img = tmp_path / "a.png"
    img.write_bytes(b"\x89PNG fake")
    assert proc._ocr_images([(1, img)]) == {}
    assert DocumentProcessor._tess_cache is False        # 한 번만 찾는다
    # 환경변수로 끄면 도구가 있어도 돌지 않는다
    monkeypatch.setattr(DocumentProcessor, "_tess_cache", ("/usr/bin/tesseract", "kor+eng"))
    monkeypatch.setenv("ZZAIMY_NO_IMAGE_OCR", "1")
    assert proc._ocr_images([(1, img)]) == {}


def test_image_ocr_calls_tesseract_tsv_and_keeps_text(tmp_path, monkeypatch):
    """도구가 있으면 kor+eng·tsv로 부르고 결과를 파일명→글자로 준다 (실행은 가짜)."""
    import subprocess

    from zzaimy.app.pipeline import DocumentProcessor

    calls: list[list[str]] = []

    class _Done:
        returncode = 0
        stdout = _TSV_HEADER + "\n5\t1\t1\t1\t1\t1\t0\t0\t10\t10\t90.0\t취업률\n"

    monkeypatch.setattr(subprocess, "run", lambda args, **kw: (calls.append(args), _Done())[1])
    monkeypatch.delenv("ZZAIMY_NO_IMAGE_OCR", raising=False)
    monkeypatch.setattr(DocumentProcessor, "_tess_cache", ("/usr/bin/tesseract", "kor+eng"))
    proc = DocumentProcessor()
    img = tmp_path / "chart.png"
    img.write_bytes(b"\x89PNG fake")   # 전처리(PIL) 실패 → 원본으로 진행
    assert proc._ocr_images([(1, img)]) == {"chart.png": "취업률"}
    assert calls[0][0] == "/usr/bin/tesseract" and "kor+eng" in calls[0] and calls[0][-1] == "tsv"


# ---- 한글 문서 구조 (HWPX / HWP) -------------------------------------------------

_HWPX_NS = ('xmlns:hp="http://www.hancom.co.kr/hwpml/2011/paragraph" '
            'xmlns:hs="http://www.hancom.co.kr/hwpml/2011/section" '
            'xmlns:hc="http://www.hancom.co.kr/hwpml/2011/core"')


def _para(text: str, style: str = "0") -> str:
    return f'<hp:p styleIDRef="{style}"><hp:run><hp:t>{text}</hp:t></hp:run></hp:p>'


def _cell(r: int, c: int, text: str, rs: int = 1, cs: int = 1, header: int = 0,
          width: int = 100) -> str:
    return (
        f'<hp:tc header="{header}"><hp:subList><hp:p><hp:run><hp:t>{text}</hp:t></hp:run>'
        f'</hp:p></hp:subList><hp:cellAddr colAddr="{c}" rowAddr="{r}"/>'
        f'<hp:cellSpan colSpan="{cs}" rowSpan="{rs}"/><hp:cellSz width="{width}" height="10"/>'
        "</hp:tc>"
    )


def _make_hwpx(path) -> None:
    section = (
        f'<?xml version="1.0" encoding="UTF-8"?><hs:sec {_HWPX_NS}>'
        + _para("1. 사업 개요", style="3")
        + _para("본 사업은 지역 산업 수요에 맞춘 교육과정을 운영한다.")
        + '<hp:p><hp:run><hp:tbl rowCnt="2" colCnt="2">'
        + '<hp:caption side="TOP"><hp:subList><hp:p><hp:run><hp:t>표 1. 연도별 예산</hp:t>'
          "</hp:run></hp:p></hp:subList></hp:caption>"
        + "<hp:tr>" + _cell(0, 0, "구분", cs=2, header=1, width=300) + "</hp:tr>"
        + "<hp:tr>" + _cell(1, 0, "국고", width=100) + _cell(1, 1, "1,200", width=200) + "</hp:tr>"
        + "</hp:tbl></hp:run></hp:p>"
        + '<hp:p><hp:run><hp:pic><hc:img binaryItemIDRef="image1"/>'
          "<hp:caption><hp:subList><hp:p><hp:run><hp:t>그림 1. 추진 체계</hp:t></hp:run></hp:p>"
          "</hp:subList></hp:caption></hp:pic></hp:run></hp:p>"
        + _para("마무리 문단이다.")
        + "</hs:sec>"
    )
    header = (
        '<?xml version="1.0" encoding="UTF-8"?>'
        '<hh:head xmlns:hh="http://www.hancom.co.kr/hwpml/2011/head"><hh:styles>'
        '<hh:style id="0" type="PARA" name="바탕글" engName="Normal"/>'
        '<hh:style id="3" type="PARA" name="개요 1" engName="Outline 1"/>'
        "</hh:styles></hh:head>"
    )
    hpf = (
        '<?xml version="1.0" encoding="UTF-8"?>'
        '<opf:package xmlns:opf="http://www.idpf.org/2007/opf/"><opf:manifest>'
        '<opf:item id="image1" href="BinData/image1.png" media-type="image/png"/>'
        "</opf:manifest></opf:package>"
    )
    with zipfile.ZipFile(path, "w") as zf:
        zf.writestr("mimetype", "application/hwp+zip")
        zf.writestr("Contents/header.xml", header)
        zf.writestr("Contents/section0.xml", section)
        zf.writestr("Contents/content.hpf", hpf)
        zf.writestr("BinData/image1.png", b"\x89PNG fake")


def test_hwpx_parser_reads_tables_pictures_and_captions(tmp_path):
    from zzaimy.ingest.parsers.hwpx import HwpxParser

    f = tmp_path / "계획서.hwpx"
    _make_hwpx(f)
    parsed = HwpxParser().parse(f, work_dir=tmp_path / "imgs")
    assert [e.kind for e in parsed.entries] == ["heading", "text", "table", "image", "text"]
    t = parsed.tables[0]
    assert (t.n_rows, t.n_cols) == (2, 2)
    by = {c.text: c for c in t.cells}
    assert by["구분"].col_span == 2 and by["구분"].is_header and by["1,200"].col == 1
    assert t.col_w == (0.3333, 0.6667)                     # 셀 폭 → 열 폭 비율
    assert parsed.entries[2].text == "표 1. 연도별 예산"
    assert parsed.entries[3].text == "그림 1. 추진 체계"
    assert parsed.images[0].path.exists() and parsed.images[0].path.name == "image1.png"


def test_pipeline_hwpx_gives_structured_chunks_and_reading_order_text(tmp_path, monkeypatch):
    from zzaimy.app.pipeline import DocumentProcessor

    monkeypatch.setenv("ZZAIMY_NO_IMAGE_OCR", "1")
    f = tmp_path / "계획서.hwpx"
    _make_hwpx(f)
    proc = DocumentProcessor()
    text = proc._parse(f)
    assert "1. 사업 개요" in text
    assert "표 1. 연도별 예산\n구분\n국고 | 1,200" in text   # 표가 제자리에, 행 평문으로
    assert "[그림] 그림 1. 추진 체계" in text
    assert proc._last_parse_note.startswith("한글(HWPX) 구조 추출 · 표 1개 · 그림 1장")
    chunks = proc._structured_chunks(do_mask=False)
    assert [c["kind"] for c in chunks] == [
        "heading", "text", "table", "image", "image_text", "text",
    ]
    data = json.loads(chunks[2]["content"])
    assert data["caption"] == "표 1. 연도별 예산" and data["col_w"] == [0.3333, 0.6667]
    assert chunks[3]["content"] == "image1.png" and chunks[4]["content"] == "그림 1. 추진 체계"


def test_hwpx_structured_falls_back_to_plain_text_on_bad_xml(tmp_path, monkeypatch):
    from zzaimy.app.pipeline import DocumentProcessor

    monkeypatch.setenv("ZZAIMY_NO_IMAGE_OCR", "1")
    f = tmp_path / "이상.hwpx"
    with zipfile.ZipFile(f, "w") as zf:   # 네임스페이스 선언이 없는 XML — 구조 파서는 실패
        zf.writestr("Contents/section0.xml", "<hp:p><hp:t>본문이다</hp:t></hp:p>")
    proc = DocumentProcessor()
    assert "본문이다" in proc._parse(f)
    assert proc._last_result is None


def test_hwp5_html_walker_reads_paragraphs_tables_and_images(tmp_path):
    """hwp5html 산출 XHTML — 바깥 문단·표(병합·셀 폭·중첩 표 평탄화)·그림을 읽기 순서로."""
    from zzaimy.ingest.parsers.hwp5 import parse_hwp5_html

    (tmp_path / "bindata").mkdir()
    (tmp_path / "bindata" / "BIN0001.jpg").write_bytes(b"\xff\xd8fake")
    src = (
        '<html><body><div class="Section-0"><div class="Page">'
        '<p class="Normal"><span>신청서 제목</span>&#13;</p>'
        '<table class="borderfill-1"><tr>'
        '<td style="width: 30mm;" rowspan="2" colspan="1"><p><span>성명</span>&#13;</p></td>'
        '<td style="width: 70mm;" rowspan="1" colspan="1"><p><span>홍길동</span>&#13;</p>'
        "<p><span>(합성)</span>&#13;</p></td>"
        '</tr><tr><td style="width: 70mm;" rowspan="1" colspan="1">'
        "<table><tr><td><p>중첩 표 글자</p></td></tr></table></td></tr></table>"
        '<p><img src="bindata/BIN0001.jpg"/></p>'
        '<p class="Normal"><span>제출일: 2026. 9. 15.</span>&#13;</p>'
        "</div></div></body></html>"
    )
    parsed = parse_hwp5_html(src, tmp_path)
    assert [e.kind for e in parsed.entries] == ["text", "table", "image", "text"]
    t = parsed.tables[0]
    assert (t.n_rows, t.n_cols) == (2, 2)
    by = {c.text.split("\n")[0]: c for c in t.cells}
    assert by["성명"].row_span == 2
    assert by["홍길동"].text == "홍길동\n(합성)"           # 셀 안 문단은 줄로
    assert "중첩 표 글자" in by and by["중첩 표 글자"].col == 1
    assert t.col_w == (0.3, 0.7)
    assert parsed.images[0].path.name == "BIN0001.jpg"
    assert parsed.entries[0].text == "신청서 제목"
    assert parsed.entries[-1].text == "제출일: 2026. 9. 15."


# ---- 뷰어·내보내기·맥락 분석 ----------------------------------------------------


def _sample_chunks() -> list[dict]:
    from zzaimy.app.pipeline import DocumentProcessor

    table = DocumentProcessor._finish_table_payload({
        "n_rows": 1, "n_cols": 1, "caption": "표 1. 예산", "note": "※ 단위: 천원",
        "cells": [[0, 0, 1, 1, 0, "1,200"]],
    })
    return [
        {"id": 7, "kind": "table", "page_no": 1, "content": table},
        {"id": 8, "kind": "image", "page_no": 1, "content": "fig.png"},
        {"id": 9, "kind": "image_text", "page_no": 1, "content": "그림 1. 체계도\n기획처 → 학과"},
    ]


def test_viewer_renders_caption_note_and_figure_text():
    from zzaimy.app.render import chunk_blocks, export_markdown

    chunks = _sample_chunks()
    html = "".join(str(b) for b in chunk_blocks(chunks, doc_id=3, asset_by_name={"fig.png": 11}))
    i_cap, i_tbl = html.index("표 1. 예산"), html.index('<table class="extract">')
    assert i_cap < i_tbl < html.index("※ 단위: 천원")
    assert "/doc/3/asset/11" in html and "그림 설명·글자" in html and "기획처 → 학과" in html
    md = export_markdown("x.pdf", chunks)
    assert "**표 1. 예산**" in md and "| 1,200 |" in md and "> 그림 1. 체계도" in md


def test_docx_export_includes_table_caption_and_note():
    import io

    from docx import Document

    from zzaimy.app.render import build_docx

    d = Document(io.BytesIO(build_docx("x", _sample_chunks())))
    texts = [p.text for p in d.paragraphs]
    assert "표 1. 예산" in texts and "※ 단위: 천원" in texts and len(d.tables) == 1


def test_analyze_feeds_table_text_not_json(tmp_path, monkeypatch):
    """맥락 분석 입력은 표 JSON이 아니라 캡션+행 평문, 그림 글자 조각도 포함 (LLM은 가짜)."""
    import sys
    import types

    from zzaimy.app.db import Database
    from zzaimy.app.pipeline import DocumentProcessor

    captured: dict = {}

    class _Completions:
        def create(self, **kw):
            captured["prompt"] = kw["messages"][0]["content"]
            msg = type("Msg", (), {"content": "문서 유형: 계획서"})()
            return type("Resp", (), {"choices": [type("Choice", (), {"message": msg})()]})()

    class _FakeClient:
        model = "fake"

        def __init__(self, *a, **kw) -> None:      # role= 로 부르는 호출부도 받는다
            chat = type("Chat", (), {"completions": _Completions()})()
            self.client = type("Client", (), {"chat": chat})()

    # 생성 클라이언트 모듈은 openai 의존 — 테스트 환경에 없어도 되게 가짜 모듈을 꽂는다
    fake_mod = types.ModuleType("zzaimy.generate.client")
    fake_mod.VllmClient = _FakeClient  # type: ignore[attr-defined]
    monkeypatch.setitem(sys.modules, "zzaimy.generate.client", fake_mod)
    db = Database(tmp_path / "t.db")
    doc_id = db.add_document(filename="a.pdf", stored_path="/tmp/x", doc_type="ocr")
    db.replace_doc_chunks(doc_id, _sample_chunks())
    DocumentProcessor().analyze(db, doc_id)
    assert "[표]\n표 1. 예산\n1,200\n※ 단위: 천원" in captured["prompt"]
    assert '"cells"' not in captured["prompt"]
    assert "[그림 설명·글자] 그림 1. 체계도" in captured["prompt"]
    assert (db.get_document(doc_id) or {}).get("ai_review") == "문서 유형: 계획서"


def test_image_document_falls_back_to_tesseract_when_docling_unavailable(tmp_path, monkeypatch):
    """오프라인 VM 실측: docling 레이아웃 모델이 없어 PNG 접수가 실패했다 → 전체 이미지 tesseract 폴백."""
    from zzaimy.app.pipeline import DocumentProcessor
    from zzaimy.ingest.parsers import docling as docling_mod

    img = tmp_path / "스크린샷.png"
    img.write_bytes(b"\x89PNG\r\n\x1a\n" + b"0" * 64)

    def boom(self, path):
        raise RuntimeError("Cannot find an appropriate cached snapshot folder")

    monkeypatch.setattr(docling_mod.DoclingParser, "parse", boom)
    proc = DocumentProcessor()
    monkeypatch.setattr(proc, "_ocr_image_text", lambda path, tmp_dir: "국가근로장학생 사전교육 안내\n일시 2026.09.01")
    text = proc._parse_inner(img)
    assert "사전교육 안내" in text
    assert proc._last_parse_note == "이미지 글자 OCR (tesseract)"
    assert proc._last_images == [(1, img)] and proc._ocr_used is True

    # tesseract 도 없으면 원래 오류를 사람이 읽을 메시지로 감싸 올린다
    monkeypatch.setattr(proc, "_ocr_image_text", lambda path, tmp_dir: None)
    with pytest.raises(RuntimeError, match="문서 판독 실패"):
        proc._parse_inner(img)


def test_sentence_is_not_a_caption():
    """스캔 문제집 실측: 표 앞 문장이 표 제목으로 붙었다 → 문장은 본문으로, 명사구·표 N 은 캡션으로."""
    from zzaimy.app.pipeline import DocumentProcessor as P

    assert P._plausible_caption("표 1. 연도별 예산 내역")
    assert P._plausible_caption("(단위: 천원)")
    assert P._plausible_caption("2022학년도 입학자 연계교육과정")
    assert not P._plausible_caption("오답피하기① ICMP Flooding은 대역폭 공격이다.")
    assert not P._plausible_caption("전자입찰 시스템의 각 구성요소들은 자율성을 보장받아야 한다")
    cap, back = P._split_caption("표 2 RR 레코드 유형\n오답피하기① AES 알고리즘은 DES보다 안전하다.")
    assert cap == "표 2 RR 레코드 유형" and "AES" in back
