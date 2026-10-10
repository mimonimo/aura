"""구글 독스 공통 양식 5종 사양 — 독스 규칙(표 8열 이하·모양 일치·제목 단계)과 지침·실문서 값 없음."""
from zzaimy.ingest import gdocs_templates as gt


def test_specs_follow_docs_rules():
    assert set(gt.SPECS) == {"plan", "report", "program_plan", "program_report", "minutes"}
    for sid, spec in gt.SPECS.items():
        assert gt.check_spec(spec) == [], sid
        heads = [b["text"] for b in spec["blocks"] if "h" in b]
        guides = [b for b in spec["blocks"] if "guide" in b]
        tables = [b for b in spec["blocks"] if "table" in b]
        assert len(heads) >= 8 and len(guides) >= 5 and len(tables) >= 5, sid
        text = gt.outline_text(spec)
        assert "영남이공" not in text and "2023" not in text           # 실문서 값은 넣지 않는다


def test_plan_has_common_backbone_and_rules():
    heads = [b["text"] for b in gt.PLAN["blocks"] if "h" in b]
    for part in ("Ⅰ. 사업 추진 배경 및 목표", "Ⅱ. 사업 추진 체계", "Ⅲ. 사업 추진 계획", "Ⅳ. 성과관리 계획", "Ⅴ. 재정투자 계획"):
        assert part in heads
    # 실문서 절 집계(계획서 5,732건)에서 여러 사업이 두는 요소 — 빠지면 공통 양식이 얇아진다
    for part in ("사업 요약", "1. 추진 배경 및 필요성", "2. 대학 중장기 발전계획과의 연계", "2. 위원회·협의체 운영", "6. 운영 규정 및 제도화",
                 "3. 대응자금 확보 계획", "3. 자율성과지표 설정 근거", "5. 자체평가·환류 계획", "7. 이전 평가·컨설팅 결과 반영"):
        assert part in heads, part
    text = gt.outline_text(gt.PLAN)
    assert "계산하지 말고" in text and "그대로 옮기고" in text           # 절대 규칙 1·5 를 지침에


def test_check_spec_catches_wide_tables():
    bad = {"blocks": [gt.H(1, "가"), gt.T([str(i) for i in range(9)], 1)]}
    assert gt.check_spec(bad)


def test_fit_widths_keeps_header_words_on_one_line():
    heads = ["구분", "지표명(단위)", "기준값", "1차년도", "2차년도", "3차년도", "4차년도", "측정 방법·증빙"]
    w = gt.fit_widths(heads, [1, 2.4, 1, 1, 1, 1, 1, 2])
    assert abs(sum(w) - gt.CONTENT_W) < 1
    assert all(x >= gt._min_width(h) - 0.1 for x, h in zip(w, heads))   # 「1차년도」가 「1차년/도」로 꺾이지 않는다
    key, val = gt.fit_widths(["사업명", "총 사업비 (국비·지방비·대응자금)"], [1, 3], kv=True)
    assert key >= gt._min_width("(국비·지방비·대응자금)") and key <= gt.CONTENT_W * gt.KV_KEY_MAX + 0.1
    bad = {"blocks": [gt.H(1, "가"), gt.KV(["총사업비(국비·지방비·대응자금·기타재원)"])]}
    assert any("항목 이름" in e for e in gt.check_spec(bad))          # 독스는 한글을 빈칸에서만 꺾는다


def test_check_spec_catches_section_without_guide():
    assert gt.check_spec({"blocks": [gt.H(1, "가"), gt.P()]}) == ["작성 지침·표가 없는 절: 가"]
    assert gt.check_spec({"blocks": [gt.H(1, "가"), gt.G("지침"), gt.P()]}) == []


def test_pick_matches_document_kind_words_only():
    assert gt.pick("2027년 RISE 사업계획서 초안을 써 줘")["id"] == "plan"
    assert gt.pick("3차년도 실적보고서 초안 작성해 줘")["id"] == "report"
    assert gt.pick("AI 특강 프로그램 실시계획서 만들어 줘")["id"] == "program_plan"
    assert gt.pick("취업 캠프 결과보고서 초안")["id"] == "program_report"
    assert gt.pick("공문 초안 써 줘") is None                       # 공통 양식이 없는 갈래는 빈 문서로


def test_sources_are_attached_to_their_sections():
    """절마다 근거 위치가 지침 끝에 붙고(지침 없는 절은 표 앞에 근거 지침), 사업 이름은 넣지 않는다."""
    blocks = gt.PLAN["blocks"]
    i = next(k for k, b in enumerate(blocks) if b.get("text") == "가. 과제 개요")
    assert "근거 — 평가편람" in blocks[i + 1].get("guide", "") and "table" in blocks[i + 2]
    j = next(k for k, b in enumerate(blocks) if b.get("text") == "1. 추진 배경 및 필요성")
    assert "근거 — 공고·기본계획" in blocks[j + 1]["guide"]
    assert not any(w in gt.outline_text(gt.PLAN) for w in ("LINC", "RISE", "영남이공", "AID"))


def test_length_notes_come_from_measured_medians():
    """분량 지침은 실측 중앙값(백 단위 반올림)에서 — 측정이 없는 절에는 붙이지 않는다."""
    assert gt._length_note("plan", "1. 추진 배경 및 필요성") == "분량 — 실문서 같은 절 중앙값 약 1,800자(표 포함)"
    assert gt._length_note("plan", "사업 요약") == ""
    assert all(k in gt._LEN[sid] for sid in gt.LENGTH_OF for k in gt.LENGTH_OF[sid].values())


def test_program_templates_have_measured_common_parts():
    """단위 프로그램 문서 1,151건(24개 사업)에서 여러 사업이 두는 요소 — 모집·선발, 협력 기관, 평가 방법, 참여 현황, 운영 사진."""
    plan = [b["text"] for b in gt.PROGRAM_PLAN["blocks"] if "h" in b]
    report = [b["text"] for b in gt.PROGRAM_REPORT["blocks"] if "h" in b]
    for h in ("3. 참여 대상 및 모집·선발", "5. 운영 체계 및 협력 기관", "6. 성과 목표 및 평가 방법", "7. 안전·개인정보 관리"):
        assert h in plan, h
    for h in ("3. 참여 현황", "5. 협력 기관 참여 실적", "8. 운영 사진"):
        assert h in report, h


def test_strip_guides_docx_removes_only_guide_paragraphs():
    import io
    from docx import Document

    d = Document()
    d.add_paragraph("1. 추진 배경")
    d.add_paragraph(gt.GUIDE_PREFIX + "공고의 목적을 쓴다.")
    d.add_paragraph("본문이다.")
    buf = io.BytesIO()
    d.save(buf)
    out, n = gt.strip_guides_docx(buf.getvalue())
    assert n == 1 and [p.text for p in Document(io.BytesIO(out)).paragraphs] == ["1. 추진 배경", "본문이다."]


def test_examples_attach_to_real_sections_and_match_table_shape():
    """예시는 실제 절에 붙고, 「표 한 줄」 예시는 그 절 첫 표의 열 수와 같다(모양을 보이려는 예시가 모양부터 틀리면 안 된다).
    값은 ○○ 로 비운다 — 실제 수치·사업 이름을 넣지 않는다(절대 규칙 1·10)."""
    import re
    from zzaimy.ingest import gdocs_templates as gt

    for sid, exs in gt.EXAMPLES.items():
        blocks = gt.SPECS[sid]["blocks"]
        heads = [b["text"] for b in blocks if "h" in b]
        for head, ex in exs.items():
            assert head in heads, (sid, head)
            assert gt.EXAMPLE_PREFIX + ex in " ".join(b.get("guide", "") for b in blocks), (sid, head)
            assert not re.search(r"\d{3,}", ex), (sid, head)                 # 세 자리 이상 숫자(금액·인원)는 넣지 않는다
            if ex.startswith("표 한 줄 — "):
                i = next(i for i, b in enumerate(blocks) if b.get("text") == head and "h" in b)
                table = next(b["table"] for b in blocks[i + 1:] if "table" in b)
                assert len(ex[len("표 한 줄 — "):].split(" | ")) == len(table["columns"]), (sid, head, table["columns"])


def test_minutes_template_picks_and_keeps_budget_rule():
    """회의록 — 「회의록」 요청이면 이 양식, 예산 변경 표는 금액을 계산하지 않는 규칙을 단다(절대 규칙 5)."""
    assert gt.pick("운영위원회 제3차 회의록 초안 써 줘")["id"] == "minutes"
    assert gt.pick("특강 결과보고서")["id"] == "program_report"
    blocks = gt.MINUTES["blocks"]
    i = next(i for i, b in enumerate(blocks) if b.get("text") == "4. 예산 변경 내역(해당 시)")
    assert "계산하지 말고" in blocks[i + 1]["guide"]


def test_unit_moves_to_caption_and_columns_align_like_forms():
    """머리말의 단위는 표 위 오른쪽 캡션으로, 금액 칸은 오른쪽·짧은 값 칸은 가운데·글 칸은 왼쪽."""
    heads, unit = gt.unit_of(["구분(단위: 백만원)", "1차년도", "합계"])
    assert heads == ["구분", "1차년도", "합계"] and unit == "(단위: 백만원)"
    assert gt.col_align("1차년도", money=True) == "END" and gt.col_align("1차년도", money=False) == "CENTER"
    assert gt.col_align("금액(원)", False) == "END" and gt.col_align("연번", False) == "CENTER"
    assert gt.col_align("주요 내용", False) == "START" and gt.col_align("달성 여부", False) == "CENTER"
    assert set(gt.ACCENT) == set(gt.SPECS)                       # 갈래마다 주조색


def test_unit_caption_is_not_body_and_stays_above_table():
    """단위 캡션은 본문 글자로 세지 않고(빈 절이 '쓴 절'로 보이지 않게), 새 글은 캡션 앞에 들어간다(캡션은 표에 붙어 있게)."""
    from zzaimy.ingest import gdocs

    def para(s, text):
        return {"startIndex": s, "endIndex": s + len(text) + 1, "paragraph": {"elements": [{"textRun": {"content": text + "\n"}}]}}
    g = para(10, gdocs.GUIDE_PREFIX + "금액은 계산하지 않는다")
    cap = para(g["endIndex"], "(단위: 백만원)")
    tbl = {"startIndex": cap["endIndex"], "endIndex": cap["endIndex"] + 30, "table": {}}
    info = {"sections": [{"index": 1, "start": 0}], "end": 100}
    assert gdocs._guide_anchor([g, cap, tbl], info, {"index": 1, "start": 0}) == g["endIndex"] - 1
    assert gdocs.UNIT_CAPTION.match("(단위: 천원)") and not gdocs.UNIT_CAPTION.match("(단위: 천원) 국비 포함")


def test_slim_docx_drops_embedded_fonts_and_uses_word_font():
    """독스가 심은 글꼴 파일을 빼고(첨부가 6MB 던 것) 글꼴 이름을 Word 기본 한글 글꼴로 — 글·모양은 그대로."""
    import io
    import zipfile
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("[Content_Types].xml", '<Types><Override PartName="/word/fonts/NotoSansKR-regular.ttf" ContentType="x"/>'
                                          '<Override PartName="/word/document.xml" ContentType="y"/></Types>')
        z.writestr("word/fonts/NotoSansKR-regular.ttf", b"\0" * 5000)
        z.writestr("word/_rels/fontTable.xml.rels", '<Relationships><Relationship Id="rId1" Type="t" Target="fonts/NotoSansKR-regular.ttf"/></Relationships>')
        z.writestr("word/fontTable.xml", '<w:fonts><w:font w:name="Noto Sans KR"><w:embedRegular r:id="rId1"/></w:font></w:fonts>')
        z.writestr("word/settings.xml", '<w:settings><w:embedTrueTypeFonts w:val="1"/></w:settings>')
        z.writestr("word/document.xml", '<w:document><w:r><w:rPr><w:rFonts w:ascii="Noto Sans KR" w:eastAsia="Noto Sans KR"/></w:rPr>'
                                        '<w:t>추진 배경</w:t></w:r></w:document>')
    out = zipfile.ZipFile(io.BytesIO(gt.slim_docx(buf.getvalue())))
    names = out.namelist()
    assert not any(n.startswith("word/fonts/") for n in names)
    doc = out.read("word/document.xml").decode()
    assert "추진 배경" in doc and "Noto Sans KR" not in doc and '"맑은 고딕"' in doc
    assert "embed" not in out.read("word/fontTable.xml").decode() and "embedTrueTypeFonts" not in out.read("word/settings.xml").decode()
    assert "fonts/" not in out.read("[Content_Types].xml").decode() and "fonts/" not in out.read("word/_rels/fontTable.xml.rels").decode()


def test_table_only_marks_overview_tables_and_minutes_records():
    """항목 표만 있는 개요 절과 회의록의 기록 표 절만 「표만 채운다」 — 계획서 분석 절은 표가 있어도 글을 쓴다."""
    def marked(spec):
        out = []
        for i, b in enumerate(spec["blocks"]):
            if "guide" in b and gt.TABLE_ONLY in b["guide"]:
                out.append(next(x["text"] for x in reversed(spec["blocks"][:i]) if "h" in x))
        return out
    assert "가. 과제 개요" in marked(gt.PLAN) and "3. 대학 여건 및 현황 분석" not in marked(gt.PLAN)
    assert "실적 요약" in marked(gt.REPORT) and "가. 추진 실적" not in marked(gt.REPORT)
    assert "1. 프로그램 개요" in marked(gt.PROGRAM_PLAN) and "4. 세부 운영 내용" not in marked(gt.PROGRAM_PLAN)
    assert "2. 상정 안건" in marked(gt.MINUTES) and "나. 논의 내용" not in marked(gt.MINUTES)


def test_toc_lists_chapters_and_sections_only():
    toc = gt.toc_lines(gt.PLAN)
    assert toc[0] == (1, "사업 개요") and (1, "Ⅰ. 사업 추진 배경 및 목표") in toc and (2, "1. 추진 배경 및 필요성") in toc
    assert all(level <= 2 for level, _ in toc) and (3, "가. 과제 개요") not in toc
    assert "plan" in gt.COVER_PAGE and "minutes" not in gt.COVER_PAGE          # 목차는 표지 한 쪽을 두는 갈래만
