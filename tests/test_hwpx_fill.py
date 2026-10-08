"""서식 보존 채우기(hwpx_fill) — 절 제목 뒤(안내 상자 뒤)에 문단·표를 끼워 넣고, 손대지 않은 것은 바이트 그대로."""

import io
import re
import zipfile

from docx import Document

from zzaimy.ingest import hwpx_docx, hwpx_fill

HEADER = """<?xml version="1.0" encoding="UTF-8"?>
<hh:head xmlns:hh="http://www.hancom.co.kr/hwpml/2011/head" xmlns:hc="http://www.hancom.co.kr/hwpml/2011/core">
<hh:refList>
<hh:fontfaces><hh:fontface lang="HANGUL"><hh:font id="0" face="맑은 고딕" type="TTF"/></hh:fontface></hh:fontfaces>
<hh:borderFills itemCnt="3">
 <hh:borderFill id="1"><hh:leftBorder type="NONE" width="0.1 mm"/><hh:rightBorder type="NONE" width="0.1 mm"/><hh:topBorder type="NONE" width="0.1 mm"/><hh:bottomBorder type="NONE" width="0.1 mm"/></hh:borderFill>
 <hh:borderFill id="2"><hh:leftBorder type="DASH" width="0.12 mm"/><hh:rightBorder type="DASH" width="0.12 mm"/><hh:topBorder type="DASH" width="0.12 mm"/><hh:bottomBorder type="DASH" width="0.12 mm"/></hh:borderFill>
 <hh:borderFill id="3"><hh:leftBorder type="SOLID" width="0.12 mm"/><hh:rightBorder type="SOLID" width="0.12 mm"/><hh:topBorder type="SOLID" width="0.12 mm"/><hh:bottomBorder type="SOLID" width="0.12 mm"/></hh:borderFill>
</hh:borderFills>
<hh:charProperties itemCnt="2">
 <hh:charPr id="0" height="1000" textColor="#000000"><hh:fontRef hangul="0"/></hh:charPr>
 <hh:charPr id="1" height="1600" textColor="#2525F5"><hh:fontRef hangul="0"/><hh:bold/></hh:charPr>
</hh:charProperties>
<hh:paraProperties itemCnt="2">
 <hh:paraPr id="0"><hh:align horizontal="JUSTIFY"/><hh:heading type="NONE" level="0"/></hh:paraPr>
 <hh:paraPr id="1"><hh:align horizontal="CENTER"/><hh:heading type="OUTLINE" level="0"/></hh:paraPr>
</hh:paraProperties>
<hh:styles itemCnt="1"><hh:style id="0" type="PARA" name="바탕글" paraPrIDRef="0" charPrIDRef="0"/></hh:styles>
</hh:refList></hh:head>"""

LINESEG = '<hp:linesegarray><hp:lineseg textpos="0" vertpos="0" vertsize="1000" textheight="1000" baseline="850" spacing="600" horzpos="0" horzsize="48188" flags="393216"/></hp:linesegarray>'


def _cell(text: str, bf: str, r: int, c: int) -> str:
    return (f'<hp:tc borderFillIDRef="{bf}"><hp:subList><hp:p paraPrIDRef="0"><hp:run charPrIDRef="0"><hp:t>{text}</hp:t></hp:run>{LINESEG}</hp:p></hp:subList>'
            f'<hp:cellAddr rowAddr="{r}" colAddr="{c}"/><hp:cellSpan rowSpan="1" colSpan="1"/><hp:cellSz width="20000" height="1000"/><hp:cellMargin left="141" right="141" top="141" bottom="141"/></hp:tc>')


def _section() -> str:
    box = (f'<hp:p id="0" paraPrIDRef="0" styleIDRef="0"><hp:run charPrIDRef="0"><hp:tbl id="900" rowCnt="1" colCnt="1" borderFillIDRef="2"><hp:sz width="40000"/><hp:tr>'
           f'{_cell("【작성방법】 1) 현황을 쓴다", "2", 0, 0)}</hp:tr></hp:tbl></hp:run>{LINESEG}</hp:p>')
    body_table = (f'<hp:p id="0" paraPrIDRef="0" styleIDRef="0"><hp:run charPrIDRef="0"><hp:tbl id="901" rowCnt="1" colCnt="2" borderFillIDRef="3"><hp:sz width="40000"/><hp:tr>'
                  f'{_cell("구분", "3", 0, 0)}{_cell("값", "3", 0, 1)}</hp:tr></hp:tbl></hp:run>{LINESEG}</hp:p>')
    return f"""<?xml version="1.0" encoding="UTF-8"?>
<hs:sec xmlns:hs="http://www.hancom.co.kr/hwpml/2011/section" xmlns:hp="http://www.hancom.co.kr/hwpml/2011/paragraph">
<hp:p id="0" paraPrIDRef="1" styleIDRef="0"><hp:run charPrIDRef="0"><hp:secPr><hp:pagePr landscape="WIDELY" width="59528" height="84188"><hp:margin left="5669" right="5669" top="4251" bottom="2834" header="2834" footer="2834" gutter="0"/></hp:pagePr></hp:secPr></hp:run><hp:run charPrIDRef="1"><hp:t>1. 사업 개요</hp:t></hp:run>{LINESEG}</hp:p>
<hp:p id="0" paraPrIDRef="1" styleIDRef="0"><hp:run charPrIDRef="1"><hp:t>1.1. 대학의 여건 분석</hp:t></hp:run>{LINESEG}</hp:p>
{box}
<hp:p id="0" paraPrIDRef="1" styleIDRef="0"><hp:run charPrIDRef="1"><hp:t>1.2. 특성화 방향</hp:t></hp:run>{LINESEG}</hp:p>
{body_table}
<hp:p id="0" paraPrIDRef="0" styleIDRef="0"><hp:run charPrIDRef="0"><hp:t>마무리 문단.</hp:t></hp:run>{LINESEG}</hp:p>
</hs:sec>"""


def _hwpx(tmp_path, two_sections: bool = False):
    p = tmp_path / "form.hwpx"
    with zipfile.ZipFile(p, "w") as zf:
        zf.writestr(zipfile.ZipInfo("mimetype"), "application/hwp+zip", compress_type=zipfile.ZIP_STORED)
        zf.writestr("Contents/header.xml", HEADER, compress_type=zipfile.ZIP_DEFLATED)
        zf.writestr("BinData/image1.png", b"\x89PNG-not-really", compress_type=zipfile.ZIP_STORED)
        zf.writestr("Contents/section0.xml", _section(), compress_type=zipfile.ZIP_DEFLATED)
        if two_sections:
            zf.writestr("Contents/section1.xml", _section().replace("1.1. 대학의 여건 분석", "3.1. 성과관리 계획"), compress_type=zipfile.ZIP_DEFLATED)
    return p


BODIES = [
    {"heading": "1.1. 대학의 여건 분석", "items": [("text", "첫 문단 <검증> & 둘.\n둘째 문단"), ("table", [["구분", "2025", "2026"], ["재학생", "1,000", "1,100"]]), ("text", "표 뒤 문단")]},
    {"heading": "1.2 특성화 방향", "items": [("text", "특성화 문단")]},
    {"heading": "9.9. 없는 절", "items": [("text", "안 들어간다")]},
]


def test_fill_inserts_after_box_and_keeps_untouched_bytes(tmp_path):
    src = _hwpx(tmp_path)
    out = tmp_path / "out.hwpx"
    rep = hwpx_fill.fill(src, BODIES, out)
    assert rep["filled"] == ["1.1. 대학의 여건 분석", "1.2 특성화 방향"] and rep["skipped"] == ["9.9. 없는 절"]
    assert rep["paragraphs"] == 4 and rep["tables"] == 1 and rep["sections_changed"] == ["Contents/section0.xml"]
    zi, zo = zipfile.ZipFile(src), zipfile.ZipFile(out)
    # 항목 순서·압축 방식 보존, mimetype 첫 항목 무압축, 안 바뀐 항목은 CRC 동일
    assert [i.filename for i in zi.infolist()] == [i.filename for i in zo.infolist()]
    assert zo.infolist()[0].filename == "mimetype" and zo.infolist()[0].compress_type == zipfile.ZIP_STORED
    assert all(a.compress_type == b.compress_type for a, b in zip(zi.infolist(), zo.infolist()))
    assert {i.filename for i in zi.infolist() if i.CRC != zo.getinfo(i.filename).CRC} == {"Contents/section0.xml"}
    xml = zo.read("Contents/section0.xml").decode()
    # 1.1 본문은 안내 상자 뒤·1.2 제목 앞, 1.2 본문은 1.2 제목 뒤(번호 끝 점 유무는 상관없다)
    i_box, i_body, i_12, i_12body = xml.index("【작성방법】"), xml.index("첫 문단"), xml.index("1.2. 특성화 방향"), xml.index("특성화 문단")
    assert i_box < i_body < i_12 < i_12body
    assert "&lt;검증&gt; &amp; 둘." in xml and "안 들어간다" not in xml
    # 줄 배치 캐시는 전부 사라지고, 새 표는 본문 표의 실선 테두리(3, 점선 안내 상자 2 가 아님)·문서 최댓값 다음 id
    assert "linesegarray" not in xml and rep["linesegs_removed"] > 0
    new_tbl = xml[i_body:i_12]
    assert 'borderFillIDRef="3"' in new_tbl and 'borderFillIDRef="2"' not in new_tbl and '<hp:tbl id="902"' in new_tbl
    assert 'rowCnt="2" colCnt="3"' in new_tbl and "재학생" in new_tbl
    # 문단·글자 모양은 바탕글(style 0)의 참조
    assert '<hp:p id="0" paraPrIDRef="0" styleIDRef="0" pageBreak="0" columnBreak="0" merged="0"><hp:run charPrIDRef="0"><hp:t>첫 문단' in xml
    # 우리 변환기로 다시 읽힌다
    data, _ = hwpx_docx.convert(out)
    d = Document(io.BytesIO(data))
    texts = [p.text for p in d.paragraphs]
    assert "첫 문단 <검증> & 둘." in texts and "표 뒤 문단" in texts and "특성화 문단" in texts
    assert any(c.text == "1,100" for t in d.tables for r in t.rows for c in r.cells)


def test_fill_only_touches_sections_it_changes_and_can_drop_boxes(tmp_path):
    src = _hwpx(tmp_path, two_sections=True)
    out = tmp_path / "out.hwpx"
    rep = hwpx_fill.fill(src, [{"heading": "3.1. 성과관리 계획", "items": [("text", "성과관리")]}], out, remove_boxes=True)
    assert rep["filled"] == ["3.1. 성과관리 계획"] and rep["boxes_removed"] == 2
    zo = zipfile.ZipFile(out)
    s0, s1 = zo.read("Contents/section0.xml").decode(), zo.read("Contents/section1.xml").decode()
    assert "【작성방법】" not in s0 and "【작성방법】" not in s1 and "성과관리" in s1 and "성과관리" not in s0
    assert rep["sections_changed"] == ["Contents/section0.xml", "Contents/section1.xml"]   # 상자를 지운 구역도 바뀐 구역
    assert "linesegarray" not in s0 and "linesegarray" not in s1
    assert "구분" in s0                                                               # 본문 표는 남는다


def test_table_border_injected_when_form_has_no_body_table(tmp_path):
    header = HEADER.replace('<hh:borderFill id="3">', '<hh:borderFill id="3"><hh:leftBorder type="NONE" width="0.1 mm"/>', 1)
    head2, bf = hwpx_fill._inject_border_fill(HEADER)
    assert bf == "4" and 'itemCnt="4"' in head2 and '<hh:borderFill id="4"' in head2 and head2.count("<hh:borderFill ") == 4
    assert hwpx_fill._solid_border_ids(HEADER) == {"3"}
    assert "3" not in hwpx_fill._solid_border_ids(header)


def _bodies_from_form_plus_edits():
    """작업본은 서식의 변환본이라 서식 글·표가 그대로 들어 있고, 거기에 에이전트가 쓴 것이 더해진다."""
    return [
        {"heading": "1.1. 대학의 여건 분석", "items": [
            ("text", "새로 쓴 문단"),
        ]},
        {"heading": "1) 강점(S)", "items": [("text", "지역 산업 연계 기반")]},          # 에이전트 소제목 — 앞 절(1.1)에 잇는다
        {"heading": "1.2 특성화 방향", "items": [
            ("table", [["구분", "값"]]),                                                   # 서식 표 그대로 — 다시 넣지 않고 지나간다
            ("text", "마무리 문단."),                                                              # 서식 문단 그대로 — 지나간다
            ("text", "표 뒤에 이어 쓴 문단"),
        ]},
        {"heading": "9.9. 없는 절", "items": [("text", "안 들어간다")]},                   # 점 번호 제목이 없으면 건너뛴다
    ]


def test_fill_skips_form_content_and_folds_agent_subheadings(tmp_path):
    src = _hwpx(tmp_path)
    out = tmp_path / "out.hwpx"
    rep = hwpx_fill.fill(src, _bodies_from_form_plus_edits(), out)
    assert rep["filled"] == ["1.1. 대학의 여건 분석", "1.2 특성화 방향"]
    assert rep["folded"] == ["1) 강점(S)"] and rep["skipped"] == ["9.9. 없는 절"]
    assert rep["tables"] == 0 and rep["existing_kept"] == 2 and rep["paragraphs"] == 4
    xml = zipfile.ZipFile(out).read("Contents/section0.xml").decode()
    assert xml.count("<hp:tbl ") == 2                                              # 서식 표 2개 그대로, 새 표 없음
    i_new, i_sub, i_subbody = xml.index("새로 쓴 문단"), xml.index("1) 강점(S)"), xml.index("지역 산업 연계 기반")
    i_box, i_12, i_tbl, i_end, i_after = xml.index("【작성방법】"), xml.index("1.2. 특성화 방향"), xml.index('<hp:tbl id="901"'), xml.index("마무리 문단."), xml.index("표 뒤에 이어 쓴 문단")
    assert i_box < i_new < i_sub < i_subbody < i_12 < i_tbl < i_end < i_after      # 서식 문단 '끝.' 뒤에 이어 썼다


def test_fill_updates_form_table_cells_in_place(tmp_path):
    src = _hwpx(tmp_path)
    out = tmp_path / "out.hwpx"
    rep = hwpx_fill.fill(src, [{"heading": "1.2 특성화 방향", "items": [("table", [["구분", "1,234"]])]}], out)
    assert rep["tables_updated"] == 1 and rep["tables"] == 0
    xml = zipfile.ZipFile(out).read("Contents/section0.xml").decode()
    assert xml.count("<hp:tbl ") == 2 and "1,234" in xml and ">값<" not in xml
    tbl = xml[xml.index('<hp:tbl id="901"'):]
    # 칸의 문단·글자 모양 참조는 서식 것 그대로
    assert '<hp:p paraPrIDRef="0"><hp:run charPrIDRef="0"><hp:t>1,234</hp:t></hp:run></hp:p>' in tbl


def test_heading_with_tab_inside_text_is_found():
    xml = '<hp:p id="1"><hp:run charPrIDRef="0"><hp:t>1.1. 대학의 재정투자 전략<hp:tab width="3036" leader="0" type="1"/></hp:t></hp:run></hp:p>'
    assert hwpx_fill._para_text(xml) == "1.1. 대학의 재정투자 전략"


def test_new_table_uses_working_copy_column_ratios(tmp_path):
    src = _hwpx(tmp_path)
    out = tmp_path / "out.hwpx"
    rep = hwpx_fill.fill(src, [{"heading": "1.1. 대학의 여건 분석", "items": [("widths", [100.0, 300.0]), ("table", [["구분", "내용"], ["가", "나"]])]}], out)
    assert rep["tables"] == 1
    xml = zipfile.ZipFile(out).read("Contents/section0.xml").decode()
    ws = [int(w) for w in re.findall(r'<hp:cellSz width="(\d+)"', xml[xml.index('<hp:tbl id="902"'):])]
    assert ws[0] * 3 - ws[1] <= 3 and ws[0] < ws[1]                       # 1:3 비율, 본문 폭 안


def test_docs_grid_with_merged_cells_maps_values_by_x_position(tmp_path):
    """독스는 병합 표를 잘게 나눈 격자로 낸다(서식 3열 → 독스 5열, 열 너비 동봉). 값 칸을 x 위치로 서식 칸에 맞춰 써 넣는다."""
    def cell(text, bf, r, c, w):
        return (f'<hp:tc borderFillIDRef="{bf}"><hp:subList><hp:p paraPrIDRef="0"><hp:run charPrIDRef="0"><hp:t>{text}</hp:t></hp:run></hp:p></hp:subList>'
                f'<hp:cellAddr rowAddr="{r}" colAddr="{c}"/><hp:cellSpan rowSpan="1" colSpan="1"/><hp:cellSz width="{w}" height="1000"/><hp:cellMargin left="141" right="141" top="141" bottom="141"/></hp:tc>')
    # 서식 표: [지표명(20000)] [단위(4000)] [기준값(6000)] — 두 행
    rows_xml = "".join("<hp:tr>" + cell(a, "3", r, 0, 20000) + cell(b, "3", r, 1, 4000) + cell(c, "3", r, 2, 6000) + "</hp:tr>"
                       for r, (a, b, c) in enumerate([("지표명", "단위", "기준값"), ("AI 이수율", "%", "")]))
    form_tbl = (f'<hp:p id="0" paraPrIDRef="0" styleIDRef="0"><hp:run charPrIDRef="0"><hp:tbl id="901" rowCnt="2" colCnt="3" borderFillIDRef="3"><hp:sz width="30000"/>'
                f'{rows_xml}</hp:tbl></hp:run>{LINESEG}</hp:p>')
    section = _section().replace(_section()[_section().index('<hp:p id="0" paraPrIDRef="0" styleIDRef="0"><hp:run charPrIDRef="0"><hp:tbl id="901"'):_section().index("<hp:p id=\"0\" paraPrIDRef=\"0\" styleIDRef=\"0\"><hp:run charPrIDRef=\"0\"><hp:t>마무리")], form_tbl)
    src = tmp_path / "form.hwpx"
    with zipfile.ZipFile(src, "w") as zf:
        zf.writestr(zipfile.ZipInfo("mimetype"), "application/hwp+zip", compress_type=zipfile.ZIP_STORED)
        zf.writestr("Contents/header.xml", HEADER); zf.writestr("Contents/section0.xml", section)
    # 독스 격자: 5열(지표명 12000+8000 로 갈라짐, 단위 4000, 기준값 4000+2000) — 값 4.6 은 넷째 칸
    bodies = [{"heading": "1.2 특성화 방향", "items": [("widths", [120.0, 80.0, 40.0, 40.0, 20.0]),
                                                  ("table", [["지표명", "", "단위", "기준값", ""], ["AI 이수율", "", "%", "4.6", ""]])]}]
    out = tmp_path / "out.hwpx"
    rep = hwpx_fill.fill(src, bodies, out)
    assert rep["tables_updated"] == 1 and rep["tables"] == 0
    xml = zipfile.ZipFile(out).read("Contents/section0.xml").decode()
    assert xml.count("<hp:tbl ") == 2 and "<hp:t>4.6</hp:t>" in xml
    tbl = xml[xml.index('<hp:tbl id="901"'):]
    assert 'rowAddr="1" colAddr="2"' in tbl and tbl.index("<hp:t>4.6</hp:t>") < tbl.index('rowAddr="1" colAddr="2"')   # 기준값 칸(2열)에 들어갔다


HPF = ('<?xml version="1.0" encoding="UTF-8"?><opf:package xmlns:opf="http://www.idpf.org/2007/opf/"><opf:manifest>'
       '<opf:item id="header" href="Contents/header.xml" media-type="application/xml"/>'
       '<opf:item id="image1" href="BinData/image1.png" media-type="image/png" isEmbeded="1"/>'
       '<opf:item id="section0" href="Contents/section0.xml" media-type="application/xml"/></opf:manifest></opf:package>')


def _png(w: int, h: int) -> bytes:
    import struct
    import zlib
    raw = b"".join(b"\x00" + b"\xff\x00\x00" * w for _ in range(h))
    def chunk(t, d):
        return struct.pack(">I", len(d)) + t + d + struct.pack(">I", zlib.crc32(t + d) & 0xFFFFFFFF)
    return (b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", struct.pack(">IIBBBBB", w, h, 8, 2, 0, 0, 0))
            + chunk(b"IDAT", zlib.compress(raw)) + chunk(b"IEND", b""))


def test_figure_goes_into_bindata_manifest_and_picture_paragraph(tmp_path):
    src = _hwpx(tmp_path)
    with zipfile.ZipFile(src, "a") as zf:
        zf.writestr("Contents/content.hpf", HPF)
    img = _png(40, 20)
    out = tmp_path / "out.hwpx"
    bodies = [{"heading": "1.1. 대학의 여건 분석", "items": [("text", "도식 앞 문단"), ("image", {"data": img, "width_pt": 900.0, "height_pt": 0}),
                                                      ("image", {"data": b"not an image"}), ("text", "도식 뒤 문단")]}]
    rep = hwpx_fill.fill(src, bodies, out)
    assert rep["images"] == 1 and rep["images_skipped"] == 1 and rep["paragraphs"] == 2
    zo = zipfile.ZipFile(out)
    # 이미 있는 image1 과 겹치지 않는 새 항목, 원본 그대로 무압축, 매니페스트 등록
    assert zo.read("BinData/image2.png") == img and zo.getinfo("BinData/image2.png").compress_type == zipfile.ZIP_STORED
    assert zo.read("BinData/image1.png") == b"\x89PNG-not-really"
    hpf = zo.read("Contents/content.hpf").decode()
    assert '<opf:item id="image2" href="BinData/image2.png" media-type="image/png" isEmbeded="1"/></opf:manifest>' in hpf
    xml = zo.read("Contents/section0.xml").decode()
    i_before, i_pic, i_after = xml.index("도식 앞 문단"), xml.index("<hp:pic "), xml.index("도식 뒤 문단")
    assert i_before < i_pic < i_after < xml.index("1.2. 특성화 방향")
    pic = xml[i_pic:xml.index("</hp:pic>")]
    # 900pt 는 본문 폭(59528-5669*2=48190)으로 줄이고 가로:세로 2:1 유지, 원본 크기는 픽셀×75
    assert 'binaryItemIDRef="image2"' in pic and '<hp:sz width="48190" widthRelTo="ABSOLUTE" height="24095"' in pic
    assert '<hp:orgSz width="3000" height="1500"/>' in pic and 'treatAsChar="1"' in pic
    assert 'xmlns:hc="http://www.hancom.co.kr/hwpml/2011/core"' in xml[:400]
    from xml.dom.minidom import parseString
    parseString(zo.read("Contents/section0.xml"))


def test_figure_without_manifest_is_reported_not_inserted(tmp_path):
    rep = hwpx_fill.fill(_hwpx(tmp_path), [{"heading": "1.2 특성화 방향", "items": [("image", {"data": _png(4, 4), "width_pt": 100.0})]}],
                         tmp_path / "out.hwpx")
    assert rep["images"] == 0 and rep["images_skipped"] == 1
    assert "BinData/image2.png" not in zipfile.ZipFile(tmp_path / "out.hwpx").namelist()


def test_control_characters_never_reach_the_section_xml(tmp_path):
    """독스의 문단 안 줄 바꿈(\\x0b)은 새 문단으로, 그 밖의 금지 제어문자는 지운다 — 한 글자만 섞여도 한글이 문서를 못 연다."""
    from xml.dom.minidom import parseString

    out = tmp_path / "out.hwpx"
    bodies = [{"heading": "1.1. 대학의 여건 분석", "items": [("text", "첫 줄\x0b둘째 줄\x01끝"), ("table", [["구분", "값\x0b둘"], ["가\x02", "1"]])]}]
    rep = hwpx_fill.fill(_hwpx(tmp_path), bodies, out)
    xml = zipfile.ZipFile(out).read("Contents/section0.xml")
    parseString(xml)
    text = xml.decode()
    assert rep["paragraphs"] == 2 and ">첫 줄<" in text and ">둘째 줄끝<" in text and "\x0b" not in text and "\x01" not in text
    assert ">값<" in text and ">둘<" in text and ">가<" in text


def test_self_check_passes_on_fill_and_catches_broken_references(tmp_path):
    out = tmp_path / "out.hwpx"
    rep = hwpx_fill.fill(_hwpx(tmp_path), BODIES, out)
    assert rep["checks"] == []
    # 없는 글자 모양·틀린 rowCnt·itemCnt 를 넣은 사본은 잡힌다
    bad = tmp_path / "bad.hwpx"
    with zipfile.ZipFile(out) as zin, zipfile.ZipFile(bad, "w") as zout:
        for info in zin.infolist():
            data = zin.read(info.filename)
            if info.filename == "Contents/section0.xml":
                data = data.decode().replace('charPrIDRef="1"', 'charPrIDRef="77"', 1).replace('rowCnt="1" colCnt="2"', 'rowCnt="3" colCnt="2"', 1).encode()
            if info.filename == "Contents/header.xml":
                data = data.decode().replace('<hh:charProperties itemCnt="2">', '<hh:charProperties itemCnt="5">').encode()
            zout.writestr(info, data)
    probs = hwpx_fill.check(bad)
    assert any("charPrIDRef 77" in p for p in probs) and any("rowCnt 3" in p for p in probs) and any("itemCnt 5" in p for p in probs)


def test_new_table_inherits_form_table_roles(tmp_path):
    """서식 본문 표가 머리 행 음영(borderFill 4·굵은 글자 1)을 쓰면 새 표 머리 행도 그 모양, 본문 칸은 서식 본문 칸 모양.
    첫 열이 라벨 같으면(짧은 글 + 나머지 수치) 라벨 열 모양도 — 음영이 있을 때만."""
    head = HEADER.replace('<hh:borderFills itemCnt="3">', '<hh:borderFills itemCnt="5">').replace(
        "</hh:borderFills>",
        '<hh:borderFill id="4"><hh:leftBorder type="SOLID" width="0.12 mm"/><hh:rightBorder type="SOLID" width="0.12 mm"/>'
        '<hh:topBorder type="SOLID" width="0.12 mm"/><hh:bottomBorder type="SOLID" width="0.12 mm"/>'
        '<hc:fillBrush><hc:winBrush faceColor="#DFE6F7" hatchColor="#000000" alpha="0"/></hc:fillBrush></hh:borderFill>'
        '<hh:borderFill id="5"><hh:leftBorder type="SOLID" width="0.12 mm"/><hh:rightBorder type="SOLID" width="0.12 mm"/>'
        '<hh:topBorder type="SOLID" width="0.12 mm"/><hh:bottomBorder type="SOLID" width="0.12 mm"/>'
        '<hc:fillBrush><hc:winBrush faceColor="#F2F2F2" hatchColor="#000000" alpha="0"/></hc:fillBrush></hh:borderFill></hh:borderFills>')
    shaded = (f'<hp:p id="0" paraPrIDRef="0" styleIDRef="0"><hp:run charPrIDRef="0"><hp:tbl id="905" rowCnt="2" colCnt="2" borderFillIDRef="3"><hp:sz width="40000"/>'
              f'<hp:tr>{_cell("항목", "4", 0, 0).replace("charPrIDRef=\"0\"", "charPrIDRef=\"1\"")}{_cell("값", "4", 0, 1).replace("charPrIDRef=\"0\"", "charPrIDRef=\"1\"")}</hp:tr>'
              f'<hp:tr>{_cell("학생", "5", 1, 0)}{_cell("10", "3", 1, 1)}</hp:tr></hp:tbl></hp:run>{LINESEG}</hp:p>')
    sec = _section().replace('<hp:p id="0" paraPrIDRef="0" styleIDRef="0"><hp:run charPrIDRef="0"><hp:t>마무리 문단.', shaded + '<hp:p id="0" paraPrIDRef="0" styleIDRef="0"><hp:run charPrIDRef="0"><hp:t>마무리 문단.')
    src = tmp_path / "form.hwpx"
    with zipfile.ZipFile(src, "w") as zf:
        zf.writestr(zipfile.ZipInfo("mimetype"), "application/hwp+zip", compress_type=zipfile.ZIP_STORED)
        zf.writestr("Contents/header.xml", head)
        zf.writestr("Contents/section0.xml", sec)
    out = tmp_path / "out.hwpx"
    rep = hwpx_fill.fill(src, [{"heading": "1.1. 대학의 여건 분석", "items": [("table", [["구분", "2025", "2026"], ["재학생", "1,000", "1,100"], ["교원", "50", "55"]])]}], out)
    assert rep["tables"] == 1 and rep["checks"] == []
    xml = zipfile.ZipFile(out).read("Contents/section0.xml").decode()
    new = xml[xml.index("구분") - 900: xml.index(">55<") + 50]
    heads = re.findall(r'borderFillIDRef="(\d+)"><hp:subList[^>]*><hp:p [^>]*><hp:run charPrIDRef="(\d+)"><hp:t>(구분|2025|재학생|1,000|교원)<', new)
    got = {t: (b, c) for b, c, t in heads}
    assert got["구분"] == ("4", "1") and got["2025"] == ("4", "1")          # 머리 행 = 서식 머리 행
    assert got["재학생"] == ("5", "0") and got["교원"] == ("5", "0")        # 라벨 열 = 서식 라벨 열(음영 있음)
    assert got["1,000"] == ("3", "0")                                        # 본문 칸


def test_new_paragraphs_follow_form_marker_styles_and_writing_slot(tmp_path):
    """서식의 '◦ 짧은 굵은 줄'(문단 1·글자 1 굵게)은 소제목 모양, 부호 없는 문장은 제목 뒤 빈 쓰기 자리(문단 1·글자 0) 모양."""
    extra = ('<hp:p id="0" paraPrIDRef="1" styleIDRef="0"><hp:run charPrIDRef="1"><hp:t>◦ 추진 전략</hp:t></hp:run></hp:p>'
             '<hp:p id="0" paraPrIDRef="0" styleIDRef="0"><hp:run charPrIDRef="0"><hp:t>◦ 이 문장은 소제목이 아니라 서른 자를 훌쩍 넘기는 긴 본문 문장이다</hp:t></hp:run></hp:p>')
    sec = _section().replace('<hp:p id="0" paraPrIDRef="1" styleIDRef="0"><hp:run charPrIDRef="1"><hp:t>1.2. 특성화 방향</hp:t>',
                             '<hp:p id="0" paraPrIDRef="1" styleIDRef="0"><hp:run charPrIDRef="0"/></hp:p>'
                             + extra + '<hp:p id="0" paraPrIDRef="1" styleIDRef="0"><hp:run charPrIDRef="1"><hp:t>1.2. 특성화 방향</hp:t>')
    src = tmp_path / "form.hwpx"
    with zipfile.ZipFile(src, "w") as zf:
        zf.writestr(zipfile.ZipInfo("mimetype"), "application/hwp+zip", compress_type=zipfile.ZIP_STORED)
        zf.writestr("Contents/header.xml", HEADER)
        zf.writestr("Contents/section0.xml", sec)
    out = tmp_path / "out.hwpx"
    lines = "◦ 인력 양성\n◦ 지역 산업 수요와 연계해 교육과정을 바꾸고 현장 실습을 늘리는 방안을 단계적으로 추진한다\n부호 없는 본문 문장이다."
    rep = hwpx_fill.fill(src, [{"heading": "1.1. 대학의 여건 분석", "items": [("text", lines)]}], out)
    assert rep["paragraphs"] == 3 and rep["checks"] == []
    xml = zipfile.ZipFile(out).read("Contents/section0.xml").decode()

    def style(text):
        m = re.search(r'<hp:p id="0" paraPrIDRef="(\d+)"[^>]*><hp:run charPrIDRef="(\d+)"><hp:t>' + re.escape(text), xml)
        return m.groups()
    assert style("◦ 인력 양성") == ("1", "1")                              # 짧은 굵은 ◦ 줄 = 서식 소제목 모양
    assert style("◦ 지역 산업") == ("0", "0")                              # 긴 ◦ 문장 = 서식 ◦ 본문 모양
    assert style("부호 없는 본문") == ("1", "0")                           # 부호 없음 = 빈 쓰기 자리 모양


def _budget_form(tmp_path, merged: bool = False):
    span = '<hp:cellSpan rowSpan="1" colSpan="1"/>'
    rows = (f'<hp:tr>{_cell("비목", "3", 0, 0)}{_cell("금액", "3", 0, 1)}</hp:tr>'
            f'<hp:tr>{_cell("", "3", 1, 0)}{_cell("", "3", 1, 1)}</hp:tr>'
            f'<hp:tr>{_cell("합계", "3", 2, 0)}{_cell("", "3", 2, 1)}</hp:tr>')
    if merged:
        rows = rows.replace(span, '<hp:cellSpan rowSpan="2" colSpan="1"/>', 1)
    tbl = (f'<hp:p id="0" paraPrIDRef="0" styleIDRef="0"><hp:run charPrIDRef="0"><hp:tbl id="910" rowCnt="3" colCnt="2" borderFillIDRef="3">'
           f'<hp:sz width="40000" height="3000"/>{rows}</hp:tbl></hp:run>{LINESEG}</hp:p>')
    sec = _section().replace('<hp:p id="0" paraPrIDRef="0" styleIDRef="0"><hp:run charPrIDRef="0"><hp:t>마무리 문단.', tbl + '<hp:p id="0" paraPrIDRef="0" styleIDRef="0"><hp:run charPrIDRef="0"><hp:t>마무리 문단.')
    src = tmp_path / "form.hwpx"
    with zipfile.ZipFile(src, "w") as zf:
        zf.writestr(zipfile.ZipInfo("mimetype"), "application/hwp+zip", compress_type=zipfile.ZIP_STORED)
        zf.writestr("Contents/header.xml", HEADER)
        zf.writestr("Contents/section0.xml", sec)
    return src


def test_form_table_grows_rows_by_cloning_blank_row(tmp_path):
    """작업본 행이 더 많으면 새 표가 아니라 서식 표의 빈 입력 행을 복제해 늘리고 합계 행은 제자리에(행 번호·rowCnt·높이 갱신)."""
    out = tmp_path / "out.hwpx"
    work = [["비목", "금액"], ["인건비", "30"], ["장학금", "20"], ["운영비", "10"], ["합계", "60"]]
    rep = hwpx_fill.fill(_budget_form(tmp_path), [{"heading": "1.2 특성화 방향", "items": [("table", work)]}], out)
    assert rep["tables_grown"] == 1 and rep["tables"] == 0 and rep["checks"] == []
    xml = zipfile.ZipFile(out).read("Contents/section0.xml").decode()
    tbl = xml[xml.index('<hp:tbl id="910"'):xml.index("</hp:tbl>", xml.index('<hp:tbl id="910"'))]
    assert 'rowCnt="5"' in tbl and 'height="5000"' in tbl
    order = re.findall(r"<hp:t>([^<]*)</hp:t>", tbl)
    assert order == ["비목", "금액", "인건비", "30", "장학금", "20", "운영비", "10", "합계", "60"]
    assert re.findall(r'rowAddr="(\d+)"', tbl) == ["0", "0", "1", "1", "2", "2", "3", "3", "4", "4"]


def _group_form(tmp_path):
    """첫 열 '영역' 칸이 두 행을 묶는 표: 머리 | 영역A(rowSpan 2)·빈 행 둘 | 합계(colSpan 2)."""
    def cell(text, r, c, rs=1, cs=1, h=1000):
        return (f'<hp:tc borderFillIDRef="3"><hp:subList><hp:p paraPrIDRef="0"><hp:run charPrIDRef="0"><hp:t>{text}</hp:t></hp:run>{LINESEG}</hp:p></hp:subList>'
                f'<hp:cellAddr rowAddr="{r}" colAddr="{c}"/><hp:cellSpan rowSpan="{rs}" colSpan="{cs}"/><hp:cellSz width="10000" height="{h}"/>'
                f'<hp:cellMargin left="141" right="141" top="141" bottom="141"/></hp:tc>')
    rows = (f'<hp:tr>{cell("영역", 0, 0)}{cell("항목", 0, 1)}{cell("금액", 0, 2)}</hp:tr>'
            f'<hp:tr>{cell("영역A", 1, 0, rs=2, h=2000)}{cell("", 1, 1)}{cell("", 1, 2)}</hp:tr>'
            f'<hp:tr>{cell("", 2, 1)}{cell("", 2, 2)}</hp:tr>'
            f'<hp:tr>{cell("합계", 3, 0, cs=2)}{cell("", 3, 2)}</hp:tr>')
    tbl = (f'<hp:p id="0" paraPrIDRef="0" styleIDRef="0"><hp:run charPrIDRef="0"><hp:tbl id="920" rowCnt="4" colCnt="3" borderFillIDRef="3">'
           f'<hp:sz width="30000" height="4000"/>{rows}</hp:tbl></hp:run>{LINESEG}</hp:p>')
    sec = _section().replace('<hp:p id="0" paraPrIDRef="0" styleIDRef="0"><hp:run charPrIDRef="0"><hp:t>마무리 문단.', tbl + '<hp:p id="0" paraPrIDRef="0" styleIDRef="0"><hp:run charPrIDRef="0"><hp:t>마무리 문단.')
    src = tmp_path / "form.hwpx"
    with zipfile.ZipFile(src, "w") as zf:
        zf.writestr(zipfile.ZipInfo("mimetype"), "application/hwp+zip", compress_type=zipfile.ZIP_STORED)
        zf.writestr("Contents/header.xml", HEADER)
        zf.writestr("Contents/section0.xml", sec)
    return src


def test_rows_grow_inside_a_vertical_merge_group(tmp_path):
    """묶음 안 빈 행을 복제하면 묶음 머리 칸(영역A)의 rowSpan·높이가 늘어 격자가 유지된다. 작업본은 독스 격자(가려진 열은 빈칸)로 온다."""
    out = tmp_path / "out.hwpx"
    work = [["영역", "항목", "금액"], ["영역A", "가", "1"], ["", "나", "2"], ["", "다", "3"], ["", "라", "4"], ["합계", "", "10"]]
    rep = hwpx_fill.fill(_group_form(tmp_path), [{"heading": "1.2 특성화 방향", "items": [("table", work)]}], out)
    assert rep["tables_grown"] == 1 and rep["checks"] == []
    xml = zipfile.ZipFile(out).read("Contents/section0.xml").decode()
    tbl = xml[xml.index('<hp:tbl id="920"'):xml.index("</hp:tbl>", xml.index('<hp:tbl id="920"'))]
    assert 'rowCnt="6"' in tbl and 'rowSpan="4"' in tbl and 'height="6000"' in tbl
    assert re.findall(r"<hp:t>([^<]*)</hp:t>", tbl) == ["영역", "항목", "금액", "영역A", "가", "1", "나", "2", "다", "3", "라", "4", "합계", "10"]


def test_value_in_a_merged_away_column_blocks_growth(tmp_path):
    out = tmp_path / "out.hwpx"
    work = [["영역", "항목", "금액"], ["영역A", "가", "1"], ["영역B", "나", "2"], ["", "다", "3"], ["합계", "", "6"]]
    rep = hwpx_fill.fill(_group_form(tmp_path), [{"heading": "1.2 특성화 방향", "items": [("table", work)]}], out)
    assert rep["tables_grown"] == 0 and rep["tables"] == 1                   # 묶음에 가려진 칸에 다른 값(영역B) — 새 표로


def test_long_marker_paragraph_gets_hanging_indent_and_auto_tab(tmp_path):
    """두 줄로 넘어갈 '□ …' 문단은 바탕 모양을 복제해 내어쓰기(intent 음수, case·default 두 벌)와 자동 탭을 붙인 새 paraPr 로,
    부호 뒤는 탭. 같은 사양은 한 번만 등록하고 itemCnt 를 올린다. 짧은 부호 줄과 자동 탭이 없는 서식은 그대로."""
    pr = ('<hh:paraPr id="{id}" tabPrIDRef="0"><hh:align horizontal="JUSTIFY"/><hp:switch><hp:case hp:required-namespace="x">'
          '<hh:margin><hc:intent value="0" unit="HWPUNIT"/><hc:left value="0" unit="HWPUNIT"/></hh:margin></hp:case><hp:default>'
          '<hh:margin><hc:intent value="0" unit="HWPUNIT"/><hc:left value="0" unit="HWPUNIT"/></hh:margin></hp:default></hp:switch></hh:paraPr>')
    head = (HEADER.replace('<hh:paraProperties itemCnt="2">', '<hh:paraProperties itemCnt="2">')
            .replace('<hh:paraPr id="0"><hh:align horizontal="JUSTIFY"/><hh:heading type="NONE" level="0"/></hh:paraPr>', pr.format(id=0))
            .replace("<hh:refList>", '<hh:refList><hh:tabProperties itemCnt="2"><hh:tabPr id="0" autoTabLeft="0" autoTabRight="0"/>'
                                     '<hh:tabPr id="1" autoTabLeft="1" autoTabRight="0"/></hh:tabProperties>')
            .replace('xmlns:hc="http://www.hancom.co.kr/hwpml/2011/core"', 'xmlns:hc="http://www.hancom.co.kr/hwpml/2011/core" xmlns:hp="http://www.hancom.co.kr/hwpml/2011/paragraph"'))
    src = tmp_path / "form.hwpx"
    with zipfile.ZipFile(src, "w") as zf:
        zf.writestr(zipfile.ZipInfo("mimetype"), "application/hwp+zip", compress_type=zipfile.ZIP_STORED)
        zf.writestr("Contents/header.xml", head)
        zf.writestr("Contents/section0.xml", _section())
    out = tmp_path / "out.hwpx"
    long1 = "□ 지역 산업 수요와 연계해 교육과정을 바꾸고 현장 실습을 늘리는 방안을 단계적으로 추진한다"
    long2 = "□ 두 번째 긴 부호 문단도 같은 사양이므로 새 문단 모양을 다시 등록하지 않고 같은 것을 쓴다"
    rep = hwpx_fill.fill(src, [{"heading": "1.1. 대학의 여건 분석", "items": [("text", f"{long1}\n{long2}\n□ 짧은 줄")]}], out)
    assert rep["checks"] == []
    z = zipfile.ZipFile(out)
    h = z.read("Contents/header.xml").decode()
    assert '<hh:paraProperties itemCnt="3">' in h
    new = re.search(r'<hh:paraPr id="2" tabPrIDRef="1">.*?</hh:paraPr>', h, re.S).group(0)
    hang = int(re.search(r'<hp:default><hh:margin><hc:intent value="(-\d+)"', new).group(1))
    case = int(re.search(r'<hp:case[^>]*><hh:margin><hc:intent value="(-\d+)"', new).group(1))
    assert hang == -1500 and case == -750                                  # 10pt × (□ 1em + 0.5em), case 는 절반
    xml = z.read("Contents/section0.xml").decode()
    assert xml.count('paraPrIDRef="2"') == 2 and '<hp:t>□<hp:tab width="1500" leader="0" type="1"/>지역 산업' in xml
    assert '<hp:t>□ 짧은 줄</hp:t>' in xml


def test_no_hanging_registry_without_auto_tab(tmp_path):
    out = tmp_path / "out.hwpx"
    rep = hwpx_fill.fill(_hwpx(tmp_path), [{"heading": "1.1. 대학의 여건 분석", "items": [("text", "□ 서른 자를 넘기는 긴 부호 문단이지만 자동 탭 설정이 없는 서식이다 그러니 그대로")]}], out)
    assert rep["checks"] == [] and "<hp:tab" not in zipfile.ZipFile(out).read("Contents/section0.xml").decode()


def test_map_cells_by_span_when_docs_grid_is_finer():
    """독스가 병합 칸을 잘게 나눠 격자가 서식과 어긋나도, 작업본 칸 가운데가 든 서식 칸으로 대응한다(리허설: 11열 서식·21열 독스)."""
    C = hwpx_fill._Cell
    def cell(col, width, text=""):
        return C(tc=(0, 0), inner=(0, 0), col=col, text=text, nested=False, p_open="", char_ref="0", width=width)
    form = hwpx_fill._FormTable(para=(0, 0), rows=[[cell(0, 300, "사업유형"), cell(1, 700, "단독형 / 연합형")],
                                                   [cell(0, 300, "사업목표"), cell(1, 700)]], col_cnt=2)
    rows = [["사업유형", "", "단독형", ""], ["사업목표", "", "AI 인재 양성", "지역 연계"]]
    pairs = hwpx_fill._map_cells(form, rows, widths=[3, 2, 3, 2])
    got = {(fc.col, fc.text): t for fc, t in pairs}
    assert got[(1, "단독형 / 연합형")] == "단독형" and got[(1, "")] == "AI 인재 양성 지역 연계"
