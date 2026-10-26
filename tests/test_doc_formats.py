"""학교 원본 같은 파일 판정(내용 지문)과 처리 대상 문서 형식 — 가짜 파일만 쓴다.

- 이름·크기가 같아도 내용이 다르면 중복이 아니다. 내용까지 같으면 중복. 지문을 모르면 예전처럼 이름·크기로.
- DGX 목록(165)이 낸 파일 사실 장부를 VM(170)이 받아 같은 판정을 한다.
- 새 처리 대상 형식(xlsm·pptm·xlt·ods·rtf·csv·txt·html·htm·mht·eml)마다 본문이 읽힌다.
- 확장자 없음·「…다운로드」는 파일 머리로 실제 형식을 보아, 문서면 처리하고 웹 페이지 부속(스크립트)은 빼낸다.
"""
from __future__ import annotations

import importlib.util
import os
import stat
import sys
import zipfile
from pathlib import Path

import pytest

from zzaimy import archive_dups, doc_formats
from zzaimy.app import pipeline

ROOT = Path(__file__).resolve().parents[1]


def _f(rel: str, size: int, mtime: int = 1, ext: str | None = None) -> dict:
    name = rel.rsplit("/", 1)[-1]
    return {"rel": rel, "filename": name, "size": size, "mtime": mtime,
            "ext": doc_formats.file_ext(name) if ext is None else ext}


def _docx(path: Path, *paras: str) -> Path:
    import docx
    d = docx.Document()
    for p in paras:
        d.add_paragraph(p)
    d.save(str(path))
    return path


def _xlsx(path: Path, rows: list[list[str]]) -> Path:
    import openpyxl
    wb = openpyxl.Workbook()
    for r in rows:
        wb.active.append(r)
    wb.save(str(path))
    return path


# ── 같은 파일 판정 ─────────────────────────────────────────────────────────

def test_same_name_and_size_but_different_content_is_not_a_duplicate():
    files = [_f("가/신청서.hwp", 100), _f("나/신청서.hwp", 100), _f("다/신청서.hwp", 100), _f("라/계획.pdf", 7)]
    facts = {"가/신청서.hwp": {"rel": "가/신청서.hwp", "size": 100, "sha256": "aaa"},
             "나/신청서.hwp": {"rel": "나/신청서.hwp", "size": 100, "sha256": "bbb"},     # 같은 양식에 다른 내용
             "다/신청서.hwp": {"rel": "다/신청서.hwp", "size": 100, "sha256": "aaa"}}     # 첫 파일과 같은 내용
    got = archive_dups.assign_dups(files, facts)
    assert [f["dup_of"] for f in files] == ["", "", "가/신청서.hwp", ""]
    assert got == {"split_by_content": 1, "unknown_digest": 0}


def test_unknown_digest_falls_back_to_name_and_size():
    files = [_f("가/양식.docx", 50), _f("나/양식.docx", 50)]
    got = archive_dups.assign_dups(files, {})
    assert files[1]["dup_of"] == "가/양식.docx" and got["unknown_digest"] == 1
    # 크기가 바뀐 판의 옛 지문은 쓰지 않는다
    stale = {"나/양식.docx": {"rel": "나/양식.docx", "size": 49, "sha256": "zzz"}}
    archive_dups.assign_dups(files, stale)
    assert files[1]["dup_of"] == "가/양식.docx"


def test_dgx_facts_ledger_gives_vm_the_same_decision(tmp_path):
    """DGX: 파일을 읽어 지문·실제 형식을 장부에 → VM: 목록(이름·크기)과 장부만으로 같은 dup_of·형식."""
    root = tmp_path / "data"
    for rel, body in (("갑/결과보고.hwp", b"A" * 64), ("을/결과보고.hwp", b"B" * 64), ("병/결과보고.hwp", b"A" * 64),
                      ("정/사진.jpg", b"J" * 10), ("무/사진.jpg", b"K" * 10)):
        (root / rel).parent.mkdir(parents=True, exist_ok=True)
        (root / rel).write_bytes(body)
    (root / "기").mkdir()
    _docx(root / "기/공고문다운로드", "모집 공고 본문")
    (root / "기/jquery.min.js.다운로드").write_text("!function(){return 1}();", encoding="utf-8")

    def listing():
        out = []
        for p in sorted(root.rglob("*")):
            if p.is_file():
                st = p.stat()
                out.append(_f(str(p.relative_to(root)), st.st_size, int(st.st_mtime)))
        return out

    dgx = listing()
    facts: dict = {}
    assert archive_dups.apply_real_ext(dgx, facts, root=root, use_mtime=True) == 1
    got = archive_dups.fill_digests(dgx, facts, root)
    assert got["hashed"] == 3                                         # 문서 묶음만 읽는다(사진 묶음은 읽지 않음)
    archive_dups.assign_dups(dgx, facts, use_mtime=True)
    ledger = tmp_path / "archive_facts.jsonl"
    archive_dups.save_facts(ledger, facts, keep={f["rel"] for f in dgx})

    vm = listing()
    loaded = archive_dups.load_facts(ledger)
    archive_dups.apply_real_ext(vm, loaded)
    archive_dups.assign_dups(vm, loaded)
    by = {f["rel"]: f for f in vm}
    assert {f["rel"]: f["dup_of"] for f in vm} == {f["rel"]: f["dup_of"] for f in dgx}
    assert by["병/결과보고.hwp"]["dup_of"] == "갑/결과보고.hwp" and by["을/결과보고.hwp"]["dup_of"] == ""
    assert by["정/사진.jpg"]["dup_of"] == "무/사진.jpg"                # 사진은 예전처럼 이름·크기
    assert by["기/공고문다운로드"]["ext"] == "docx"                     # 실제 형식이 문서 → 처리 대상
    assert by["기/jquery.min.js.다운로드"]["ext"] == "다운로드"          # 스크립트는 문서가 아니다

    # 다음 회차: 같은 판은 다시 읽지 않는다
    again = listing()
    assert archive_dups.fill_digests(again, loaded, root)["hashed"] == 0


def test_hash_budget_leaves_rest_for_next_round(tmp_path):
    for d in ("a", "b"):
        (tmp_path / d).mkdir()
        (tmp_path / d / "x.pdf").write_bytes(d.encode() * 8)
    files = [_f("a/x.pdf", 8), _f("b/x.pdf", 8)]
    got = archive_dups.fill_digests(files, {}, tmp_path, budget_s=-1)
    assert got == {"hashed": 0, "pending": 2, "failed": 0}


def test_vm_sync_uses_dgx_facts(monkeypatch):
    spec = importlib.util.spec_from_file_location("vm_sync_facts", ROOT / "scripts/170_vm_sync.py")
    job = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(job)
    files = [_f("가/보고.pdf", 9), _f("나/보고.pdf", 9), _f("다/첨부다운로드", 5)]
    facts = {"가/보고.pdf": {"rel": "가/보고.pdf", "size": 9, "sha256": "1"},
             "나/보고.pdf": {"rel": "나/보고.pdf", "size": 9, "sha256": "2"},
             "다/첨부다운로드": {"rel": "다/첨부다운로드", "size": 5, "real": "pdf"}}
    got = job.mark_dups(files, facts)
    assert [f["dup_of"] for f in files] == ["", "", ""] and files[2]["ext"] == "pdf"
    assert got["split_by_content"] == 1 and got["real_ext"] == 1


def test_audit_covers_only_same_content_group(tmp_path, monkeypatch):
    """175 — 같은 내용 묶음의 다른 파일이 들어왔으면 처리된 것, 이름·크기만 같은 다른 내용은 따로 센다."""
    from zzaimy.app import archive
    from zzaimy.app.db import Database
    spec = importlib.util.spec_from_file_location("audit_dups", ROOT / "scripts/175_pipeline_audit.py")
    job = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(job)
    plat = tmp_path / "platform"
    plat.mkdir()
    monkeypatch.setattr(job, "PLAT", plat)
    db = Database(plat / "platform.db")
    rows = [dict(rel="가/계획.hwp", size=10, mtime=1, ext="hwp", dup_of=""),
            dict(rel="나/계획.hwp", size=10, mtime=1, ext="hwp", dup_of="가/계획.hwp"),   # 같은 내용 — 이쪽이 들어옴
            dict(rel="다/계획.hwp", size=10, mtime=1, ext="hwp", dup_of="")]              # 이름·크기만 같은 다른 내용
    d7 = db.add_document("계획.hwp", "x", doc_type="grant")         # 반입 기록이 가리키는 문서(지금 있는 문서만 잇는다)
    archive.load(db, rows, {"나/계획.hwp": d7})
    with db._conn() as conn:
        conn.execute("UPDATE archive_files SET seen_at = '2000-01-01 00:00'")
    monkeypatch.setenv("ZZAIMY_PLATFORM_SQLITE_PATH", str(plat / "platform.db"))
    monkeypatch.setattr(sys, "argv", ["175_pipeline_audit.py"])
    assert job.main() == 0
    import json
    out = json.loads((plat / "pipeline_audit.json").read_text(encoding="utf-8"))
    assert out["targets"] == 2 and out["in_docbox"] == 1 and out["missed"] == 1
    assert out["composition"]["같은 파일 중복"] == 1


# ── 형식 목록 한 곳 ───────────────────────────────────────────────────────

def test_format_lists_come_from_one_place():
    from zzaimy.app import archive_state, vocab
    assert archive_state.DOC_EXT is doc_formats.DOC_EXT
    assert set(doc_formats.DOC_EXT) <= set(vocab.original_formats())       # 보관소 화면 형식 갈래에 모두 든다
    for e in ("xlsm", "pptm", "xlt", "ods", "rtf", "csv", "txt", "html", "htm", "mht", "eml"):
        assert e in doc_formats.DOC_EXT and vocab.format_group(e)
    spec = importlib.util.spec_from_file_location("dgx_parse_fmt", ROOT / "scripts/167_dgx_parse.py")
    job = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(job)
    assert job.TEXT_EXT == set(doc_formats.DOC_EXT)


def test_needs_sniff_and_real_format(tmp_path):
    assert doc_formats.needs_sniff("공고문다운로드") and doc_formats.needs_sniff("보고서.pdf다운로드")
    assert doc_formats.needs_sniff("README") and not doc_formats.needs_sniff("계획.hwp")
    p = tmp_path / "첨부"
    p.write_bytes(b"%PDF-1.7\n...")
    assert doc_formats.sniff_doc_ext(p) == "pdf"
    t = tmp_path / "style.css.다운로드"
    t.write_text("body{margin:0}", encoding="utf-8")
    assert doc_formats.sniff_doc_ext(t) == ""                              # 글 파일은 문서로 치지 않는다
    assert pipeline.sniff_suffix is doc_formats.sniff_suffix


# ── 새 형식마다 본문이 읽힌다 ─────────────────────────────────────────────

def _read(path: Path) -> str:
    return pipeline.DocumentProcessor()._parse_inner(path)


def test_xlsm_and_pptm(tmp_path, monkeypatch):
    monkeypatch.setenv("ZZAIMY_NO_OCR_FALLBACK", "1")
    x = _xlsx(tmp_path / "예산.xlsx", [["항목", "금액"], ["교육비", "300"]]).rename(tmp_path / "예산.xlsm")
    assert "교육비" in _read(x)
    from pptx import Presentation
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[5])
    s.shapes.title.text = "성과 발표 요약"
    src = tmp_path / "발표.pptx"
    prs.save(str(src))
    pptm = tmp_path / "발표.pptm"                                          # 매크로 묶음의 내용 형식으로 바꿔 쓴다
    with zipfile.ZipFile(src) as zin, zipfile.ZipFile(pptm, "w") as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "[Content_Types].xml":
                data = data.replace(b"application/vnd.openxmlformats-officedocument.presentationml.presentation.main+xml",
                                    b"application/vnd.ms-powerpoint.presentation.macroEnabled.main+xml")
                assert b"macroEnabled" in data
            zout.writestr(item, data)
    assert "성과 발표 요약" in _read(pptm)


def test_text_csv_html_mht_eml(tmp_path):
    t = tmp_path / "메모.txt"
    t.write_bytes("회의 결과 정리".encode("cp949"))                           # 한글 윈도 옛 인코딩
    assert _read(t) == "회의 결과 정리"
    c = tmp_path / "명단.csv"
    c.write_bytes("과정,인원\n기초,12\n".encode("cp949"))
    assert _read(c) == "과정 | 인원\n기초 | 12"
    h = tmp_path / "안내.html"
    h.write_text('<html><head><meta charset="utf-8"><title>사업 안내</title><style>p{color:red}</style>'
                 '<script>var x="숨은 글";</script></head><body><p>참여 대상</p><table><tr><td>구분</td><td>내용</td></tr>'
                 '</table></body></html>', encoding="utf-8")
    got = _read(h)
    assert got.splitlines()[0] == "사업 안내" and "참여 대상" in got and "구분 | 내용" in got
    assert "숨은 글" not in got and "color" not in got and "<" not in got
    hm = tmp_path / "옛 안내.htm"
    hm.write_bytes('<html><head><meta charset="euc-kr"></head><body>공지 사항</body></html>'.encode("cp949"))
    assert _read(hm) == "공지 사항"
    mht = tmp_path / "보관.mht"
    mht.write_text('MIME-Version: 1.0\r\nContent-Type: multipart/related; boundary="B"\r\n\r\n--B\r\n'
                   'Content-Type: text/html; charset="utf-8"\r\nContent-Transfer-Encoding: 8bit\r\n\r\n'
                   '<html><body><h1>성과 공유회</h1><p>일정 안내</p></body></html>\r\n--B\r\n'
                   'Content-Type: image/png\r\nContent-Location: a.png\r\n\r\nPNGDATA\r\n--B--\r\n', encoding="utf-8")
    assert _read(mht) == "성과 공유회\n일정 안내"
    from email.message import EmailMessage
    m = EmailMessage()
    m["Subject"] = "자료 제출 요청"
    m["From"] = "a@example.org"
    m.set_content("기한 안에 제출 바랍니다.")
    m.add_attachment(b"secret-attachment-body", maintype="application", subtype="octet-stream", filename="첨부.bin")
    e = tmp_path / "요청.eml"
    e.write_bytes(m.as_bytes())
    got = _read(e)
    assert got.startswith("제목: 자료 제출 요청") and "기한 안에 제출" in got and "secret" not in got


@pytest.fixture
def fake_soffice(tmp_path, monkeypatch):
    """LibreOffice 대신 — 받은 대상과 같은 이름의 새 형식 파일을 미리 만든 견본으로 낸다(변환 경로만 확인)."""
    made = tmp_path / "made"
    made.mkdir()
    _docx(made / "out.docx", "서식 있는 글 본문")
    _xlsx(made / "out.xlsx", [["열린 표", "값"]])
    exe = tmp_path / "soffice"
    exe.write_text(f"""#!{sys.executable}
import shutil, sys
from pathlib import Path
a = sys.argv
fmt = a[a.index("--convert-to") + 1]
out = Path(a[a.index("--outdir") + 1])
src = Path(a[-1])
shutil.copyfile(Path({str(made)!r}) / ("out." + fmt), out / (src.stem + "." + fmt))
""", encoding="utf-8")
    exe.chmod(exe.stat().st_mode | stat.S_IEXEC)
    monkeypatch.setenv("ZZAIMY_SOFFICE", str(exe))
    monkeypatch.setenv("ZZAIMY_NO_OCR_FALLBACK", "1")
    monkeypatch.setenv("ZZAIMY_KORDOC_OFF", "1")
    return exe


@pytest.mark.parametrize("name,want", [("안내.rtf", "서식 있는 글 본문"), ("집계.ods", "열린 표"), ("서식.xlt", "열린 표")])
def test_libreoffice_formats(tmp_path, fake_soffice, name, want):
    p = tmp_path / name
    p.write_bytes(b"{\\rtf1 x}" if name.endswith(".rtf") else b"\xd0\xcf\x11\xe0" + b"\x00" * 60)
    proc = pipeline.DocumentProcessor()
    assert want in proc._parse_inner(p)
    assert "LibreOffice" in proc._last_parse_note and "옛" not in proc._last_parse_note


def test_download_and_extensionless_files(tmp_path, monkeypatch):
    monkeypatch.setenv("ZZAIMY_NO_OCR_FALLBACK", "1")
    monkeypatch.setenv("ZZAIMY_KORDOC_OFF", "1")
    d = _docx(tmp_path / "x.docx", "선정 결과 공고").rename(tmp_path / "결과 공고다운로드")
    proc = pipeline.DocumentProcessor()
    assert "선정 결과 공고" in proc._parse_inner(d)
    assert "DOCX" in proc._last_parse_note
    n = _docx(tmp_path / "y.docx", "확장자 없는 계획").rename(tmp_path / "계획")
    assert "확장자 없는 계획" in pipeline.DocumentProcessor()._parse_inner(n)
    js = tmp_path / "analytics.js.다운로드"
    js.write_text("window.x = 1;", encoding="utf-8")
    with pytest.raises(RuntimeError, match="문서 형식이 아님"):
        pipeline.DocumentProcessor()._parse_inner(js)
