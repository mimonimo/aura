"""마크다운 → docx(독스용) — 제목 스타일, 개조식 부호·내어쓰기, 파이프·HTML 표(병합), 머리 행·라벨 열 음영, 그림, 코드."""

from __future__ import annotations

import base64
import io
import struct
import zlib

from docx import Document
from docx.oxml.ns import qn

from zzaimy.ingest import md_docx


def _png(w=4, h=2) -> bytes:
    def ch(t, d):
        return struct.pack(">I", len(d)) + t + d + struct.pack(">I", zlib.crc32(t + d) & 0xFFFFFFFF)
    raw = b"".join(b"\x00" + b"\x10\x20\x30" * w for _ in range(h))
    return b"\x89PNG\r\n\x1a\n" + ch(b"IHDR", struct.pack(">IIBBBBB", w, h, 8, 2, 0, 0, 0)) + ch(b"IDAT", zlib.compress(raw)) + ch(b"IEND", b"")


MD = """# 2027 산학협력 계획

## 1. 추진 배경

지역 산업 **수요**가 늘고 있다. 자세한 내용은 [공고](https://example.org/a)를 본다.

- 첫째 과제
  - 세부 과제 가
    - 더 세부
- 둘째 과제는 한 줄을 넘길 만큼 긴 문장으로 써서 내어쓰기가 필요한지 볼 수 있게 한다

1. 준비
2. 실행
   1. 가 단계

- [x] 완료한 일
- [ ] 남은 일

| 구분 | 2025 | 2026 | 비고 |
|---|---:|---:|---|
| 재학생 | 1,000 | 1,100 | 증가 |
| 교원 | 50 | 55 | 유지 |

<table><tr><th>항목</th><th colspan="2">금액</th></tr><tr><td rowspan="2">장학금</td><td>국고</td><td>10</td></tr><tr><td>대응</td><td>5</td></tr></table>

> 인용한 문장

```
print("hi")
```

---

![도식](data:image/png;base64,IMG)
![없는 그림](missing.png)
""".replace("IMG", base64.b64encode(_png()).decode())


def _doc():
    data, stats = md_docx.convert(MD)
    return Document(io.BytesIO(data)), stats


def test_headings_use_word_heading_styles_and_inline_marks():
    d, stats = _doc()
    heads = [(p.style.name, p.text) for p in d.paragraphs if p.style.name.startswith("Heading")]
    assert heads == [("Heading 1", "2027 산학협력 계획"), ("Heading 2", "1. 추진 배경")]
    body = next(p for p in d.paragraphs if p.text.startswith("지역 산업"))
    assert any(r.bold and r.text == "수요" for r in body.runs)
    assert "공고" in body._p.xml and "w:hyperlink" in body._p.xml          # 링크는 하이퍼링크로
    assert stats["headings"] == 2


def test_lists_get_gongmun_markers_and_hanging_indent():
    d, stats = _doc()
    texts = [p.text for p in d.paragraphs]
    assert "□\t첫째 과제" in texts and "○\t세부 과제 가" in texts and "-\t더 세부" in texts
    assert "1.\t준비" in texts and "2.\t실행" in texts and "가.\t가 단계" in texts
    assert "□\t☑ 완료한 일" in texts and "□\t☐ 남은 일" in texts
    p = next(p for p in d.paragraphs if p.text.startswith("○\t"))
    pf = p.paragraph_format
    assert pf.first_line_indent < 0 and pf.left_indent > 0 and abs(pf.first_line_indent) < pf.left_indent   # 둘째 단계는 더 들어간다
    assert stats["list_items"] == 9


def test_pipe_table_head_shading_alignment_and_label_column():
    d, _ = _doc()
    t = d.tables[0]
    head = t.cell(0, 0)._tc.xml
    assert md_docx.HEAD_FILL in head and "w:b" in t.cell(0, 0)._tc.xml
    assert md_docx.LABEL_FILL in t.cell(1, 0)._tc.xml                     # 첫 열이 짧은 글·나머지가 수치 → 라벨 열
    assert md_docx.LABEL_FILL not in t.cell(1, 1)._tc.xml
    assert t.cell(1, 1).paragraphs[0].alignment == 2                       # ---: 오른쪽
    widths = [int(gc.get(qn("w:w"))) for gc in t._tbl.tblGrid.findall(qn("w:gridCol"))]
    assert len(widths) == 4 and max(widths) <= 0.55 * sum(widths) + 1


def test_html_table_merges_cells():
    d, stats = _doc()
    t = d.tables[1]
    assert len(t.rows) == 3 and len(t.columns) == 3
    assert t.cell(0, 1)._tc is t.cell(0, 2)._tc                            # colspan
    assert t.cell(1, 0)._tc is t.cell(2, 0)._tc and t.cell(1, 0).text == "장학금"   # rowspan, 빈 문단 없이
    assert stats["merged_cells"] == 2 and stats["tables"] == 2


def test_quote_code_rule_and_images():
    d, stats = _doc()
    q = next(p for p in d.paragraphs if p.text == "인용한 문장")
    assert "w:pBdr" in q._p.xml
    assert any('print("hi")' in c.text for t in d.tables for c in t._cells) and stats["code_blocks"] == 1
    assert stats["images"] == 1 and stats["images_missing"] == 1
    assert any("[그림: 없는 그림]" in p.text for p in d.paragraphs)


def test_relative_images_stay_inside_base_dir(tmp_path):
    (tmp_path / "fig.png").write_bytes(_png())
    (tmp_path.parent / "secret.png").write_bytes(_png())
    _data, stats = md_docx.convert("![a](fig.png)\n\n![b](../secret.png)\n", base_dir=tmp_path)
    assert stats["images"] == 1 and stats["images_missing"] == 1


def test_markdown_document_goes_to_docs_as_converted_docx(tmp_path):
    from zzaimy.ingest import gdrive_files

    src = tmp_path / "보고.md"
    src.write_text("# 제목\n\n- 항목\n", encoding="utf-8")
    data, name, mime, target = gdrive_files.bytes_for_view(None, {"stored_path": str(src), "filename": "보고.md", "id": 1})
    assert name == "보고.docx" and target == "application/vnd.google-apps.document" and mime.endswith("wordprocessingml.document")
    assert "□\t항목" in [p.text for p in Document(io.BytesIO(data)).paragraphs]
