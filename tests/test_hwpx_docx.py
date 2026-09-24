"""HWPX → DOCX 충실 변환 — 글자 모양·문단 정렬·표(병합·테두리·바탕색·중첩)·쪽 설정이 옮겨진다."""

import io
import zipfile

from docx import Document
from docx.oxml.ns import qn

from zzaimy.ingest import hwpx_docx

HEADER = """<?xml version="1.0" encoding="UTF-8"?>
<hh:head xmlns:hh="http://www.hancom.co.kr/hwpml/2011/head" xmlns:hc="http://www.hancom.co.kr/hwpml/2011/core">
<hh:refList>
<hh:fontfaces><hh:fontface lang="HANGUL"><hh:font id="0" face="맑은 고딕" type="TTF"/></hh:fontface></hh:fontfaces>
<hh:borderFills itemCnt="3">
 <hh:borderFill id="1"><hh:leftBorder type="NONE" width="0.1 mm"/><hh:rightBorder type="NONE" width="0.1 mm"/><hh:topBorder type="NONE" width="0.1 mm"/><hh:bottomBorder type="NONE" width="0.1 mm"/></hh:borderFill>
 <hh:borderFill id="2"><hh:leftBorder type="SOLID" width="0.12 mm"/><hh:rightBorder type="SOLID" width="0.12 mm"/><hh:topBorder type="SOLID" width="0.12 mm"/><hh:bottomBorder type="SOLID" width="0.12 mm"/></hh:borderFill>
 <hh:borderFill id="3"><hh:leftBorder type="SOLID" width="0.12 mm"/><hh:rightBorder type="SOLID" width="0.12 mm"/><hh:topBorder type="SOLID" width="0.12 mm"/><hh:bottomBorder type="SOLID" width="0.12 mm"/><hc:fillBrush><hc:winBrush faceColor="#D6D6D6" hatchColor="#000000" alpha="0"/></hc:fillBrush></hh:borderFill>
</hh:borderFills>
<hh:charProperties itemCnt="2">
 <hh:charPr id="0" height="1000" textColor="#000000"><hh:fontRef hangul="0"/></hh:charPr>
 <hh:charPr id="1" height="1600" textColor="#2525F5"><hh:fontRef hangul="0"/><hh:bold/></hh:charPr>
</hh:charProperties>
<hh:paraProperties itemCnt="2">
 <hh:paraPr id="0"><hh:align horizontal="JUSTIFY"/><hh:heading type="NONE" level="0"/></hh:paraPr>
 <hh:paraPr id="1"><hh:align horizontal="CENTER"/><hh:heading type="OUTLINE" level="0"/></hh:paraPr>
</hh:paraProperties>
<hh:styles itemCnt="1"><hh:style id="0" type="PARA" name="바탕글" paraPrIDRef="0" charPrIDRef="0"/></hh:styles>
</hh:refList></hh:head>"""

SECTION = """<?xml version="1.0" encoding="UTF-8"?>
<hs:sec xmlns:hs="http://www.hancom.co.kr/hwpml/2011/section" xmlns:hp="http://www.hancom.co.kr/hwpml/2011/paragraph">
<hp:p paraPrIDRef="1" styleIDRef="0"><hp:run charPrIDRef="0"><hp:secPr><hp:pagePr landscape="WIDELY" width="59528" height="84188"><hp:margin left="5669" right="5669" top="4251" bottom="2834" header="2834" footer="2834" gutter="0"/></hp:pagePr></hp:secPr></hp:run>
 <hp:run charPrIDRef="1"><hp:t>사업 개요</hp:t></hp:run></hp:p>
<hp:p paraPrIDRef="0" styleIDRef="0"><hp:run charPrIDRef="0"><hp:t>첫 줄<hp:lineBreak/>둘째 줄</hp:t></hp:run></hp:p>
<hp:p paraPrIDRef="0" styleIDRef="0"><hp:run charPrIDRef="0"><hp:tbl rowCnt="2" colCnt="2" borderFillIDRef="2"><hp:sz width="40000"/>
 <hp:tr>
  <hp:tc borderFillIDRef="3"><hp:subList vertAlign="CENTER"><hp:p paraPrIDRef="0"><hp:run charPrIDRef="1"><hp:t>제목 칸</hp:t></hp:run></hp:p></hp:subList><hp:cellAddr rowAddr="0" colAddr="0"/><hp:cellSpan rowSpan="1" colSpan="2"/><hp:cellSz width="40000" height="1000"/></hp:tc>
 </hp:tr>
 <hp:tr>
  <hp:tc borderFillIDRef="2"><hp:subList vertAlign="TOP"><hp:p paraPrIDRef="0"><hp:run charPrIDRef="0"><hp:tbl rowCnt="1" colCnt="1" borderFillIDRef="2"><hp:sz width="10000"/><hp:tr><hp:tc borderFillIDRef="2"><hp:subList><hp:p paraPrIDRef="0"><hp:run charPrIDRef="0"><hp:t>안쪽 표</hp:t></hp:run></hp:p></hp:subList><hp:cellAddr rowAddr="0" colAddr="0"/><hp:cellSpan rowSpan="1" colSpan="1"/><hp:cellSz width="10000" height="500"/></hp:tc></hp:tr></hp:tbl></hp:run></hp:p></hp:subList><hp:cellAddr rowAddr="1" colAddr="0"/><hp:cellSpan rowSpan="1" colSpan="1"/><hp:cellSz width="10000" height="90000"/></hp:tc>
  <hp:tc borderFillIDRef="2"><hp:subList vertAlign="TOP"><hp:p paraPrIDRef="0"><hp:run charPrIDRef="0"><hp:t>오른쪽</hp:t></hp:run></hp:p></hp:subList><hp:cellAddr rowAddr="1" colAddr="1"/><hp:cellSpan rowSpan="1" colSpan="1"/><hp:cellSz width="30000" height="1000"/></hp:tc>
 </hp:tr>
</hp:tbl></hp:run></hp:p>
</hs:sec>"""


def _hwpx(tmp_path):
    p = tmp_path / "t.hwpx"
    with zipfile.ZipFile(p, "w") as zf:
        zf.writestr("mimetype", "application/hwp+zip")
        zf.writestr("Contents/header.xml", HEADER)
        zf.writestr("Contents/section0.xml", SECTION)
    return p


def test_convert_keeps_styles_tables_and_page(tmp_path):
    data, stats = hwpx_docx.convert(_hwpx(tmp_path))
    assert stats["tables"] == 1 and stats["nested_tables"] == 1
    d = Document(io.BytesIO(data))
    # 개요 문단은 제목 스타일, 굵고 파란 16pt 글자
    h = d.paragraphs[0]
    assert h.style.name.startswith("Heading") and h.runs[0].bold and h.runs[0].font.size.pt == 16
    assert str(h.runs[0].font.color.rgb) == "2525F5"
    # 줄 바꿈은 같은 문단 안의 br
    assert len(d.paragraphs[1]._p.findall(".//" + qn("w:br"))) == 1
    # 표: 첫 행 병합 + 회색 바탕, 둘째 행 왼쪽 칸에 중첩 표, 오른쪽 칸 글
    t = d.tables[0]
    assert len(t.rows) == 2 and len(t.columns) == 2
    head = t.cell(0, 0)
    assert head.text == "제목 칸" and head._tc is t.cell(0, 1)._tc
    shd = head._tc.tcPr.find(qn("w:shd"))
    assert shd is not None and shd.get(qn("w:fill")) == "D6D6D6"
    assert t.cell(1, 0).tables and t.cell(1, 0).tables[0].cell(0, 0).text == "안쪽 표"
    assert t.cell(1, 1).text == "오른쪽"
    # 배치용 큰 행 높이는 상한(5인치)으로 — 빈 쪽을 만들지 않게. 중첩 표가 든 행은 높이를 정하지 않는다
    trh = t.rows[1]._tr.find(qn("w:trPr"))
    assert trh is None or trh.find(qn("w:trHeight")) is None
    # 쪽: A4 세로, 여백
    sec = d.sections[0]
    assert round(sec.page_width.inches, 2) == 8.27 and round(sec.page_height.inches, 2) == 11.69
    assert round(sec.left_margin.inches, 2) == 0.79


def test_encrypted_section_raises(tmp_path):
    p = tmp_path / "e.hwpx"
    with zipfile.ZipFile(p, "w") as zf:
        zf.writestr("Contents/header.xml", HEADER)
        zf.writestr("Contents/section0.xml", b"\x00\x01garbage")
    import pytest

    with pytest.raises(ValueError):
        hwpx_docx.convert(p)
