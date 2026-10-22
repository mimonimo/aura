"""사업 관리 — 가져오기 단계·되돌리기, 내보내기 갈래·거르기, 권한별 고치기·신청·승인·반려·알림, 변경 기록·되돌리기,
합치기·사업 아님, 사람 찾기, 연도별 값 표 저장, 금액·기간 꼴, 파일 상한·수식 막기. 엑셀·사람은 모두 가짜 값."""

from __future__ import annotations

import io
import json
import zipfile

import openpyxl
import pytest

from tests.test_staff_grades import _add, _client, _no_google  # noqa: F401  (구글 끄기 고정물)
from tests.test_program_sheet import LEDGER, R25, R26, book_full, world  # noqa: F401  (고정물)
from zzaimy.app import program_admin, program_sheet

KEY = "program:가나다혁신사업"          # 기획처 사업
OTHER = "program:마바사지원사업"         # 교무처 사업(엑셀에서 더한 줄)


def _preview(c, *paths, **form):
    files = [("files", (p.name, p.read_bytes(), "application/octet-stream")) for p in paths]
    return c.post("/programs/admin/sheet/preview", data=form, files=files)


def _token(r):
    return r.text.split('name="token" value="')[1].split('"')[0]


@pytest.fixture()
def filled(world):
    app, tmp_path, *_ = world
    x25, x26 = tmp_path / "f25.xlsx", tmp_path / "f26.xlsx"
    book_full(x25, 2025, R25)
    book_full(x26, 2026, R26, fiscal=True)
    dev = _client(app, "zzdev", "devpass")
    r = dev.post("/programs/admin/sheet/apply", data={"token": _token(_preview(dev, x25, x26))}, follow_redirects=False)
    assert "ok=" in r.headers["location"]
    _add(app, "clerk1", dept="기획처", grade="clerk")
    _add(app, "boss2", dept="교무처", grade="head")
    return app, tmp_path, dev, x25, x26


# ---------------------------------------------------------------- 가져오기

def test_import_steps_by_year_not_sheet_and_undo(world):
    app, tmp_path, *_ = world
    dev = _client(app, "zzdev", "devpass")
    x25 = tmp_path / "f25.xlsx"
    book_full(x25, 2025, R25)
    other = tmp_path / "개설과목.xlsx"                                  # 사업명 머리줄이 없는 엑셀
    wb = openpyxl.Workbook()
    wb.active.append(["과목명", "학점"])
    wb.active.append(["가짜 과목", 3])
    wb.save(other)
    r = _preview(dev, x25, other)
    assert r.status_code == 200 and "data-pa-preview" in r.text and "data-pa-byyear" in r.text
    assert f"개설과목.xlsx: {program_sheet.NOT_SHEET}" in r.text            # 그 파일만 빼고 알린다
    assert "<td>2025</td>" in r.text and "손으로 고친 칸도 덮어쓰기" in r.text
    assert not (tmp_path / "program_admin.json").exists()
    # 3단계 — 적용하면 완료 요약과 되돌리기
    loc = dev.post("/programs/admin/sheet/apply", data={"token": _token(r)}, follow_redirects=False).headers["location"]
    assert "done=" in loc and "h=" in loc
    page = dev.get(loc).text
    assert "data-pa-done" in page and "data-pa-undo" in page
    stamp = loc.split("done=")[1].split("&")[0]
    h = loc.split("h=")[1].split("&")[0]
    assert program_admin.load_store(tmp_path)["rows"]
    r = dev.post("/programs/admin/sheet/undo", data={"stamp": stamp, "h": h}, follow_redirects=False)
    assert "ok=" in r.headers["location"]
    assert program_admin.load_store(tmp_path)["rows"] == {}                # 가져오기 앞 판(첫 저장 앞은 빈 판)으로


def test_undo_refused_after_other_save(filled):
    app, tmp_path, dev, x25, _ = filled
    loc = dev.post("/programs/admin/sheet/apply", data={"token": _token(_preview(dev, x25))}, follow_redirects=False).headers["location"]
    stamp, h = loc.split("done=")[1].split("&")[0], loc.split("h=")[1].split("&")[0]
    dev.post("/programs/admin/hide", data={"key": OTHER, "hidden": "1"})
    r = dev.post("/programs/admin/sheet/undo", data={"stamp": stamp, "h": h}, follow_redirects=False)
    assert "err=" in r.headers["location"]
    assert _client(app, "boss").post("/programs/admin/sheet/undo", data={"stamp": stamp, "h": h}).status_code == 403


def test_doc_list_only_program_sheets_with_cache(world):
    app, tmp_path, *_ = world
    dev = _client(app, "zzdev", "devpass")
    x25 = tmp_path / "f25.xlsx"
    book_full(x25, 2025, R25)
    other = tmp_path / "과목.xlsx"
    wb = openpyxl.Workbook()
    wb.active.append(["과목명"])
    wb.save(other)
    db = app.state.db
    good = db.add_document("현황 2025", str(x25), doc_type="grant")
    bad = db.add_document("개설과목 관리", str(other), doc_type="grant")
    g, o = program_admin.sheet_docs(db, tmp_path)
    assert [d["id"] for d in g] == [good] and g[0]["year"] == 2025 and [d["id"] for d in o] == [bad]
    cache = json.loads((tmp_path / "cache" / program_admin.PROBE_CACHE).read_text())
    assert cache[str(good)]["ok"] and not cache[str(bad)]["ok"]
    page = dev.get("/programs/admin").text
    good_box = page.split("data-pa-gooddocs")[1].split("</div>")[0]
    assert "현황 2025" in good_box and "개설과목 관리" not in good_box and "data-pa-otherdocs" in page
    book_full(other, 2024, R25[:1])                                       # 파일이 바뀌면(지문) 다시 판정
    g, o = program_admin.sheet_docs(db, tmp_path)
    assert {d["id"] for d in g} == {good, bad}


# ---------------------------------------------------------------- 내보내기

def _sheet_rows(content):
    wb = openpyxl.load_workbook(io.BytesIO(content))
    return wb, {ws.title: [ws.cell(i, 5 if ws.cell(4, 1).value == "기존,신규,종료" and ws.cell(4, 3).value == "회계연도" else 4).value
                           for i in range(5, ws.max_row)] for ws in wb.worksheets}


def test_export_kinds_follow_filters(filled):
    app, tmp_path, dev, *_ = filled
    r = dev.get("/programs/admin.xlsx", params={"all": "1", "tab": "all"})
    wb, subs = _sheet_rows(r.content)
    assert wb.sheetnames == ["2025", "2026"] and "1단계 과제" in subs["2025"] and len(subs["2026"]) == 2
    r = dev.get("/programs/admin.xlsx", params={"all": "1", "tab": "all", "dept": "교무처"})
    _, subs = _sheet_rows(r.content)
    assert subs["2025"] == ["마바사 지원사업"]                            # 거르기(부서)를 그대로
    r = dev.get("/programs/admin.xlsx", params={"year": "2025", "tab": "all", "st": "종료"})
    _, subs = _sheet_rows(r.content)
    assert subs["2025"] == ["아자차 과제"]                                # 구분 거르기
    csv_text = dev.get("/programs/admin.csv", params={"tab": "all", "src": "noexcel"}).content.decode("utf-8-sig")
    assert "가나다 혁신사업" not in csv_text and "라마바 지원사업" in csv_text   # 출처: 엑셀에 없는 사업만
    q = dev.get("/programs/admin.csv", params={"tab": "all", "q": "마바사"}).content.decode("utf-8-sig")
    assert "마바사 지원사업" in q and "가나다" not in q


def test_export_formula_injection_blocked(filled):
    app, tmp_path, dev, *_ = filled
    dev.post("/programs/admin/save", data={"key": KEY, "name": "가나다 혁신사업", "manager": "=HYPERLINK(\"x\")", "dept": "기획처"})
    csv_text = dev.get("/programs/admin.csv", params={"tab": "all"}).content.decode("utf-8-sig")
    assert "'=HYPERLINK" in csv_text and ",=HYPERLINK" not in csv_text
    lines = [{"name": "=1+1", "sub": "@SUM(A1)", "manager": "-2", "budget_won": 5}]
    ws = openpyxl.load_workbook(io.BytesIO(program_sheet.write_book(2025, lines))).active
    cells = [ws.cell(5, j) for j in range(1, ws.max_column + 1)]
    texts = [c for c in cells if isinstance(c.value, str)]
    assert {c.value for c in texts} >= {"=1+1", "@SUM(A1)", "-2"} and all(c.data_type == "s" for c in texts)
    assert program_sheet.csv_safe("+cmd") == "'+cmd" and program_sheet.csv_safe(12) == 12


# ---------------------------------------------------------------- 파일 상한

def test_wide_formatted_sheet_reads_only_used_columns(tmp_path):
    p = tmp_path / "넓은.xlsx"
    book_full(p, 2025, R25[:1])
    wb = openpyxl.load_workbook(p)
    wb.active.cell(1, 16000).number_format = "#,##0"                     # 서식만 있는 먼 열(원본처럼 max_column 이 큼)
    wb.active.cell(1, 16000).value = None
    wb.save(p)
    items, metas = program_sheet.read_sheets(p)
    assert len(items) == 1 and metas[0]["total"]["budget_won"] == sum(R25[0][-1])


def test_limits_refused(tmp_path, monkeypatch):
    big = tmp_path / "큰.xlsx"
    big.write_bytes(b"PK\x03\x04" + b"0" * 64)
    monkeypatch.setattr(program_sheet, "MAX_BYTES", 32)
    with pytest.raises(program_sheet.SheetError, match="MB"):
        program_sheet.read_sheets(big)
    monkeypatch.setattr(program_sheet, "MAX_BYTES", 10 * 1024 * 1024)
    bomb = tmp_path / "폭탄.xlsx"                                        # 풀면 큰 항목 하나(압축률 큼)
    with zipfile.ZipFile(bomb, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("[Content_Types].xml", "<x/>")
        z.writestr("xl/big.bin", b"\0" * (8 * 1024 * 1024))
    with pytest.raises(program_sheet.SheetError, match="압축률"):
        program_sheet.read_sheets(bomb)
    many = tmp_path / "시트많음.xlsx"
    wb = openpyxl.Workbook()
    for n in range(program_sheet.MAX_SHEETS + 1):
        wb.create_sheet(f"s{n}")
    wb.save(many)
    with pytest.raises(program_sheet.SheetError, match="시트가 너무 많"):
        program_sheet.read_sheets(many)


def test_upload_size_limit_message(world, monkeypatch):
    app, tmp_path, *_ = world
    monkeypatch.setattr(program_admin, "MAX_UPLOAD", 100)
    p = tmp_path / "f25.xlsx"
    book_full(p, 2025, R25)
    r = _preview(_client(app, "zzdev", "devpass"), p)
    assert "MB 보다 큽니다" in r.text or "err=" in str(r.url)


# ---------------------------------------------------------------- 권한·신청·기록

def test_head_edits_own_dept_only_and_requests(filled):
    app, tmp_path, dev, *_ = filled
    boss = _client(app, "boss")                                        # 기획처 총괄 책임자
    page = boss.get("/programs/admin?tab=all&year=2025").text
    assert "마바사 지원사업" in page                                      # 다른 부서 사업도 읽는다
    row_of = lambda k: [tr for tr in page.split("<tr") if f'data-key="{k}"' in tr][0]
    assert "data-pa-edit-open" in row_of(KEY) and "data-pa-req-open" not in row_of(KEY)
    assert "data-pa-edit-open" not in row_of(OTHER) and 'data-pa-req-open="edit"' in row_of(OTHER)
    r = boss.post("/programs/admin/save", data={"key": KEY, "name": "가나다 혁신사업", "dept": "기획처", "l3": "혁신", "leads": ["lead1"]},
                  follow_redirects=False)
    assert "ok=" in r.headers["location"]
    assert boss.post("/programs/admin/save", data={"key": OTHER, "name": "마바사 지원사업", "l3": "x"}).status_code == 403
    assert boss.post("/programs/admin/save", data={"key": "", "name": "새 사업"}).status_code == 403
    r = boss.post("/programs/admin/save", data={"key": KEY, "name": "가나다 혁신사업", "dept": "기획처", "heads": ["boss2"]}, follow_redirects=False)
    assert "boss2" not in program_admin.load_store(tmp_path)["rows"][KEY].get("heads", [])   # 총괄 책임자 지정은 관리자만
    # 수정 신청 → 관리자 알림 → 승인 → 반영·기록·신청자 알림
    r = boss.post("/programs/admin/request", data={"key": OTHER, "kind": "edit", "field": "y:manager", "year": "2025", "value": "새책임",
                                                    "reason": "담당 교체"}, follow_redirects=False)
    assert "ok=" in r.headers["location"]
    st = program_admin.load_store(tmp_path)
    q = st["requests"][0]
    assert q["status"] == "wait" and q["field"] == "manager" and q["year"] == "2025"
    assert any("새 신청 #1" in n["text"] for n in st["notes"]["zzdev"])
    assert "처리할 신청" in dev.get("/programs/admin").text
    reqs = dev.get("/programs/admin?view=requests").text
    assert "최새로 → 새책임" in reqs and "담당 교체" in reqs
    assert boss.post("/programs/admin/request/decide", data={"id": "1", "decision": "ok"}).status_code == 403
    dev.post("/programs/admin/request/decide", data={"id": "1", "decision": "ok"})
    st = program_admin.load_store(tmp_path)
    y = st["rows"][OTHER]["years"]["2025"]
    assert y["manager"] == "새책임" and "manager" in y["by_hand"]
    assert st["rows"][OTHER]["log"][-1]["via"] == "신청 #1 승인"
    assert any("신청 #1" in n["text"] and "승인" in n["text"] for n in st["notes"]["boss"])
    assert "신청 #1" in boss.get("/programs/admin").text                  # 신청자 알림 띠
    boss.post("/programs/admin/notes/read")
    assert not [n for n in program_admin.load_store(tmp_path)["notes"]["boss"] if not n["seen"]]
    # 지정 신청(나를 총괄 책임자로) → 반려는 사유가 있어야
    boss.post("/programs/admin/request", data={"key": OTHER, "kind": "assign", "role": "heads", "reason": "공동 운영"})
    assert "err=" in dev.post("/programs/admin/request/decide", data={"id": "2", "decision": "no"}, follow_redirects=False).headers["location"]
    dev.post("/programs/admin/request/decide", data={"id": "2", "decision": "no", "note": "부서 협의 먼저"})
    st = program_admin.load_store(tmp_path)
    assert st["requests"][1]["status"] == "no" and "boss" not in (st["rows"][OTHER].get("heads") or [])
    assert any("반려" in n["text"] and "부서 협의 먼저" in n["text"] for n in st["notes"]["boss"])
    # 다른 부서 사람을 사업 책임자로 신청하면 거절
    r = boss.post("/programs/admin/request", data={"key": OTHER, "kind": "assign", "role": "leads", "users": ["boss2"], "reason": "x"},
                  follow_redirects=False)
    assert "err=" in r.headers["location"]
    assert _client(app, "lead1").post("/programs/admin/request", data={"key": OTHER, "kind": "edit"}).status_code == 403


def test_change_log_and_revert(filled):
    app, tmp_path, dev, *_ = filled
    dev.post("/programs/admin/save", data={"key": KEY, "name": "가나다 혁신사업", "dept": "기획처", "l3": "첫 값"})
    dev.post("/programs/admin/save", data={"key": KEY, "name": "가나다 혁신사업", "dept": "기획처", "l3": "둘째 값"})
    d = dev.get("/programs/admin/row", params={"key": KEY}).json()
    e = next(x for x in d["log"] if x["field"] == "l3" and x["after"] == "둘째 값")
    assert e["via"] == "직접" and e["by"] == "zzdev" and e["can_revert"]
    first = next(x for x in d["log"] if x["field"] == "l3" and x["after"] == "첫 값")
    r = dev.post("/programs/admin/revert", data={"key": KEY, "i": str(first["i"])}, follow_redirects=False)
    assert "err=" in r.headers["location"]                                # 그 뒤 다시 바뀐 칸
    r = dev.post("/programs/admin/revert", data={"key": KEY, "i": str(e["i"])}, follow_redirects=False)
    assert "ok=" in r.headers["location"]
    rows = {r["key"]: r for r in program_admin.merged_rows(tmp_path)}
    assert rows[KEY]["l3"] == "첫 값" and rows[KEY]["log"][-1]["via"].startswith("되돌리기")
    assert any(x["via"] == "가져오기" for x in d["log"])                  # 가져오기도 기록
    assert _client(app, "boss").post("/programs/admin/revert", data={"key": KEY, "i": "0"}).status_code == 403
    assert not _client(app, "boss").get("/programs/admin/row", params={"key": KEY}).json()["log"][0]["can_revert"]


def test_merge_and_not_program(filled):
    app, tmp_path, dev, x25, _ = filled
    pid = app.state.db.create_project("grant", "마바사 프로젝트", owner="zzdev", program=OTHER + "|2025")
    r = dev.post("/programs/admin/merge", data={"key": OTHER, "into": KEY}, follow_redirects=False)
    assert "ok=" in r.headers["location"]
    rows = {r["key"]: r for r in program_admin.merged_rows(tmp_path)}
    assert rows[OTHER]["hidden"] and rows[OTHER]["merged_into"] == KEY and "마바사 지원사업" in rows[KEY]["terms"]
    y25 = rows[KEY]["years"]["2025"]
    assert len(y25["tasks"]) == 3 and y25["company_won"] == 30_000_000      # 겹치는 해는 과제를 잇고 금액을 더함
    assert app.state.db.get_project(pid)["program"] == KEY + "|2025"
    assert rows[KEY]["log"][-1]["field"] == "합치기"
    # 다시 가져와도 합친 이름은 받은 줄로 간다
    sm, _, _ = program_admin._plan(type("R", (), {"app": app})(), [(x25, "f25.xlsx")], False, "zzdev")
    assert not sm["new"] and not sm["ambiguous"]
    # 사업 아님으로 숨기면 다음 가져오기에서 건너뛴다
    dev.post("/programs/admin/hide", data={"key": "program:아자차연계사업", "hidden": "1", "not_program": "1"})
    sm, _, _ = program_admin._plan(type("R", (), {"app": app})(), [(x25, "f25.xlsx")], False, "zzdev")
    assert "아자차 연계사업" in sm["not_program"]
    page = dev.get("/programs/admin?tab=all&year=0").text
    assert "아자차 연계사업" not in page.split("<tbody>")[1]
    assert _client(app, "boss").post("/programs/admin/merge", data={"key": OTHER, "into": KEY}).status_code == 403


# ---------------------------------------------------------------- 사람 찾기

def test_people_api_scope_privacy_filters(filled):
    app, tmp_path, dev, *_ = filled
    app.state.accounts["clerk1"]["email"] = "fake@example.com"
    app.state.accounts["clerk1"]["phone"] = "010-0000-0000"
    r = dev.get("/api/people", params={"q": "clerk"}).json()
    assert [p["user"] for p in r["people"]] == ["clerk1"]
    assert set(r["people"][0]) == {"user", "name", "dept", "grade", "grade_label"}       # 개인정보 없음
    got = {p["user"] for p in dev.get("/api/people", params={"grade": "lead"}).json()["people"]}
    assert "clerk1" not in got and {"lead1", "boss"} <= got
    got = {p["user"] for p in dev.get("/api/people", params={"dept": "교무처"}).json()["people"]}
    assert got == {"boss2"}
    assert _client(app, "lead1").get("/api/people").status_code == 403
    d = dev.get("/api/people", params={"q": "기획", "kind": "dept"}).json()["depts"]
    assert d and d[0]["user"] == "dept:기획처"
    # 저장 검사 — 등급·부서가 맞지 않으면 이유 한 줄
    boss = _client(app, "boss")
    r = boss.post("/programs/admin/save", data={"key": KEY, "name": "가나다 혁신사업", "dept": "기획처", "leads": ["boss2"]},
                  follow_redirects=False)
    assert "err=" in r.headers["location"] and "%EC%82%AC%EC%97%85" in r.headers["location"]   # 「사업 책임자로 지정할 수 없습니다」
    dev.post("/programs/admin/save", data={"key": KEY, "name": "가나다 혁신사업", "dept": "기획처", "heads": ["dept:산학협력단"]})
    assert program_admin.load_store(tmp_path)["rows"][KEY]["heads"] == ["dept:산학협력단"]
    assert "산학협력단(부서)" in dev.get("/programs/admin?tab=all&year=0").text


def test_covered_depts_head_by_dept_grade(filled):
    """겸직 — 소속 등급은 사업 책임자여도 부서별 등급으로 어떤 부서의 총괄 책임자이면 그 부서 사업을 고친다."""
    app, tmp_path, dev, *_ = filled
    _add(app, "dual", dept="기획처", grade="lead")
    app.state.accounts["dual"]["dept_grades"] = {"교무처": "head"}
    c = _client(app, "dual")
    assert c.get("/programs/admin").status_code == 200
    r = c.post("/programs/admin/save", data={"key": OTHER, "name": "마바사 지원사업", "l3": "겸직 고침"}, follow_redirects=False)
    assert r.status_code == 303 and "ok=" in r.headers["location"]
    assert c.post("/programs/admin/save", data={"key": KEY, "name": "가나다 혁신사업", "l3": "x"}).status_code == 403


# ---------------------------------------------------------------- 연도별 값 표

def test_year_table_save_parse_sum_validate(filled):
    app, tmp_path, dev, *_ = filled
    table = {"2026": {"status": "기존", "year_period": "20260301-20270228", "manager": "홍길순", "officer": "강행정", "dept": "기획처",
                      "money": {"national": "1.2억", "local": "3천만", "company": "1,000,000", "school": "", "other": "", "incentive": ""}},
             "2027": {"status": "신규", "year_period": "2027.3.1~2028.2.28", "manager": "", "officer": "", "dept": "",
                      "tasks": [{"sub": "가 과제", "money": {"national": "1억"}}, {"sub": "나 과제", "money": {"national": 5000, "local": "2만"}}]}}
    r = dev.post("/programs/admin/save", data={"key": KEY, "name": "가나다 혁신사업", "dept": "기획처", "years_json": json.dumps(table)},
                 follow_redirects=False)
    assert "ok=" in r.headers["location"], r.headers["location"]
    ys = program_admin.load_store(tmp_path)["rows"][KEY]["years"]
    a = ys["2026"]
    assert a["national_won"] == 120_000_000 and a["local_won"] == 30_000_000 and a["company_won"] == 1_000_000
    assert a["budget_won"] == 151_000_000 and a["year_period"] == "2026.03.01~2027.02.28" and "national" in a["by_hand"]
    b = ys["2027"]
    assert b["national_won"] == 100_005_000 and b["local_won"] == 20_000 and b["budget_won"] == 100_025_000 and len(b["tasks"]) == 2
    for bad, msg in (({"2026": {"year_period": "2026.03.01~2025.01.01"}}, "%EB%81%9D"),              # 끝이 시작보다 앞
                     ({"2026": {"money": {"national": "많이"}}}, "%EA%B8%88%EC%95%A1"),                # 금액을 읽지 못함
                     ({"abcd": {"manager": "x"}}, "%EC%97%B0%EB%8F%84")):                              # 연도
        r = dev.post("/programs/admin/save", data={"key": KEY, "name": "가나다 혁신사업", "dept": "기획처", "years_json": json.dumps(bad)},
                     follow_redirects=False)
        assert "err=" in r.headers["location"] and msg in r.headers["location"]
    r = dev.post("/programs/admin/save", data={"key": KEY, "name": "가나다 혁신사업", "dept": "기획처",
                                                "years_json": json.dumps({"2027": {"remove": True}})}, follow_redirects=False)
    assert "2027" not in program_admin.load_store(tmp_path)["rows"][KEY]["years"]


def test_money_and_period_formats():
    f = program_sheet.fmt_unit
    assert [f(n) for n in (99_999_999, 100_000_000, 0, -150_000_000, None, 9_999, 81_900_000, 5_310_000_000, 5_300_000_000)] == \
        ["1억", "1억", "0", "-1.5억", "", "9,999", "8,190만", "53.1억", "53억"]
    assert f(5_310_000_000, "eok") == "53.1" and f(1_234_567_000, "won") == "1,234,567,000" and f(1_260_000, "million") == "1.3"
    p = program_sheet.parse_won
    assert (p("1.2억"), p("3천만"), p("12억 3,000만 원"), p("1,200"), p("abc")) == (120_000_000, 30_000_000, 1_230_000_000, 1200, None)
    n = program_sheet.norm_period
    assert n("2025.3.1~2026.2.28") == ("2025.03.01~2026.02.28", "") and n("20250301-20260228")[0] == "2025.03.01~2026.02.28"
    assert n("2025-03-01 ~ 2026-02-28")[0] == "2025.03.01~2026.02.28" and n("2025~")[0] == "2025~" and n("2025~2027")[0] == "2025~2027"
    assert n("2026.03.01~2025.01.01")[1] and n("2025.02.30")[1] and n("2027~2025")[1]
    assert program_sheet.short_period("2025.03.01~2030.02.28") == "2025.03~2030.02"


# ---------------------------------------------------------------- 보안 점검(권한 범위·상한 우회)

def test_link_and_new_project_need_write_scope(filled):
    """총괄 책임자는 다른 부서 사업을 읽기만 — 프로젝트 잇기·새 프로젝트도 403."""
    app, tmp_path, dev, *_ = filled
    boss = _client(app, "boss")
    pid = app.state.db.create_project("grant", "내 프로젝트", owner="boss")
    assert boss.post("/programs/admin/link", data={"key": OTHER, "project_id": str(pid), "year": "2026"}).status_code == 403
    assert boss.post("/programs/admin/new-project", data={"key": OTHER, "year": "2026"}).status_code == 403
    assert app.state.db.get_project(pid)["program"] == ""
    r = boss.post("/programs/admin/link", data={"key": KEY, "project_id": str(pid), "year": "2026"}, follow_redirects=False)
    assert r.status_code == 303 and app.state.db.get_project(pid)["program"] == KEY + "|2026"
    row = [tr for tr in boss.get("/programs/admin?tab=all&year=2026").text.split("<tr") if f'data-key="{OTHER}"' in tr][0]
    assert "data-pa-link-open" not in row


def test_head_import_cannot_touch_meta_or_other_rows(filled):
    """총괄 책임자 가져오기 — 연도 정보(sheets)·다른 부서 줄·새 줄을 쓰지 못한다."""
    app, tmp_path, dev, x25, _ = filled
    before = program_admin.load_store(tmp_path)
    fake = tmp_path / "f25b.xlsx"
    rows = [list(r) for r in R25]
    rows[2][7] = "바뀐사람"                                              # 교무처 사업의 책임자
    rows.append(("신규", "완전히 새 사업", "새 과제", "교육부", "x", "", "", "누군가", "", "", "기획처", "", [1, 0, 0, 0, 0, 0]))
    book_full(fake, 2025, [tuple(r) for r in rows])
    boss = _client(app, "boss")
    boss.post("/programs/admin/sheet/apply", data={"token": _token(_preview(boss, fake))})
    after = program_admin.load_store(tmp_path)
    assert after["sheets"] == before["sheets"]
    assert after["rows"][OTHER]["years"]["2025"]["manager"] == before["rows"][OTHER]["years"]["2025"]["manager"]
    assert set(after["added"]) == set(before["added"])


def test_import_doc_ids_respect_visibility_and_token_owner(filled):
    app, tmp_path, dev, x25, _ = filled
    db = app.state.db
    hidden = db.add_document("남의 부서 현황", str(x25), doc_type="grant")
    with db._conn() as conn:                                            # 산학협력단 부서 제한 문서로
        conn.execute("UPDATE documents SET dept = ?, owner = ?, access_level = ? WHERE id = ?", ("산학협력단", "zzdev", "dept", hidden))
    from zzaimy.app import access_policy as _ap
    assert not _ap.visible(db.get_document(hidden), dept="기획처", user="boss", role="staff")
    boss = _client(app, "boss")
    r = boss.post("/programs/admin/sheet/preview", data={"doc_ids": str(hidden)}, follow_redirects=False)
    assert r.status_code == 303 and "err=" in r.headers["location"]
    assert "남의 부서 현황" not in boss.get("/programs/admin").text
    # 남이 올린 파일 열쇠로는 읽지 못한다
    token = _token(_preview(dev, x25))
    r = boss.post("/programs/admin/sheet/preview", data={"token": token}, follow_redirects=False)
    assert r.status_code == 303 and "err=" in r.headers["location"]


def test_approval_rechecks_at_decision_time(filled):
    app, tmp_path, dev, *_ = filled
    boss = _client(app, "boss")
    boss.post("/programs/admin/request", data={"key": OTHER, "kind": "assign", "role": "leads", "users": ["lead1"], "reason": "공동"})
    app.state.accounts["lead1"]["active"] = False                       # 신청 뒤 계정이 멈춤
    r = dev.post("/programs/admin/request/decide", data={"id": "1", "decision": "ok"}, follow_redirects=False)
    assert "err=" in r.headers["location"]
    st = program_admin.load_store(tmp_path)
    assert st["requests"][0]["status"] == "wait" and "lead1" not in (st["rows"][OTHER].get("leads") or [])
    app.state.accounts["lead1"]["active"] = True
    app.state.accounts["boss"]["grade"] = "member"                      # 신청자가 더는 총괄 책임자가 아님
    r = dev.post("/programs/admin/request/decide", data={"id": "1", "decision": "ok"}, follow_redirects=False)
    assert "err=" in r.headers["location"]
    assert dev.post("/programs/admin/request/decide", data={"id": "99", "decision": "ok"}).status_code == 404


def test_people_api_head_sees_covered_depts_only(filled):
    app, tmp_path, dev, *_ = filled
    got = {p["user"] for p in _client(app, "boss").get("/api/people").json()["people"]}
    assert "boss2" not in got and {"lead1", "clerk1"} <= got


def test_reader_caps_merges_spans_entities(tmp_path):
    html = "<table><tr><td>사업명</td></tr>" + "".join(f'<tr><td rowspan="5000" colspan="200">x{i}</td></tr>' for i in range(40)) + "</table>"
    p = tmp_path / "span.html"
    p.write_text(html, encoding="utf-8")
    with pytest.raises(program_sheet.SheetError):
        program_sheet.read_sheets(p)
    x = tmp_path / "ent.xml"
    x.write_text('<?xml version="1.0"?><!DOCTYPE a [<!ENTITY e "x">]><Workbook xmlns="urn:schemas-microsoft-com:office:spreadsheet"></Workbook>')
    with pytest.raises(program_sheet.SheetError, match="DOCTYPE"):
        program_sheet.read_sheets(x)
    m = tmp_path / "merge.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "사업명"
    ws.cell(1, 150, "끝 열")
    for n in range(20):
        ws.merge_cells(start_row=2 + n * 300, start_column=1, end_row=301 + n * 300, end_column=150)
    ws.cell(6100, 150, 1)
    wb.save(m)
    with pytest.raises(program_sheet.SheetError, match="병합"):
        program_sheet.read_sheets(m)


def test_upload_count_and_total_caps(world, monkeypatch):
    app, tmp_path, *_ = world
    dev = _client(app, "zzdev", "devpass")
    p = tmp_path / "f25.xlsx"
    book_full(p, 2025, R25)
    monkeypatch.setattr(program_admin, "MAX_FILES", 2)
    r = _preview(dev, p, p, p)
    assert "한 번에 2개까지" in r.text
    monkeypatch.setattr(program_admin, "MAX_TOTAL", p.stat().st_size + 10)
    r = _preview(dev, p, p)
    assert "합계가 너무 큽니다" in r.text
    tmp = tmp_path / "tmp" / "program_sheet"
    _preview(dev, p)
    _preview(dev, p)                                                    # 거듭 불러도 앞 미리보기 파일은 지운다
    assert len(list(tmp.glob("*.who"))) == 1
