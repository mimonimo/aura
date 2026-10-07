"""오피스 파일 직접 읽기 — python-docx·openpyxl·python-pptx 로 우리 ParseResult 를 낸다(ADR-0051).

오피스 읽기 대결(scripts/176, 2026-10-06, 형식마다 실제 원본 24~25건): docling 은 어느 형식에서도 1등이 아니었다 —
엑셀 표 개수 맞음 42~56%, 발표자료 표 68%. 순서: docx·xlsx 는 kordoc 먼저(이 모듈은 물러날 곳), pptx 는 이 모듈이 1순위.
"""
from __future__ import annotations

import time
from pathlib import Path

from zzaimy.ingest.parsers.base import ParsedEntry, ParsedPage, ParsedTable, ParseResult, TableCell


def _table(page_no: int, rows: list[list[str]]) -> ParsedTable | None:
    n_rows = len(rows)
    n_cols = max((len(r) for r in rows), default=0)
    if not n_rows or not n_cols:
        return None
    cells = tuple(TableCell(row=i, col=j, text=str(v or "").strip()) for i, r in enumerate(rows) for j, v in enumerate(r))
    return ParsedTable(page_no=page_no, n_rows=n_rows, n_cols=n_cols, cells=cells)


def _entry(page_no: int, kind: str, text: str = "", ref: int = -1) -> ParsedEntry:
    return ParsedEntry(page_no=page_no, kind=kind, text=text, ref=ref)


def parse_docx(path: Path) -> ParseResult:
    import docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    t0 = time.time()
    d = docx.Document(str(path))
    entries, tables, text = [], [], []
    for child in d.element.body.iterchildren():           # 본문 순서대로 문단·표
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            p = Paragraph(child, d)
            if p.text.strip():
                kind = "heading" if (p.style is not None and str(p.style.name).lower().startswith(("heading", "제목"))) else "text"
                entries.append(_entry(1, kind, p.text.strip()))
                text.append(p.text.strip())
        elif tag == "tbl":
            tb = Table(child, d)
            rows = []
            for r in tb.rows:
                seen, row = set(), []
                for c in r.cells:
                    row.append("" if id(c._tc) in seen else c.text)    # 병합 칸은 한 번만
                    seen.add(id(c._tc))
                rows.append(row)
            t = _table(1, rows)
            if t:
                entries.append(_entry(1, "table", ref=len(tables)))
                tables.append(t)
                text.append("\n".join(" | ".join(x for x in r if x) for r in rows))
    return ParseResult(parser="python-docx", elapsed_s=time.time() - t0, pages=[ParsedPage(page_no=1, text="\n".join(text))],
                       tables=tables, entries=entries)


def parse_xlsx(path: Path) -> ParseResult:
    import openpyxl
    t0 = time.time()
    wb = openpyxl.load_workbook(str(path), data_only=True, read_only=True)
    pages, tables, entries = [], [], []
    for i, ws in enumerate(wb.worksheets, 1):
        rows = [[("" if v is None else str(v)) for v in r] for r in ws.iter_rows(values_only=True)]
        rows = [r for r in rows if any(x.strip() for x in r)]
        if not rows:
            continue
        width = max(j + 1 for r in rows for j, x in enumerate(r) if x.strip())
        rows = [r[:width] for r in rows]
        t = _table(i, rows)
        if t:
            entries += [_entry(i, "heading", ws.title), _entry(i, "table", ref=len(tables))]
            tables.append(t)
        pages.append(ParsedPage(page_no=i, text=ws.title + "\n" + "\n".join(" | ".join(x for x in r if x.strip()) for r in rows)))
    return ParseResult(parser="openpyxl", elapsed_s=time.time() - t0, pages=pages, tables=tables, entries=entries)


def parse_pptx(path: Path) -> ParseResult:
    from pptx import Presentation
    t0 = time.time()
    prs = Presentation(str(path))
    pages, tables, entries = [], [], []
    for i, slide in enumerate(prs.slides, 1):
        texts = []

        def walk(shapes):
            for sh in shapes:
                if getattr(sh, "shape_type", None) == 6 and hasattr(sh, "shapes"):     # 그룹 도형
                    walk(sh.shapes)
                if getattr(sh, "has_text_frame", False) and sh.has_text_frame and sh.text_frame.text.strip():
                    kind = "heading" if getattr(sh, "is_placeholder", False) and "title" in str(
                        getattr(sh.placeholder_format, "type", "")).lower() else "text"
                    entries.append(_entry(i, kind, sh.text_frame.text.strip()))
                    texts.append(sh.text_frame.text.strip())
                if getattr(sh, "has_table", False) and sh.has_table:
                    rows = [[c.text for c in r.cells] for r in sh.table.rows]
                    t = _table(i, rows)
                    if t:
                        entries.append(_entry(i, "table", ref=len(tables)))
                        tables.append(t)
                        texts.append("\n".join(" | ".join(x for x in r if x) for r in rows))
        try:
            walk(slide.shapes)
        except Exception:
            # python-pptx 가 모르는 요소(호환 묶음 mc:AlternateContent 등)를 도형으로 바꾸다 실패하면 문서 전체가 실패했다
            # (2026-10-06 실패 13건 'has_ph_elm'). 그 슬라이드는 XML 의 글자(a:t)를 문단 단위로 직접 모은다
            texts.extend(_xml_paragraphs(slide._element))
            entries.extend(_entry(i, "text", t) for t in texts if t)
        try:
            if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip():
                texts.append(slide.notes_slide.notes_text_frame.text.strip())       # 발표자 메모
        except Exception:
            pass
        pages.append(ParsedPage(page_no=i, text="\n".join(texts)))
    return ParseResult(parser="python-pptx", elapsed_s=time.time() - t0, pages=pages, tables=tables, entries=entries)


_A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"


def _xml_paragraphs(elm) -> list[str]:
    """슬라이드 XML 에서 문단(a:p)마다 글자(a:t)를 이어 붙인다 — 도형 해석이 안 될 때의 물러날 곳."""
    out = []
    for p in elm.iter(_A + "p"):
        t = "".join(x.text or "" for x in p.iter(_A + "t")).strip()
        if t:
            out.append(t)
    return out


PARSERS = {".docx": parse_docx, ".xlsx": parse_xlsx, ".pptx": parse_pptx}
