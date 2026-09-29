"""마크다운 → docx(구글 독스로 올릴 것) — 공문서 모양으로(2026-09-30 사용자 지시 "md → 독스").

독스로 가져가는 길은 한글과 같다: 우리 docx → 드라이브 변환 업로드(gdrive_files.bytes_for_view). 드라이브에 마크다운을 그대로 올리면
평문이 되어 '#'·'|'·'**' 가 글자로 남는다. 모양은 kordoc 의 markdownToHwpx(보고서 프리셋)에서 원리만 가져왔다 — 흐름이 아니라 기술만:

- 개조식 부호 위계: 글머리 목록 단계마다 □ → ○ → - → ㆍ, 번호 목록은 1. → 가. → 1) → 가). 부호는 글자로 넣는다(우리 절 본문과
  hwpx_fill 이 부호를 글로 다룬다). 부호 뒤는 탭, 문단은 내어쓰기(첫 줄 −걸이, 탭 자리 = 왼쪽 여백)라 둘째 줄이 첫 줄 내용 시작에 맞는다.
- 표: 머리 행 음영(#DFE6F7)·굵게·가운데, 첫 열이 짧은 글이고 나머지가 수치·긴 글이면 라벨 열 음영(#F2F4F8). 열 폭은 칸 글 길이로
  나누되 한 열이 55% 를 넘지 않고, 짧은 열(수치·비고)은 좁게. GFM 파이프 표와 HTML 표(rowspan·colspan 병합) 둘 다 받는다 —
  kordoc --html-tables 출력이나 병합이 필요한 표를 그대로 받기 위해서다.
- 제목은 워드 제목 스타일(독스 개요에 잡힌다), 글꼴은 독스에 있는 나눔바른고딕, 줄 간격은 독스 규칙(ADR-0036: 비율 ÷ 글꼴 자연 행 높이).

정답을 코드에 넣지 않는다 — 어떤 문서든 같은 규칙으로 바뀐다.
"""

from __future__ import annotations

import base64
import io
import re
from html.parser import HTMLParser
from pathlib import Path

from zzaimy.ingest.hwpx_docx import (DOCS_LINE_EM_BY_FONT, FONT_MAP, BorderFill, _set_cell_borders, _set_cell_margins,
                                     _table_fixed_layout)

BULLETS = ("□", "○", "-", "ㆍ")
HEAD_FILL = "DFE6F7"
LABEL_FILL = "F2F4F8"
CODE_FILL = "F5F5F5"
BODY_PT = 11.0
LINE_PCT = 1.6                       # 한글 공문서 흔한 줄 간격 160%
HEADING_PT = {1: 16.0, 2: 14.0, 3: 12.5, 4: 11.5, 5: 11.0, 6: 11.0}
INDENT_PT = 14.0                     # 목록 한 단계 들여쓰기
PAGE_MM = (210, 297)
MARGIN_MM = (20, 20, 20, 15)         # 왼·오·위·아래
_HANGUL = ("가", "나", "다", "라", "마", "바", "사", "아", "자", "차", "카", "타", "파", "하")
_NUMERIC = re.compile(r"^[\s\-+±△▲▼()%.,:~/원천만억개명건회년월일시분초점배%]*\d[\d\s\-+±△▲▼()%.,:~/원천만억개명건회년월일시분초점배%]*$")


def _font() -> str:
    return FONT_MAP["고딕"]


def _line_mult(font: str) -> float:
    return round(LINE_PCT / DOCS_LINE_EM_BY_FONT.get(font, 1.2), 3)


def _number(level: int, n: int) -> str:
    """번호 목록의 단계별 번호 글자 — 1. → 가. → 1) → 가)."""
    kind = level % 4
    han = _HANGUL[(n - 1) % len(_HANGUL)]
    return (f"{n}.", f"{han}.", f"{n})", f"{han})")[kind]


def _set_run_font(run, font: str, size: float | None = None, bold: bool | None = None, color: str | None = None) -> None:
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor

    run.font.name = font
    rpr = run._element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        from docx.oxml import OxmlElement

        rf = OxmlElement("w:rFonts")
        rpr.insert(0, rf)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rf.set(qn(attr), font)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def _base_styles(doc, font: str) -> None:
    from docx.oxml.ns import qn
    from docx.shared import Mm, Pt, RGBColor

    sec = doc.sections[0]
    sec.page_width, sec.page_height = Mm(PAGE_MM[0]), Mm(PAGE_MM[1])
    sec.left_margin, sec.right_margin, sec.top_margin, sec.bottom_margin = (Mm(m) for m in MARGIN_MM)
    normal = doc.styles["Normal"]
    normal.font.name = font
    normal.font.size = Pt(BODY_PT)
    normal.element.get_or_add_rPr()
    rf = normal.element.rPr.find(qn("w:rFonts"))
    if rf is not None:
        for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
            rf.set(qn(attr), font)
    pf = normal.paragraph_format
    pf.line_spacing = _line_mult(font)
    pf.space_after = Pt(4)
    for lvl, size in HEADING_PT.items():
        st = doc.styles[f"Heading {lvl}"]
        st.font.name = font
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.italic = False
        st.font.color.rgb = RGBColor(0, 0, 0)
        rpr = st.element.get_or_add_rPr()
        rf = rpr.find(qn("w:rFonts"))
        if rf is not None:
            for attr in list(rf.attrib):
                del rf.attrib[attr]
            for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
                rf.set(qn(attr), font)
        st.paragraph_format.space_before = Pt(14 if lvl <= 2 else 10)
        st.paragraph_format.space_after = Pt(6)
        st.paragraph_format.line_spacing = 1.0


def _text_width_pt() -> float:
    return (PAGE_MM[0] - MARGIN_MM[0] - MARGIN_MM[1]) / 25.4 * 72


def _add_hyperlink(paragraph, url: str, text: str, font: str, bold=False, italic=False) -> None:
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    r_id = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    rf = OxmlElement("w:rFonts")
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rf.set(qn(attr), font)
    rpr.append(rf)
    if bold:
        rpr.append(OxmlElement("w:b"))
    if italic:
        rpr.append(OxmlElement("w:i"))
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "1155CC")
    rpr.append(color)
    u = OxmlElement("w:u")
    u.set(qn("w:val"), "single")
    rpr.append(u)
    run.append(rpr)
    t = OxmlElement("w:t")
    t.text = text
    t.set(qn("xml:space"), "preserve")
    run.append(t)
    link.append(run)
    paragraph._p.append(link)


class _Inline:
    """markdown-it 인라인 토큰을 docx run 으로 — 굵게·기울임·취소선·코드·링크·줄바꿈·그림."""

    def __init__(self, conv: "Converter") -> None:
        self.c = conv

    def emit(self, para, children, size: float | None = None, bold_all: bool = False, color: str | None = None) -> None:
        bold = italic = strike = False
        link: str | None = None
        for t in children or []:
            kind = t.type
            if kind == "strong_open":
                bold = True
            elif kind == "strong_close":
                bold = False
            elif kind == "em_open":
                italic = True
            elif kind == "em_close":
                italic = False
            elif kind == "s_open":
                strike = True
            elif kind == "s_close":
                strike = False
            elif kind == "link_open":
                link = t.attrs.get("href") or ""
            elif kind == "link_close":
                link = None
            elif kind in ("softbreak",):
                self._run(para, " ", size, bold or bold_all, italic, strike, color)
            elif kind == "hardbreak":
                para.add_run().add_break()
            elif kind == "code_inline":
                r = para.add_run(t.content)
                _set_run_font(r, "Courier New", (size or BODY_PT) - 1)
            elif kind == "image":
                self.c.image(para, t.attrs.get("src") or "", t.content or "")
            elif kind == "html_inline":
                if re.fullmatch(r"<br\s*/?>", t.content.strip(), re.I):
                    para.add_run().add_break()
            elif kind == "text":
                if link and re.match(r"https?://", link):
                    _add_hyperlink(para, link, t.content, self.c.font, bold or bold_all, italic)
                else:
                    self._run(para, t.content, size, bold or bold_all, italic, strike, color)

    def _run(self, para, text, size, bold, italic, strike, color) -> None:
        if not text:
            return
        r = para.add_run(text)
        _set_run_font(r, self.c.font, size, bold if bold else None, color)
        if italic:
            r.font.italic = True
        if strike:
            r.font.strike = True


class _HtmlTable(HTMLParser):
    """HTML 표 → 칸 격자. 칸 = {text, rowspan, colspan, head}. 중첩 표는 글로 펴서 바깥 칸에 넣는다."""

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.rows: list[list[dict]] = []
        self.cell: dict | None = None
        self.depth = 0

    def handle_starttag(self, tag, attrs):
        a = dict(attrs)
        if tag == "table":
            self.depth += 1
        elif self.depth > 1:
            if tag in ("tr", "br", "p") and self.cell is not None:
                self.cell["text"] += "\n"
        elif tag == "tr":
            self.rows.append([])
        elif tag in ("td", "th"):
            if not self.rows:
                self.rows.append([])
            self.cell = {"text": "", "rowspan": max(1, int(a.get("rowspan") or 1)), "colspan": max(1, int(a.get("colspan") or 1)),
                         "head": tag == "th"}
            self.rows[-1].append(self.cell)
        elif tag in ("br", "p", "li") and self.cell is not None and self.cell["text"]:
            self.cell["text"] += "\n"

    def handle_endtag(self, tag):
        if tag == "table":
            self.depth -= 1
        elif tag in ("td", "th") and self.depth == 1:
            if self.cell is not None:
                self.cell["text"] = re.sub(r"[ \t]+", " ", self.cell["text"]).strip()
            self.cell = None

    def handle_data(self, data):
        if self.cell is not None:
            self.cell["text"] += data


def _grid(rows: list[list[dict]]) -> tuple[list[list[dict | None]], list[tuple[int, int, int, int]]]:
    """rowspan·colspan 을 펴서 격자(칸 머리만 dict, 덮인 자리는 None)와 병합 목록(r, c, rs, cs)."""
    grid: list[list[dict | None]] = []
    taken: dict[tuple[int, int], bool] = {}
    merges = []
    for r, row in enumerate(rows):
        while len(grid) <= r:
            grid.append([])
        c = 0
        for cell in row:
            while taken.get((r, c)):
                c += 1
            rs, cs = cell["rowspan"], cell["colspan"]
            for dr in range(rs):
                for dc in range(cs):
                    taken[(r + dr, c + dc)] = True
            while len(grid[r]) <= c:
                grid[r].append(None)
            grid[r][c] = cell
            if rs > 1 or cs > 1:
                merges.append((r, c, rs, cs))
            c += cs
    n_rows = max([r + 1 for (r, _c) in taken] + [len(rows)])
    n_cols = max([c + 1 for (_r, c) in taken] + [0])
    out = [[None] * n_cols for _ in range(n_rows)]
    for r, row in enumerate(grid):
        for c, cell in enumerate(row):
            if cell is not None:
                out[r][c] = cell
    return out, merges


def _col_widths(texts: list[list[str]], total_pt: float) -> list[float]:
    """칸 글 길이로 열 폭을 나눈다 — 한 열이 55% 를 넘지 않고, 가장 좁은 열도 글자 네 자는 들어가게."""
    n = max((len(r) for r in texts), default=0)
    if not n:
        return []
    weight = []
    for c in range(n):
        col = [r[c] for r in texts if c < len(r) and r[c]]
        longest = max((max((len(ln) for ln in t.split("\n")), default=0) for t in col), default=1)
        mean = sum(len(t) for t in col) / max(len(col), 1)
        weight.append(max(4.0, min(float(longest), mean * 1.6 + 4)))
    s = sum(weight)
    share = [w / s for w in weight]
    cap = 0.55 if n > 1 else 1.0
    for _ in range(3):                                   # 넘친 몫을 나머지에 나눠 준다
        over = sum(max(0.0, x - cap) for x in share)
        share = [min(x, cap) for x in share]
        room = [cap - x for x in share]
        if over <= 1e-9 or sum(room) <= 0:
            break
        share = [x + over * r / sum(room) for x, r in zip(share, room)]
    floor = min(1.0 / n, 4 * BODY_PT / total_pt)
    share = [max(x, floor) for x in share]
    s = sum(share)
    return [total_pt * x / s for x in share]


class Converter:
    def __init__(self, base_dir: Path | None = None, images: dict[str, bytes] | None = None) -> None:
        from docx import Document

        self.doc = Document()
        self.font = _font()
        _base_styles(self.doc, self.font)
        self.base_dir = base_dir
        self.images = images or {}
        self.inline = _Inline(self)
        self.stats = {"headings": 0, "paragraphs": 0, "list_items": 0, "tables": 0, "merged_cells": 0, "images": 0,
                      "images_missing": 0, "code_blocks": 0}

    # -- 블록 ------------------------------------------------------------------
    def run(self, md: str) -> bytes:
        from markdown_it import MarkdownIt

        mdi = MarkdownIt("commonmark", {"html": True}).enable(["table", "strikethrough"])
        tokens = mdi.parse(md)
        self._blocks(tokens, 0, len(tokens), self.doc, list_stack=[])
        self._drop_leading_empty()
        buf = io.BytesIO()
        self.doc.save(buf)
        return buf.getvalue()

    def _drop_leading_empty(self) -> None:
        body = self.doc.element.body
        first = body[0] if len(body) else None
        if first is not None and first.tag.endswith("}p") and not "".join(first.itertext()).strip() and len(body) > 2:
            body.remove(first)

    def _close_of(self, tokens, i: int) -> int:
        """열린 토큰 i 의 짝 닫힘 위치."""
        depth = 0
        for j in range(i, len(tokens)):
            depth += tokens[j].nesting
            if depth == 0:
                return j
        return len(tokens) - 1

    def _blocks(self, tokens, start: int, end: int, container, list_stack: list[dict]) -> None:
        i = start
        while i < end:
            t = tokens[i]
            if t.type == "heading_open":
                lvl = int(t.tag[1])
                p = container.add_paragraph(style=f"Heading {min(lvl, 6)}") if container is self.doc else container.add_paragraph()
                self.inline.emit(p, tokens[i + 1].children, HEADING_PT[min(lvl, 6)] if container is not self.doc else None,
                                 bold_all=container is not self.doc)
                self.stats["headings"] += 1
                i += 3
            elif t.type == "paragraph_open":
                inline = tokens[i + 1]
                if list_stack:
                    self._list_para(container, inline, list_stack, first=not list_stack[-1]["started"])
                    list_stack[-1]["started"] = True
                else:
                    self._plain_para(container, inline)
                i += 3
            elif t.type in ("bullet_list_open", "ordered_list_open"):
                close = self._close_of(tokens, i)
                start_n = int(t.attrs.get("start") or 1) if t.type == "ordered_list_open" else 1
                list_stack.append({"ordered": t.type == "ordered_list_open", "n": start_n - 1, "started": False})
                self._blocks(tokens, i + 1, close, container, list_stack)
                list_stack.pop()
                i = close + 1
            elif t.type == "list_item_open":
                close = self._close_of(tokens, i)
                list_stack[-1]["n"] += 1
                list_stack[-1]["started"] = False
                self._blocks(tokens, i + 1, close, container, list_stack)
                if not list_stack[-1]["started"]:                         # 빈 항목
                    self._list_line(container, "", list_stack)
                i = close + 1
            elif t.type == "table_open":
                close = self._close_of(tokens, i)
                self._md_table(tokens, i, close, container)
                i = close + 1
            elif t.type == "html_block":
                if "<table" in t.content.lower():
                    self._html_table(t.content, container)
                elif re.search(r"<img\s", t.content, re.I):
                    for src, alt in re.findall(r"<img[^>]*?src=\"([^\"]+)\"[^>]*?(?:alt=\"([^\"]*)\")?", t.content, re.I):
                        self.image(container.add_paragraph(), src, alt)
                i += 1
            elif t.type in ("fence", "code_block"):
                self._code(container, t.content.rstrip("\n"))
                i += 1
            elif t.type == "blockquote_open":
                close = self._close_of(tokens, i)
                self._quote(tokens, i + 1, close, container)
                i = close + 1
            elif t.type == "hr":
                self._rule(container)
                i += 1
            else:
                i += 1

    def _plain_para(self, container, inline) -> None:
        p = container.add_paragraph()
        self.inline.emit(p, inline.children)
        self.stats["paragraphs"] += 1

    def _list_marker(self, list_stack: list[dict]) -> str:
        lvl = len(list_stack) - 1
        top = list_stack[-1]
        if top["ordered"]:
            depth = sum(1 for s in list_stack if s["ordered"]) - 1
            return _number(depth, top["n"])
        depth = sum(1 for s in list_stack if not s["ordered"]) - 1
        return BULLETS[depth % len(BULLETS)] if lvl >= 0 else ""

    def _hanging(self, p, list_stack: list[dict], marker: str) -> None:
        """내어쓰기 — 왼쪽 여백 = 단계 × 들여쓰기 + 걸이, 첫 줄 = −걸이, 탭 자리 = 왼쪽 여백(부호 뒤 탭이 내용 시작으로 간다)."""
        from docx.shared import Pt

        lvl = len(list_stack) - 1
        hang = max(1.0, len(marker)) * BODY_PT * (0.6 if marker.isascii() else 1.0) + BODY_PT * 0.5
        left = lvl * INDENT_PT + hang
        pf = p.paragraph_format
        pf.left_indent = Pt(left)
        pf.first_line_indent = Pt(-hang)
        pf.tab_stops.add_tab_stop(Pt(left))
        pf.space_after = Pt(2)

    def _list_line(self, container, text: str, list_stack: list[dict]):
        p = container.add_paragraph()
        marker = self._list_marker(list_stack)
        r = p.add_run(marker + "\t")
        _set_run_font(r, self.font)
        self._hanging(p, list_stack, marker)
        self.stats["list_items"] += 1
        return p

    def _list_para(self, container, inline, list_stack: list[dict], first: bool) -> None:
        children = list(inline.children or [])
        task = None
        if first and children and children[0].type == "text":
            m = re.match(r"^\[([ xX])\]\s+", children[0].content)
            if m:
                task = "☑" if m.group(1).lower() == "x" else "☐"
                children[0].content = children[0].content[m.end():]
        if first:
            p = self._list_line(container, "", list_stack)
            if task:
                r = p.add_run(task + " ")
                _set_run_font(r, self.font)
        else:                                                   # 같은 항목의 둘째 문단 — 부호 없이 내용 자리에 맞춘다
            from docx.shared import Pt

            p = container.add_paragraph()
            marker = self._list_marker(list_stack)
            self._hanging(p, list_stack, marker)
            p.paragraph_format.first_line_indent = Pt(0)
        self.inline.emit(p, children)

    def _quote(self, tokens, start: int, end: int, container) -> None:
        """인용 — 왼쪽 굵은 선과 들여쓰기(독스가 문단 테두리를 받는다)."""
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        from docx.shared import Pt

        before = len(container.paragraphs) if hasattr(container, "paragraphs") else 0
        self._blocks(tokens, start, end, container, list_stack=[])
        for p in container.paragraphs[before:]:
            p.paragraph_format.left_indent = Pt(12)
            ppr = p._p.get_or_add_pPr()
            bdr = OxmlElement("w:pBdr")
            left = OxmlElement("w:left")
            for k, v in (("w:val", "single"), ("w:sz", "18"), ("w:space", "8"), ("w:color", "9AA7B8")):
                left.set(qn(k), v)
            bdr.append(left)
            ppr.append(bdr)
            for r in p.runs:
                r.font.color.rgb = __import__("docx.shared", fromlist=["RGBColor"]).RGBColor(0x44, 0x4B, 0x55)

    def _rule(self, container) -> None:
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn

        p = container.add_paragraph()
        ppr = p._p.get_or_add_pPr()
        bdr = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        for k, v in (("w:val", "single"), ("w:sz", "6"), ("w:space", "1"), ("w:color", "9AA7B8")):
            bottom.set(qn(k), v)
        bdr.append(bottom)
        ppr.append(bdr)

    def _code(self, container, text: str) -> None:
        """코드 — 음영 한 칸 표(독스가 문단 음영보다 칸 음영을 안정되게 받는다)."""
        from docx.shared import Pt

        t = container.add_table(rows=1, cols=1)
        _table_fixed_layout(t, [int(_text_width_pt() * 20)])
        cell = t.cell(0, 0)
        _set_cell_borders(cell, BorderFill(sides={s: ("SOLID", 0.12) for s in ("left", "right", "top", "bottom")}, fill=CODE_FILL))
        _set_cell_margins(cell, 400)
        lines = text.split("\n") or [""]
        cell.paragraphs[0].text = ""
        for k, ln in enumerate(lines):
            p = cell.paragraphs[0] if k == 0 else cell.add_paragraph()
            p.paragraph_format.line_spacing = 1.0
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(ln)
            _set_run_font(r, "Courier New", BODY_PT - 1.5)
        self.stats["code_blocks"] += 1
        container.add_paragraph()

    # -- 표 ----------------------------------------------------------------------
    def _md_table(self, tokens, start: int, end: int, container) -> None:
        rows: list[list[dict]] = []
        aligns: list[str] = []
        head_rows = 0
        cur: list[dict] | None = None
        in_head = False
        for k in range(start, end):
            t = tokens[k]
            if t.type == "thead_open":
                in_head = True
            elif t.type == "thead_close":
                in_head = False
            elif t.type == "tr_open":
                cur = []
                rows.append(cur)
                if in_head:
                    head_rows += 1
            elif t.type in ("th_open", "td_open"):
                style = (t.attrs.get("style") or "")
                aligns_here = "center" if "center" in style else "right" if "right" in style else "left"
                if len(rows) == 1:
                    aligns.append(aligns_here)
                inline = tokens[k + 1]
                cur.append({"inline": inline.children, "text": inline.content, "rowspan": 1, "colspan": 1, "head": in_head})
        self._table(container, rows, head_rows, aligns)

    def _html_table(self, html: str, container) -> None:
        parser = _HtmlTable()
        parser.feed(html)
        rows = [r for r in parser.rows if r]
        if not rows:
            return
        head_rows = 0
        for r in rows:
            if all(c["head"] for c in r):
                head_rows += 1
            else:
                break
        self._table(container, rows, head_rows, [])

    def _table(self, container, rows: list[list[dict]], head_rows: int, aligns: list[str]) -> None:
        from docx.enum.table import WD_ALIGN_VERTICAL
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.shared import Pt

        grid, merges = _grid(rows)
        n_rows, n_cols = len(grid), max((len(r) for r in grid), default=0)
        if not n_rows or not n_cols:
            return
        texts = [[(c["text"] if c else "") for c in r] for r in grid]
        widths = _col_widths(texts[head_rows:] or texts, _text_width_pt())
        # 라벨 열 — 첫 열 칸이 모두 짧은 글이고, 나머지 열에 수치나 긴 글이 있을 때만(표 모양에서 판정하는 일반 규칙)
        body = texts[head_rows:]
        first = [r[0] for r in body if r and r[0]]
        rest = [x for r in body for x in r[1:] if x]
        label_col = (n_cols > 1 and len(body) >= 2 and first and all(len(x) <= 16 and not _NUMERIC.match(x) for x in first)
                     and rest and (sum(1 for x in rest if _NUMERIC.match(x) or len(x) > 16) / len(rest)) >= 0.5)
        table = container.add_table(rows=n_rows, cols=n_cols)
        _table_fixed_layout(table, [int(w * 20) for w in widths])
        thin = {s: ("SOLID", 0.12) for s in ("left", "right", "top", "bottom")}
        for r in range(n_rows):
            is_head = r < head_rows
            for c in range(n_cols):
                cell = table.cell(r, c)
                cell.width = Pt(widths[c])
                fill = HEAD_FILL if is_head else (LABEL_FILL if label_col and c == 0 else "")
                _set_cell_borders(cell, BorderFill(sides=dict(thin), fill=fill))
                _set_cell_margins(cell, {"top": 100, "bottom": 100, "left": 280, "right": 280})
                cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                info = grid[r][c] if c < len(grid[r]) else None
                if info is None:
                    continue
                p = cell.paragraphs[0]
                p.paragraph_format.line_spacing = 1.15
                p.paragraph_format.space_after = Pt(0)
                align = aligns[c] if c < len(aligns) else ("right" if _NUMERIC.match(info["text"] or "") and not is_head else "left")
                p.alignment = (WD_ALIGN_PARAGRAPH.CENTER if is_head or align == "center" else
                               WD_ALIGN_PARAGRAPH.RIGHT if align == "right" else WD_ALIGN_PARAGRAPH.LEFT)
                size = BODY_PT - 1
                if info.get("inline") is not None:
                    self.inline.emit(p, info["inline"], size=size, bold_all=is_head or bool(label_col and c == 0))
                else:
                    for k, ln in enumerate((info["text"] or "").split("\n")):
                        q = p if k == 0 else cell.add_paragraph()
                        if k:
                            q.alignment = p.alignment
                            q.paragraph_format.space_after = Pt(0)
                        run = q.add_run(ln)
                        _set_run_font(run, self.font, size, True if (is_head or (label_col and c == 0)) else None)
        for r, c, rs, cs in merges:
            a = table.cell(r, c)
            b = table.cell(min(r + rs, n_rows) - 1, min(c + cs, n_cols) - 1)
            keep = [p for p in a.paragraphs]
            merged = a.merge(b)
            # python-docx 는 덮인 칸의 빈 문단을 이어 붙인다 — 머리 칸 글만 남긴다
            for p in list(merged.paragraphs)[len(keep):]:
                if not p.text.strip():
                    p._p.getparent().remove(p._p)
            self.stats["merged_cells"] += 1
        self.stats["tables"] += 1
        if container is self.doc:
            container.add_paragraph().paragraph_format.space_after = Pt(2)

    # -- 그림 --------------------------------------------------------------------
    def _image_bytes(self, src: str) -> bytes | None:
        if src in self.images:
            return self.images[src]
        m = re.match(r"data:image/[\w.+-]+;base64,(.+)$", src, re.S)
        if m:
            try:
                return base64.b64decode(m.group(1))
            except ValueError:
                return None
        if self.base_dir and not re.match(r"^[a-z]+://", src):
            path = (self.base_dir / src).resolve()
            if self.base_dir.resolve() in path.parents and path.is_file():
                return path.read_bytes()
        return None

    def image(self, para, src: str, alt: str) -> None:
        from docx.shared import Pt

        from zzaimy.ingest.hwpx_docx import normalize_image

        data = self._image_bytes(src)
        if data is None:
            r = para.add_run(f"[그림: {alt or src[:40]}]")
            _set_run_font(r, self.font, color="666666")
            self.stats["images_missing"] += 1
            return
        try:
            data = normalize_image(data)
            para.add_run().add_picture(io.BytesIO(data), width=Pt(_text_width_pt()) if _wide(data) else None)
            self.stats["images"] += 1
        except Exception:
            r = para.add_run(f"[그림: {alt or '읽을 수 없음'}]")
            _set_run_font(r, self.font, color="666666")
            self.stats["images_missing"] += 1


def _wide(data: bytes) -> bool:
    """그림이 본문 폭(96dpi 기준)보다 넓은가 — 넓으면 본문 폭으로 줄인다."""
    from zzaimy.ingest.hwpx_fill import _image_px

    px = _image_px(data)
    return bool(px and px[0] * 0.75 > _text_width_pt())


def convert(md: str, base_dir: Path | str | None = None, images: dict[str, bytes] | None = None) -> tuple[bytes, dict]:
    """마크다운 글 → (docx 바이트, 통계). base_dir 은 상대 경로 그림을 찾을 곳(그 밖으로는 나가지 않는다)."""
    conv = Converter(Path(base_dir) if base_dir else None, images)
    data = conv.run(md)
    return data, conv.stats
