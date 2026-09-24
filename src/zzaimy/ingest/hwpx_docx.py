"""HWPX(OWPML) → DOCX 충실 변환 — 구글 독스에서 여러 사람과 에이전트가 같이 작업하기 위한 열람·편집본.

추출 조각으로 되살리던 이전 경로(render.build_docx)는 표 테두리·바탕색·정렬·글자 크기가 사라지고 중첩 표의 글이
두 번 들어갔다(실측 2026-09-24, AID 사업계획서 작성서식). 여기서는 `Contents/header.xml` 의 글자 모양(charPr)·
문단 모양(paraPr)·테두리/배경(borderFill)과 `Contents/section*.xml` 의 문단·표·글상자·그림을 그대로 옮긴다.

옮기는 것: 쪽 크기·여백·가로세로, 문단 정렬·들여쓰기·쪽 나눔, 글자 크기·굵기·기울임·밑줄·색, 표(열 너비·행 높이·
병합·셀 테두리·바탕색·세로 정렬·중첩 표), 글상자(hp:rect 안의 글), 그림(BinData), 줄 바꿈·탭, 개요 문단은 제목 스타일.
안 옮기는 것: 머리말·꼬리말·쪽 번호·각주(독스로 넘어가도 편집 대상이 아니다), 그림 위 좌표 배치(흐름 순서로 넣는다).

단위: HWPUNIT = 1/7200 인치. charPr height = 1/100 pt.
"""

from __future__ import annotations

import io
import re
import zipfile
from dataclasses import dataclass, field
from pathlib import Path
from xml.etree import ElementTree as ET

try:  # XML 폭탄·외부 개체 방어
    from defusedxml.ElementTree import fromstring as _fromstring
except Exception:  # pragma: no cover
    from xml.etree.ElementTree import fromstring as _fromstring

_SECTION_RE = re.compile(r"Contents/section(\d+)\.xml$")
HWPUNIT_PER_INCH = 7200
EMU_PER_HWPUNIT = 914400 / HWPUNIT_PER_INCH          # 127
TWIPS_PER_HWPUNIT = 1440 / HWPUNIT_PER_INCH          # 0.2


# 구글 독스에 있는 한글 글꼴로 맞춘다 — 한글 문서의 글꼴 이름(맑은 고딕·휴먼명조·HY헤드라인M …)을 그대로 두면 독스가 Arial 로
# 대신 그려 한글 모양이 달라진다(실측 2026-09-24). 고딕 계열은 Nanum Gothic, 명조·바탕 계열은 Nanum Myeongjo.
_SERIF_HINT = ("명조", "바탕", "신명조", "Batang", "Myeongjo", "Myungjo", "궁서", "Gungsuh")
FONT_MAP = {"고딕": "Nanum Gothic", "명조": "Nanum Myeongjo"}


def docs_font(face: str) -> str:
    if not face:
        return ""
    f = face.strip()
    if f in ("Nanum Gothic", "Nanum Myeongjo", "Noto Sans KR", "Noto Serif KR", "Arial", "Times New Roman"):
        return f
    if any(h.lower() in f.lower() for h in _SERIF_HINT):
        return FONT_MAP["명조"]
    return FONT_MAP["고딕"]


def _local(tag: str) -> str:
    return tag.rsplit("}", 1)[-1] if "}" in tag else tag.split(":")[-1]


def _children(el: ET.Element, name: str):
    return [ch for ch in el if _local(ch.tag) == name]


def _first(el: ET.Element, name: str) -> ET.Element | None:
    for node in el.iter():
        if _local(node.tag) == name:
            return node
    return None


def _child(el: ET.Element, name: str) -> ET.Element | None:
    """직접 자식만 — 셀(hp:tc)의 cellAddr·cellSpan·cellSz 는 subList(중첩 표 포함) 뒤에 오므로 iter() 로 찾으면
    중첩 표의 셀 값을 집는다(실측 2026-09-24)."""
    for ch in el:
        if _local(ch.tag) == name:
            return ch
    return None


# ---- header.xml: 스타일 사전 ------------------------------------------------------------------------

@dataclass
class CharStyle:
    size_pt: float = 10.0
    bold: bool = False
    italic: bool = False
    underline: bool = False
    color: str = "000000"
    font: str = ""
    superscript: bool = False
    subscript: bool = False


@dataclass
class ParaStyle:
    align: str = "JUSTIFY"
    left_hu: int = 0          # 왼쪽 여백(HWPUNIT)
    indent_hu: int = 0        # 첫 줄 들여쓰기(음수면 내어쓰기)
    before_hu: int = 0
    after_hu: int = 0
    line_pct: int = 0         # 줄 간격 %
    outline_level: int = 0    # 1~ 이면 개요(제목)
    bullet: str = ""          # 자동 글머리표 문자(한글은 글에 없고 문단 모양에 있다)


@dataclass
class BorderFill:
    sides: dict = field(default_factory=dict)   # left/right/top/bottom → (type, width_mm)
    fill: str = ""                              # 바탕색 hex(없으면 "")


@dataclass
class Styles:
    fonts: dict = field(default_factory=dict)
    chars: dict = field(default_factory=dict)
    paras: dict = field(default_factory=dict)
    borders: dict = field(default_factory=dict)
    heading_styles: dict = field(default_factory=dict)   # styleIDRef → outline level
    bullets: dict = field(default_factory=dict)          # bullet id → 문자


def _hu(value: str | None) -> int:
    try:
        return int(float(value or 0))
    except ValueError:
        return 0


def load_styles(root: ET.Element) -> Styles:
    st = Styles()
    for f in root.iter():
        if _local(f.tag) == "font" and f.get("id") is not None:
            st.fonts[str(f.get("id"))] = f.get("face") or ""
    for b in root.iter():
        if _local(b.tag) == "bullet" and b.get("id") is not None:
            st.bullets[str(b.get("id"))] = (b.get("char") or "").strip()
    for cp in root.iter():
        name = _local(cp.tag)
        if name == "charPr":
            cs = CharStyle()
            cs.size_pt = _hu(cp.get("height")) / 100.0 or 10.0
            color = (cp.get("textColor") or "#000000").lstrip("#")
            cs.color = color if re.fullmatch(r"[0-9A-Fa-f]{6}", color) else "000000"
            for ch in cp:
                n = _local(ch.tag)
                if n == "bold":
                    cs.bold = True
                elif n == "italic":
                    cs.italic = True
                elif n == "underline" and (ch.get("type") or "NONE") != "NONE":
                    cs.underline = True
                elif n == "supscript":
                    cs.superscript = True
                elif n == "subscript":
                    cs.subscript = True
                elif n == "fontRef":
                    cs.font = st.fonts.get(str(ch.get("hangul")), "")
            st.chars[str(cp.get("id"))] = cs
        elif name == "paraPr":
            ps = ParaStyle()
            for ch in cp.iter():
                n = _local(ch.tag)
                if n == "align":
                    ps.align = ch.get("horizontal") or "JUSTIFY"
                elif n == "heading" and (ch.get("type") or "NONE") == "OUTLINE":
                    ps.outline_level = int(ch.get("level") or 0) + 1
                elif n == "heading" and (ch.get("type") or "NONE") == "BULLET":
                    ps.bullet = st.bullets.get(str(ch.get("idRef")), "")
                elif n == "left":
                    ps.left_hu = _hu(ch.get("value"))
                elif n == "intent":
                    ps.indent_hu = _hu(ch.get("value"))
                elif n == "prev":
                    ps.before_hu = _hu(ch.get("value"))
                elif n == "next":
                    ps.after_hu = _hu(ch.get("value"))
                elif n == "lineSpacing" and (ch.get("type") or "") == "PERCENT":
                    ps.line_pct = _hu(ch.get("value"))
            st.paras[str(cp.get("id"))] = ps
        elif name == "borderFill":
            bf = BorderFill()
            for ch in cp.iter():
                n = _local(ch.tag)
                if n in ("leftBorder", "rightBorder", "topBorder", "bottomBorder"):
                    w = ch.get("width") or "0.12 mm"
                    try:
                        mm = float(w.replace("mm", "").strip())
                    except ValueError:
                        mm = 0.12
                    bf.sides[n[:-6]] = ((ch.get("type") or "NONE").upper(), mm)
                elif n == "winBrush":
                    face = (ch.get("faceColor") or "").lstrip("#")
                    alpha = ch.get("alpha") or "0"
                    # 한글은 채우기 없음을 검정(#000000)·alpha 0 으로도 적는다 — 검정 바탕은 표에 쓰지 않으므로 없음으로 본다
                    if re.fullmatch(r"[0-9A-Fa-f]{6}", face) and face.upper() not in ("000000", "FFFFFF") and alpha in ("0", "255"):
                        bf.fill = face.upper()
            st.borders[str(cp.get("id"))] = bf
        elif name == "style" and cp.get("id") is not None:
            pid = str(cp.get("paraPrIDRef"))
            if pid in st.paras and st.paras[pid].outline_level:
                st.heading_styles[str(cp.get("id"))] = st.paras[pid].outline_level
    return st


# ---- DOCX 저수준 도우미 --------------------------------------------------------------------------------

_TCPR_ORDER = ("cnfStyle", "tcW", "gridSpan", "hMerge", "vMerge", "tcBorders", "shd", "noWrap", "tcMar", "textDirection", "tcFitText",
               "vAlign", "hideMark")


def _set_tcpr(tcPr, el) -> None:
    from docx.oxml.ns import qn

    name = el.tag.split("}")[1]
    for old in tcPr.findall(qn(f"w:{name}")):
        tcPr.remove(old)
    rank = _TCPR_ORDER.index(name) if name in _TCPR_ORDER else len(_TCPR_ORDER)
    for i, ch in enumerate(list(tcPr)):
        cname = ch.tag.split("}")[1]
        crank = _TCPR_ORDER.index(cname) if cname in _TCPR_ORDER else len(_TCPR_ORDER)
        if crank > rank:
            tcPr.insert(i, el)
            return
    tcPr.append(el)


def _set_cell_borders(cell, bf: BorderFill | None) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    tcPr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for side in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{side}")
        kind, mm = (bf.sides.get(side) if bf else None) or ("SOLID", 0.12)
        if kind == "NONE":
            el.set(qn("w:val"), "nil")
        else:
            val = {"DASH": "dashed", "DOT": "dotted", "DOUBLE_SLIM": "double", "SLIM_THICK": "thinThickSmallGap",
                   "THICK_SLIM": "thickThinSmallGap"}.get(kind, "single")
            el.set(qn("w:val"), val)
            el.set(qn("w:sz"), str(max(2, int(round(mm / 25.4 * 72 * 8)))))   # 1/8 pt
            el.set(qn("w:space"), "0")
            el.set(qn("w:color"), "000000")
        borders.append(el)
    _set_tcpr(tcPr, borders)
    if bf and bf.fill:
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), bf.fill)
        _set_tcpr(tcPr, shd)


def _set_cell_valign(cell, valign: str) -> None:
    from docx.enum.table import WD_ALIGN_VERTICAL

    cell.vertical_alignment = {"CENTER": WD_ALIGN_VERTICAL.CENTER, "BOTTOM": WD_ALIGN_VERTICAL.BOTTOM}.get(
        valign, WD_ALIGN_VERTICAL.TOP)


def _set_cell_margins(cell, hu: int = 141) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    tcPr = cell._tc.get_or_add_tcPr()
    mar = OxmlElement("w:tcMar")
    for side in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:w"), str(int(hu * TWIPS_PER_HWPUNIT)))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    _set_tcpr(tcPr, mar)


# w:tblPr 자식의 규격(OOXML) 순서 — 순서가 틀리거나 같은 요소가 둘이면 구글 독스는 표 속성을 통째로 버려 표가 가늘게 찌그러진다
# (실측 2026-09-25: tblLayout 둘 + tblBorders 가 뒤에 → 독스 199쪽, 표가 한 줄 폭). 리브레오피스는 눈감아 준다.
_TBLPR_ORDER = ("tblStyle", "tblpPr", "tblOverlap", "bidiVisual", "tblStyleRowBandSize", "tblStyleColBandSize", "tblW", "jc",
                "tblCellSpacing", "tblInd", "tblBorders", "shd", "tblLayout", "tblCellMar", "tblLook", "tblCaption", "tblDescription")


def _set_tblpr(tblPr, el) -> None:
    """같은 이름의 기존 요소는 지우고 규격 순서 자리에 넣는다."""
    from docx.oxml.ns import qn

    name = el.tag.split("}")[1]
    for old in tblPr.findall(qn(f"w:{name}")):
        tblPr.remove(old)
    rank = _TBLPR_ORDER.index(name) if name in _TBLPR_ORDER else len(_TBLPR_ORDER)
    for i, ch in enumerate(list(tblPr)):
        cname = ch.tag.split("}")[1]
        crank = _TBLPR_ORDER.index(cname) if cname in _TBLPR_ORDER else len(_TBLPR_ORDER)
        if crank > rank:
            tblPr.insert(i, el)
            return
    tblPr.append(el)


def _table_fixed_layout(table, col_twips: list[int]) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    tblPr = table._tbl.tblPr
    total = sum(col_twips)
    tblw = OxmlElement("w:tblW")
    tblw.set(qn("w:w"), str(total))
    tblw.set(qn("w:type"), "dxa")
    _set_tblpr(tblPr, tblw)
    # 표 자체 테두리는 셀마다 정하므로 표 기본 테두리는 없앤다
    borders = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:val"), "nil")
        borders.append(el)
    _set_tblpr(tblPr, borders)
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    _set_tblpr(tblPr, layout)
    grid = table._tbl.tblGrid
    for gc, w in zip(grid.findall(qn("w:gridCol")), col_twips):
        gc.set(qn("w:w"), str(w))


MAX_ROW_HU = 5 * HWPUNIT_PER_INCH     # 행 높이 상한 5인치 — 쪽 전체를 차지하는 배치용 표 행이 빈 쪽을 만든다(실측 2026-09-24)


def _row_height(row, hu: int) -> None:
    from docx.enum.table import WD_ROW_HEIGHT_RULE
    from docx.shared import Emu

    if hu > 0:
        row.height = Emu(int(min(hu, MAX_ROW_HU) * EMU_PER_HWPUNIT))
        row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST


# ---- 본문 걷기 -------------------------------------------------------------------------------------

class Converter:
    def __init__(self, styles: Styles, zf: zipfile.ZipFile, bin_map: dict[str, str]) -> None:
        self.st = styles
        self.zf = zf
        self.bin_map = bin_map
        from docx import Document

        self.doc = Document()
        self._first_section = True
        self._trailing_empty: list = []
        self._just_sectioned = False
        self._prev_pb_empty = None          # 직전 쪽 나눔 문단이 빈 문단이면 그 요소
        self._content_since_pb = True       # 직전 쪽 나눔 뒤에 글·표·그림이 있었는가
        self.stats = {"paragraphs": 0, "tables": 0, "nested_tables": 0, "images": 0, "textboxes": 0}

    # -- 문단 ---------------------------------------------------------------------------------------
    def _apply_para_style(self, para, p_el: ET.Element) -> None:
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.shared import Emu, Pt

        ps = self.st.paras.get(str(p_el.get("paraPrIDRef")))
        if ps is None:
            return
        if ps.bullet:
            para._zz_bullet = ps.bullet
        para.alignment = {"LEFT": WD_ALIGN_PARAGRAPH.LEFT, "CENTER": WD_ALIGN_PARAGRAPH.CENTER,
                          "RIGHT": WD_ALIGN_PARAGRAPH.RIGHT, "DISTRIBUTE": WD_ALIGN_PARAGRAPH.DISTRIBUTE,
                          "DISTRIBUTE_SPACE": WD_ALIGN_PARAGRAPH.DISTRIBUTE}.get(ps.align, WD_ALIGN_PARAGRAPH.JUSTIFY)
        pf = para.paragraph_format
        # 한글의 내어쓰기(intent<0)는 첫 줄이 왼쪽 여백에서 시작하고 둘째 줄부터 |intent| 만큼 들어간다 — 워드로는
        # 왼쪽 여백을 |intent| 만큼 늘리고 첫 줄을 그만큼 되돌린다(그대로 옮기면 첫 줄이 칸 밖으로 나간다, 실측 2026-09-24 평가편람)
        left, first = ps.left_hu, ps.indent_hu
        if first < 0:
            left, first = left + abs(first), -abs(first)
        if left:
            pf.left_indent = Emu(int(left * EMU_PER_HWPUNIT))
        if first:
            pf.first_line_indent = Emu(int(first * EMU_PER_HWPUNIT))
        pf.space_before = Emu(int(ps.before_hu * EMU_PER_HWPUNIT))
        pf.space_after = Emu(int(ps.after_hu * EMU_PER_HWPUNIT))
        if ps.line_pct:
            # 한글의 줄 간격 %는 글자 크기 기준이다(10pt·160% = 16pt). 워드의 배수는 글꼴 행 높이 기준이라 맑은 고딕에서 2할쯤
            # 커져 쪽이 넘친다(실측: 표지 뒤 빈 쪽) — 글자 크기 × 비율을 고정 값으로 준다
            from docx.enum.text import WD_LINE_SPACING

            # 한글 양식은 여백용 빈 문단을 0.5pt 글자로 둔다(표지, 실측 2026-09-24) — 하한을 크게 두면 표지가 넘친다
            size_pt = self._para_font_pt(p_el)
            pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
            pf.line_spacing = Pt(max(1.0, size_pt * max(0.8, min(ps.line_pct / 100.0, 3.0))))
        if p_el.get("pageBreak") == "1" and not self._just_sectioned:
            pf.page_break_before = True

    def _para_font_pt(self, p_el: ET.Element) -> float:
        """문단의 글자 크기 — 첫 런의 글자 모양(없으면 10pt)."""
        for run in _children(p_el, "run"):
            cs = self.st.chars.get(str(run.get("charPrIDRef")))
            if cs is not None:
                return max(cs.size_pt, 0.5)
        return 10.0

    def _apply_run_style(self, run, char_id: str | None) -> None:
        from docx.oxml.ns import qn
        from docx.shared import Pt, RGBColor

        cs = self.st.chars.get(str(char_id))
        if cs is None:
            return
        run.font.size = Pt(max(cs.size_pt, 4.0))
        run.font.bold = cs.bold
        run.font.italic = cs.italic
        if cs.underline:
            run.font.underline = True
        if cs.color and cs.color.upper() != "000000":
            run.font.color.rgb = RGBColor.from_string(cs.color.upper())
        if cs.superscript:
            run.font.superscript = True
        if cs.subscript:
            run.font.subscript = True
        face = docs_font(cs.font) if cs.font else FONT_MAP["고딕"]
        run.font.name = face
        rpr = run._r.get_or_add_rPr()
        rfonts = rpr.find(qn("w:rFonts"))
        if rfonts is not None:
            rfonts.set(qn("w:eastAsia"), face)

    def _emit_text(self, para, t_el: ET.Element, char_id: str | None) -> None:
        """hp:t 안의 글·줄바꿈·탭을 런으로."""
        bullet = getattr(para, "_zz_bullet", "")
        if bullet and not para.runs:
            self._apply_run_style(para.add_run(bullet + " "), char_id)
        if t_el.text:
            self._apply_run_style(para.add_run(t_el.text), char_id)
        for ch in t_el:
            n = _local(ch.tag)
            if n == "lineBreak":
                para.add_run().add_break()
            elif n == "tab":
                para.add_run("\t")
            if ch.tail:
                self._apply_run_style(para.add_run(ch.tail), char_id)

    def paragraph(self, p_el: ET.Element, container, heading_ok: bool = True) -> None:
        """hp:p 하나 → 문단(그 안의 표·글상자·그림은 문단 뒤에 이어서)."""
        level = self.st.heading_styles.get(str(p_el.get("styleIDRef")), 0)
        ps = self.st.paras.get(str(p_el.get("paraPrIDRef")))
        if not level and ps and ps.outline_level:
            level = ps.outline_level
        pending: list[tuple[str, ET.Element]] = []
        para = None
        is_pb = container is self.doc and p_el.get("pageBreak") == "1" and not self._just_sectioned
        if is_pb:
            # 쪽 나눔 앞의 빈 문단들은 쪽 끝 여백일 뿐이다. 직전 쪽 나눔 뒤에 아무 내용도 없었다면 그 쪽은 빈 쪽 — 앞 쪽 나눔 문단까지 지운다
            for old_p in self._trailing_empty:
                el = old_p._element
                if el.getparent() is not None:
                    el.getparent().remove(el)
            self._trailing_empty = []
            if not self._content_since_pb and self._prev_pb_empty is not None:
                el = self._prev_pb_empty._element
                if el.getparent() is not None:
                    el.getparent().remove(el)
            self._prev_pb_empty = None
            self._content_since_pb = False
        for run in _children(p_el, "run"):
            char_id = run.get("charPrIDRef")
            for obj in run:
                n = _local(obj.tag)
                if n == "t":
                    if para is None:
                        para = self._new_para(container, level if heading_ok else 0)
                        self._apply_para_style(para, p_el)
                    self._emit_text(para, obj, char_id)
                elif n == "pic" and (_child(obj, "pos") is not None and (_child(obj, "pos").get("treatAsChar") or "0") == "1"):
                    # 글자처럼 놓인 그림(표지의 로고 여러 개)은 같은 문단 안에 나란히 — 문단마다 따로 두면 쪽이 넘친다(실측 2026-09-24)
                    if para is None:
                        para = self._new_para(container, 0)
                        self._apply_para_style(para, p_el)
                    self.picture(obj, container, para=para)
                elif n == "line":
                    continue                                   # 장식 선 — 내용이 아니다(간지 쪽의 선만 있는 문단이 내용으로 잡혀 빈 쪽을 만들었다)
                elif n in ("tbl", "pic"):
                    pending.append((n, obj))
                elif n in ("rect", "container", "ellipse", "polygon", "curve", "arc", "ole", "equation"):
                    has_text = any((t.text or "").strip() for t in obj.iter() if _local(t.tag) == "t")
                    has_obj = any(_local(x.tag) in ("tbl", "pic") for x in obj.iter())
                    if has_text or has_obj:
                        pending.append((n, obj))
                elif n in ("secPr",):
                    self._section(obj)
                    self._just_sectioned = True          # 구역 시작이 이미 새 쪽이다 — 같은 문단의 쪽 나눔은 겹치지 않게
        if (para is None or not para.text.strip()) and not pending:
            # 빈 문단(공백만 있는 문단 포함) — 원본의 줄 간격을 지킨다
            if para is None:
                para = self._new_para(container, 0)
                self._apply_para_style(para, p_el)
            if container is self.doc:
                if is_pb:
                    self._prev_pb_empty = para
                else:
                    self._trailing_empty.append(para)
            self._just_sectioned = False
            return
        if container is self.doc:
            self._trailing_empty = []
            self._content_since_pb = True
        for n, obj in pending:
            if n == "tbl":
                self.table(obj, container)
            elif n == "pic":
                self.picture(obj, container)
            elif n == "line":
                continue
            else:
                self.shape(obj, container)
        if pending and container is self.doc:
            self._trailing_empty = []
        self._just_sectioned = False

    def _new_para(self, container, level: int):
        if level and container is self.doc:
            self.stats["paragraphs"] += 1
            return self.doc.add_heading("", level=min(level, 4))
        self.stats["paragraphs"] += 1
        return container.add_paragraph()

    # -- 구역(쪽 설정) ------------------------------------------------------------------------------
    def _section(self, sec_el: ET.Element) -> None:
        from docx.enum.section import WD_ORIENT
        from docx.shared import Emu

        page = _first(sec_el, "pagePr")
        if page is None:
            return
        if not self._first_section:
            for old_p in self._trailing_empty:           # 구역 끝의 빈 문단은 쪽 끝 여백일 뿐 — 남기면 빈 쪽
                el = old_p._element
                if el.getparent() is not None:
                    el.getparent().remove(el)
            self._trailing_empty = []
        section = self.doc.sections[-1] if self._first_section else self.doc.add_section()
        self._first_section = False
        w, h = _hu(page.get("width")), _hu(page.get("height"))
        if w and h:
            section.page_width, section.page_height = Emu(int(w * EMU_PER_HWPUNIT)), Emu(int(h * EMU_PER_HWPUNIT))
            section.orientation = WD_ORIENT.LANDSCAPE if w > h else WD_ORIENT.PORTRAIT
        m = _first(page, "margin")
        if m is not None:
            section.left_margin = Emu(int((_hu(m.get("left")) + _hu(m.get("gutter"))) * EMU_PER_HWPUNIT))
            section.right_margin = Emu(int(_hu(m.get("right")) * EMU_PER_HWPUNIT))
            section.top_margin = Emu(int((_hu(m.get("top")) + _hu(m.get("header"))) * EMU_PER_HWPUNIT))
            section.bottom_margin = Emu(int((_hu(m.get("bottom")) + _hu(m.get("footer"))) * EMU_PER_HWPUNIT))

    # -- 표 ----------------------------------------------------------------------------------------
    def table(self, tbl: ET.Element, container) -> None:
        from docx.shared import Emu

        rows = _children(tbl, "tr")
        cells_info = []
        widths: dict[int, int] = {}
        heights: dict[int, int] = {}
        for r_i, tr in enumerate(rows):
            for tc in _children(tr, "tc"):
                addr = _child(tc, "cellAddr")
                span = _child(tc, "cellSpan")
                sz = _child(tc, "cellSz")
                row = _hu(addr.get("rowAddr")) if addr is not None else r_i
                col = _hu(addr.get("colAddr")) if addr is not None else 0
                rs = max(_hu(span.get("rowSpan")) if span is not None else 1, 1)
                cs = max(_hu(span.get("colSpan")) if span is not None else 1, 1)
                w = _hu(sz.get("width")) if sz is not None else 0
                hgt = _hu(sz.get("height")) if sz is not None else 0
                if cs == 1 and w:
                    widths.setdefault(col, w)
                if rs == 1 and hgt:
                    heights.setdefault(row, hgt)
                cells_info.append((row, col, rs, cs, tc))
        n_rows = max([r + rs for r, _, rs, _, _ in cells_info] + [_hu(tbl.get("rowCnt"))])
        if n_rows < 1:
            return
        # 한글 표는 행마다 열 경계가 달라도 된다(배치용 표, 실측 2026-09-24 작성서식의 4×13 표). 워드·독스는 한 그리드를 요구하므로
        # 모든 행의 셀 경계(x 좌표)를 모아 공통 그리드를 만들고 셀을 그 그리드 열에 얹는다. 폭 정보가 없으면 cellAddr 로 돌아간다.
        grid = _column_grid(rows, cells_info)
        if grid is not None:
            col_hu, cells_info = grid
            n_cols = len(col_hu)
        else:
            n_cols = max([c + cs for _, c, _, cs, _ in cells_info] + [_hu(tbl.get("colCnt"))])
            for row, col, rs, cs, _tc in cells_info:
                if cs > 1:
                    w = _hu(_child(_tc, "cellSz").get("width")) if _child(_tc, "cellSz") is not None else 0
                    missing = [c for c in range(col, col + cs) if c not in widths]
                    if missing and w:
                        known = sum(widths.get(c, 0) for c in range(col, col + cs))
                        share = max((w - known) // len(missing), 100)
                        for c in missing:
                            widths[c] = share
            total_hu = _hu((_child(tbl, "sz") or ET.Element("x")).get("width")) or sum(widths.values())
            col_hu = [widths.get(c, max(total_hu // n_cols, 100)) for c in range(n_cols)]
        if n_cols < 1:
            return
        table = container.add_table(rows=n_rows, cols=n_cols)
        table.autofit = False
        _table_fixed_layout(table, [int(w * TWIPS_PER_HWPUNIT) for w in col_hu])
        for c in range(n_cols):
            for row in table.rows:
                row.cells[c].width = Emu(int(col_hu[c] * EMU_PER_HWPUNIT))
        nested_rows = {row for row, _c, _rs, _cs, tc in cells_info if _first(tc, "tbl") is not None}
        for r_i, row in enumerate(table.rows):
            if r_i not in nested_rows:
                _row_height(row, heights.get(r_i, 0))
        covered: set[tuple[int, int]] = set()
        default_bf = self.st.borders.get(str(tbl.get("borderFillIDRef")))
        for row, col, rs, cs, tc in sorted(cells_info, key=lambda t: (t[0], t[1])):
            if row >= n_rows or col >= n_cols:
                continue
            if (row, col) in covered:
                # 공통 그리드로 옮기다 앞 셀의 병합 영역과 겹쳤다(경계 반올림) — 겹치지 않는 첫 칸으로 민다
                shift = col
                while shift < n_cols and (row, shift) in covered:
                    shift += 1
                if shift >= n_cols:
                    continue
                cs = max(1, cs - (shift - col))
                col = shift
            end_r, end_c = min(row + rs - 1, n_rows - 1), min(col + cs - 1, n_cols - 1)
            # 병합 영역이 이미 덮인 칸과 겹치면 겹치기 직전까지로 줄인다(워드는 겹치는 병합을 허용하지 않는다)
            while end_c > col and any((rr, end_c) in covered for rr in range(row, end_r + 1)):
                end_c -= 1
            while end_r > row and any((end_r, cc) in covered for cc in range(col, end_c + 1)):
                end_r -= 1
            cell = table.cell(row, col)
            if (end_r, end_c) != (row, col):
                try:
                    cell = cell.merge(table.cell(end_r, end_c))
                except ValueError:
                    end_r, end_c = row, col
            for rr in range(row, end_r + 1):
                for cc in range(col, end_c + 1):
                    covered.add((rr, cc))
            bf = self.st.borders.get(str(tc.get("borderFillIDRef"))) or default_bf
            _set_cell_borders(cell, bf)
            _set_cell_margins(cell)
            sub = _child(tc, "subList")
            _set_cell_valign(cell, (sub.get("vertAlign") if sub is not None else "TOP") or "TOP")
            # 셀의 첫 빈 문단을 지우고 원본 문단으로 채운다
            first_p = cell.paragraphs[0]
            if sub is not None:
                paras = _children(sub, "p")
                for i, p_el in enumerate(paras):
                    if i == 0:
                        self._fill_para_into(first_p, p_el, cell)
                    else:
                        self.paragraph(p_el, cell, heading_ok=False)
        # 병합에 덮이지 않은 빈 칸도 테두리를 준다
        for rr in range(n_rows):
            for cc in range(n_cols):
                if (rr, cc) not in covered:
                    _set_cell_borders(table.cell(rr, cc), default_bf)
        if container is self.doc:
            self.stats["tables"] += 1
            self._content_since_pb = True
            self._trailing_empty = []
        else:
            self.stats["nested_tables"] += 1
        # 표 뒤에 문단 하나(한글은 표 다음에 항상 빈 줄 없이 이어지지만 워드는 표 사이에 문단이 필요)
        if container is not self.doc:
            return

    def _fill_para_into(self, para, p_el: ET.Element, cell) -> None:
        """이미 있는 문단(셀의 첫 문단)에 hp:p 내용을 채운다."""
        self._apply_para_style(para, p_el)
        pending = []
        for run in _children(p_el, "run"):
            char_id = run.get("charPrIDRef")
            for obj in run:
                n = _local(obj.tag)
                if n == "t":
                    self._emit_text(para, obj, char_id)
                elif n in ("tbl", "pic", "rect", "container"):
                    pending.append((n, obj))
        for n, obj in pending:
            if n == "tbl":
                self.table(obj, cell)
            elif n == "pic":
                self.picture(obj, cell)
            else:
                self.shape(obj, cell)

    # -- 글상자·도형 --------------------------------------------------------------------------------
    def shape(self, el: ET.Element, container) -> None:
        """글상자(hp:rect 등)의 글은 그 자리에 문단으로. 테두리가 있으면 한 칸 표로 감싼다."""
        draw = _child(el, "drawText")
        if draw is None:
            # 묶음 도형(container)은 글이 안쪽 도형에 있다 — 자식 도형을 차례로(간지의 장 제목 상자, 실측 2026-09-24)
            handled = False
            for ch in el:
                if _local(ch.tag) in ("rect", "container", "ellipse", "polygon", "curve", "arc"):
                    self.shape(ch, container)
                    handled = True
                elif _local(ch.tag) == "tbl":
                    self.table(ch, container)
                    handled = True
                elif _local(ch.tag) == "pic":
                    self.picture(ch, container)
                    handled = True
            if not handled:
                for node in el.iter():
                    if _local(node.tag) == "tbl":
                        self.table(node, container)
            return
        sub = _child(draw, "subList")
        if sub is None:
            return
        line = _child(el, "lineShape")
        boxed = line is not None and (line.get("style") or "NONE") != "NONE"
        paras = _children(sub, "p")
        if not paras:
            return
        self.stats["textboxes"] += 1
        if boxed:
            # 글상자 너비를 표에 준다 — 너비 없는 한 칸 표는 독스가 가장 좁게 그려 쪽 높이의 가는 상자가 됐다(실측 2026-09-25 사업계획서)
            w_hu = 0
            for tag in ("curSz", "sz", "orgSz"):
                sz = _child(el, tag)
                if sz is not None and _hu(sz.get("width")) > 0:
                    w_hu = _hu(sz.get("width"))
                    break
            w_hu = min(max(w_hu, 4 * HWPUNIT_PER_INCH // 4), 47000)       # 최소 1인치, 최대 본문 폭쯤
            box = container.add_table(rows=1, cols=1)
            _table_fixed_layout(box, [int(w_hu * TWIPS_PER_HWPUNIT)])
            cell = box.cell(0, 0)
            cell.width = __import__("docx.shared", fromlist=["Emu"]).Emu(int(w_hu * EMU_PER_HWPUNIT))
            _set_cell_borders(cell, BorderFill(sides={s: ("SOLID", 0.12) for s in ("left", "right", "top", "bottom")}))
            _set_cell_margins(cell)
            for i, p_el in enumerate(paras):
                if i == 0:
                    self._fill_para_into(cell.paragraphs[0], p_el, cell)
                else:
                    self.paragraph(p_el, cell, heading_ok=False)
        else:
            for p_el in paras:
                self.paragraph(p_el, container, heading_ok=False)

    # -- 그림 --------------------------------------------------------------------------------------
    def picture(self, pic: ET.Element, container, para=None) -> None:
        from docx.shared import Emu

        ref = ""
        for node in pic.iter():
            if _local(node.tag) == "img" and node.get("binaryItemIDRef"):
                ref = node.get("binaryItemIDRef") or ""
                break
        member = self.bin_map.get(ref) or self.bin_map.get(Path(ref).stem)
        if not member:
            return
        try:
            data = normalize_image(self.zf.read(member))
        except KeyError:
            return
        cur = _child(pic, "curSz") or _child(pic, "orgSz")
        w = _hu(cur.get("width")) if cur is not None else 0
        if not w:
            org = _child(pic, "orgSz")
            w = _hu(org.get("width")) if org is not None else 0
        width = Emu(int(min(max(w, 2000), 45000) * EMU_PER_HWPUNIT))
        if para is None:
            para = container.add_paragraph()
        try:
            para.add_run().add_picture(io.BytesIO(data), width=width)
            self.stats["images"] += 1
        except Exception as e:
            self.stats.setdefault("image_errors", []).append(f"{type(e).__name__}: {e} member={member} head={data[:6]!r} len={len(data)}"[:200])
            para.add_run("[그림]")
        if container is self.doc:
            self._content_since_pb = True


GRID_TOL = 60     # 경계 좌표 합치기 허용치(HWPUNIT, 약 0.2mm)


def _column_grid(rows, cells_info):
    """모든 행의 셀 x 경계로 공통 열 그리드를 만든다 → (열 너비 목록, 그리드 열로 옮긴 cells_info). 폭이 없으면 None."""
    # 행마다 셀을 colAddr 순으로 놓고 폭을 누적해 x 범위를 구한다
    by_row: dict[int, list] = {}
    for row, col, rs, cs, tc in cells_info:
        sz = _child(tc, "cellSz")
        w = _hu(sz.get("width")) if sz is not None else 0
        if w <= 0:
            return None
        by_row.setdefault(row, []).append((col, w, rs, cs, tc))
    spans: list[tuple[int, int, int, int, object]] = []      # (row, x0, x1, rs, tc)
    bounds: list[int] = [0]
    # 세로 병합 셀은 아래 행에 자리를 차지하지만 XML 에 없다 — 아래 행의 x 누적을 맞추려면 그 자리를 건너뛰어야 한다
    occupied: dict[int, list[tuple[int, int]]] = {}         # row → [(x0, x1)] 세로 병합이 덮는 구간
    for row in sorted(by_row):
        x = 0
        for col, w, rs, cs, tc in sorted(by_row[row], key=lambda t: t[0]):
            for ox0, ox1 in sorted(occupied.get(row, [])):
                if abs(ox0 - x) <= GRID_TOL:
                    x = ox1
            x0, x1 = x, x + w
            spans.append((row, x0, x1, rs, tc))
            bounds += [x0, x1]
            for rr in range(row + 1, row + rs):
                occupied.setdefault(rr, []).append((x0, x1))
            x = x1
    # 경계 합치기
    uniq: list[int] = []
    for b in sorted(bounds):
        if not uniq or b - uniq[-1] > GRID_TOL:
            uniq.append(b)
    if len(uniq) < 2:
        return None

    def idx(v: int) -> int:
        return min(range(len(uniq)), key=lambda i: abs(uniq[i] - v))

    col_hu = [uniq[i + 1] - uniq[i] for i in range(len(uniq) - 1)]
    out = []
    for row, x0, x1, rs, tc in spans:
        c0, c1 = idx(x0), idx(x1)
        if c1 <= c0:
            c1 = c0 + 1
        out.append((row, c0, rs, c1 - c0, tc))
    return col_hu, out


IMAGE_MAX_SIDE = 1600          # 문서 그림의 긴 변 상한(px) — 한글 문서의 그림은 2~8MB 가 흔해 독스 문서가 커진다(내보내기 10MB 한도)
IMAGE_BIG_BYTES = 400_000


def normalize_image(data: bytes) -> bytes:
    """워드·독스가 받는 그림으로 — BMP 등은 PNG 로, JFIF·Exif 머리가 아닌 JPEG(ICC 프로필 APP2 로 시작, 실측 2026-09-24 사업계획서
    10장)는 다시 저장하고, 큰 그림은 긴 변 1600px·JPEG(투명하면 PNG)로 줄인다. python-docx 는 머리를 보고 형식을 알아낸다."""
    head = data[:4]
    ok_head = head.startswith(b"\x89PNG") or head in (b"\xff\xd8\xff\xe0", b"\xff\xd8\xff\xe1")
    if ok_head and len(data) <= IMAGE_BIG_BYTES:
        return data
    try:
        from PIL import Image

        im = Image.open(io.BytesIO(data))
        im.load()
        w, h = im.size
        scale = min(1.0, IMAGE_MAX_SIDE / max(w, h))
        if scale < 1.0:
            im = im.resize((max(1, int(w * scale)), max(1, int(h * scale))))
        elif ok_head and len(data) <= IMAGE_BIG_BYTES * 4 and (im.format or "").upper() in ("PNG", "JPEG"):
            return data                                    # 크지 않은 정상 그림은 그대로
        buf = io.BytesIO()
        transparent = im.mode in ("RGBA", "LA", "P") and (im.mode != "P" or "transparency" in im.info)
        if transparent:
            im.convert("RGBA").save(buf, format="PNG", optimize=True)
        else:
            im.convert("RGB").save(buf, format="JPEG", quality=88, optimize=True)
        out = buf.getvalue()
        return out if len(out) < len(data) or not ok_head else data
    except Exception:
        return data


def _binary_map(zf: zipfile.ZipFile, names: list[str]) -> dict[str, str]:
    out: dict[str, str] = {}
    hpf = next((n for n in names if n.endswith("content.hpf")), None)
    if hpf:
        try:
            root = _fromstring(zf.read(hpf))
            for item in root.iter():
                if _local(item.tag) == "item" and item.get("id") and item.get("href"):
                    out[str(item.get("id"))] = str(item.get("href")).lstrip("/")
        except ET.ParseError:
            pass
    for n in names:
        if n.startswith("BinData/"):
            out.setdefault(Path(n).stem, n)
    return out


def convert(path: Path | str) -> tuple[bytes, dict]:
    """HWPX 파일 → (docx 바이트, 통계). 암호화(배포용)면 ValueError."""
    path = Path(path)
    with zipfile.ZipFile(path) as zf:
        names = zf.namelist()
        header = next((n for n in names if n.endswith("Contents/header.xml")), None)
        styles = load_styles(_fromstring(zf.read(header))) if header else Styles()
        conv = Converter(styles, zf, _binary_map(zf, names))
        sections = sorted((n for n in names if _SECTION_RE.search(n)),
                          key=lambda n: int(_SECTION_RE.search(n).group(1)))  # type: ignore[union-attr]
        for sec in sections:
            raw = zf.read(sec)
            if not raw.lstrip().startswith(b"<"):
                raise ValueError(f"배포용/암호화 HWPX로 보임: {path.name}")
            root = _fromstring(raw)
            for p_el in _children(root, "p"):
                conv.paragraph(p_el, conv.doc)
        # 문서 기본 글꼴(독스가 없는 글꼴은 바꿔 쓴다)
        from docx.shared import Pt

        conv.doc.styles["Normal"].font.size = Pt(10)
        buf = io.BytesIO()
        conv.doc.save(buf)
        return buf.getvalue(), conv.stats
