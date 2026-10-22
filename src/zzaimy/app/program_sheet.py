"""사업 현황 엑셀 → 사업 관리 표의 연도별 값(ADR-0066).

학교가 해마다 내는 재정지원사업 현황 표처럼 「한 줄 = 그해 진행한 사업 하나」인 엑셀을 읽어 사업 관리 표(program_admin)의
줄과 잇고, 줄마다 연도별 값(예산·책임자·주관 부서·기간 …)을 채운다. 장부(kg_external.json)와 그래프는 건드리지 않는다.

- 읽기(read_book): 머리줄은 「사업명」 칸이 있는 줄. 열은 머리줄 이름의 낱말로 맞추므로 해마다 열 순서·개수가 달라도 된다.
  병합 칸은 왼쪽 위 값으로 채우되(가로로 합친 머리줄 「총사업기간」 = 시작·끝 두 칸) 금액 칸은 겹셈을 막으려 왼쪽 위에서만 읽는다.
  빈 줄·합계 줄·여러 칸을 가로로 합친 메모 줄은 건너뛴다. 연도는 제목의 「NNNN학년도·년도」, 없으면 시트 이름·문서 이름의 연도.
- 잇기(match): 이름 열쇠(program_names.group_key)가 줄의 이름·별칭 열쇠와 같으면 그 줄, 괄호 속 사업 이름이 같으면 그 줄,
  아니면 이름 속에 든 가장 긴 별칭. 가장 긴 겹침이 둘 이상의 줄이면 모호로 남기고 채우지 않는다. 못 찾으면 「엑셀에서 온 줄」로 더한다.
- 채우기(plan): 줄마다 years[연도] 에 값을 적는다. 손으로 고친 칸(by_hand)은 덮지 않는 것이 기본이다.
"""
from __future__ import annotations

import copy
import datetime as _dt
import re
import time
from pathlib import Path

from zzaimy.app.program_names import group_key
from zzaimy.app.program_register import slug

# 머리줄 낱말(띄어쓰기·줄바꿈을 지운 뒤 비교) — 앞에 적은 열쇠가 먼저 칸을 가져간다
HEAD = (
    ("name", ("사업명",)),
    ("sub", ("세부사업", "과제명")),
    ("year_period", ("당해연도사업기간", "당해연도기간")),
    ("period", ("총사업기간", "사업기간")),
    ("manager", ("사업책임자", "책임자")),
    ("officer", ("행정담당자", "담당자")),
    ("dept", ("업무진행부서", "주관부서", "담당부서", "업무부서")),
    ("account_dept", ("회계처리부서",)),
    ("majors", ("참여학과", "대상학과")),
    ("ministry", ("지원기관",)),
    ("agency", ("주관",)),
    ("status", ("기존", "신규", "종료", "구분")),
    ("national", ("국고",)),
)
MONEY = ("budget", "national")
# 화면·저장에 쓰는 연도별 칸
YEAR_FIELDS = ("budget", "manager", "dept", "period", "year_period", "sub", "ministry", "agency", "status", "majors",
               "account_dept", "officer", "national")
LABELS = {"budget": "예산", "manager": "책임자", "dept": "주관 부서", "period": "기간", "year_period": "그해 기간",
          "sub": "세부 사업", "ministry": "지원기관", "agency": "주관 기관", "status": "구분", "majors": "참여 학과",
          "account_dept": "회계 부서", "officer": "행정 담당자", "national": "국고"}

_YEAR_TITLE = re.compile(r"((?:19|20)\d{2})\s*(?:학년도|년도|회계연도|년)")
_YEAR_ANY = re.compile(r"(?<!\d)((?:19|20)\d{2})(?!\d)")
_LEAD_YEAR = re.compile(r"^\s*(?:19|20)\d{2}\s*(?:년도|학년도|년)?\s*(?:상반기|하반기)?\s*")
_TOTAL = re.compile(r"^\s*(?:합\s*계|소\s*계|총\s*계|계|누\s*계)\s*$")
_NOTE = ("※", "*", "주)", "비고")
_DATE = re.compile(r"((?:19|20)\d{2})[.\-/년\s]*(\d{1,2})[.\-/월\s]*(\d{1,2})")


def _flat(t) -> str:
    return re.sub(r"\s+", "", str(t or ""))


def _text(v) -> str:
    if v is None:
        return ""
    if isinstance(v, (_dt.datetime, _dt.date)):
        return f"{v.year}.{v.month:02d}.{v.day:02d}"
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return " ".join(str(v).split())


def _dates(t: str) -> list[str]:
    """글 속 날짜들 → 「YYYY.MM.DD」. 「20220301-20270228」「2022.3.1~2027.2.28」 꼴."""
    out = []
    for m in re.finditer(r"(?<!\d)((?:19|20)\d{2})(\d{2})(\d{2})(?!\d)", t):
        out.append((m.start(), f"{m.group(1)}.{m.group(2)}.{m.group(3)}"))
    if not out:
        for m in _DATE.finditer(t):
            mo, d = int(m.group(2)), int(m.group(3))
            if 1 <= mo <= 12 and 1 <= d <= 31:
                out.append((m.start(), f"{m.group(1)}.{mo:02d}.{d:02d}"))
    return [d for _, d in sorted(out)]


def fmt_period(parts: list) -> str:
    """기간 칸(하나 또는 시작·끝 두 칸) → 「2022.03.01~2027.02.28」. 날짜를 못 읽으면 글 그대로."""
    texts = [_text(p) for p in parts if _text(p)]
    ds = [d for t in texts for d in _dates(t)]
    if ds:
        return ds[0] if len(ds) == 1 else f"{ds[0]}~{ds[-1]}"
    return " ~ ".join(texts)


def money(v) -> int | None:
    if isinstance(v, bool) or v is None:
        return None
    if isinstance(v, (int, float)):
        return int(round(v))
    t = re.sub(r"[,\s원]", "", str(v))
    return int(t) if re.fullmatch(r"-?\d+", t) else None


def won(n: int | None) -> str:
    return "" if n is None else f"{n:,}원"


def clean_name(t: str) -> str:
    return _LEAD_YEAR.sub("", " ".join(str(t or "").split())).strip()


def _columns(head: list[str]) -> dict[str, list[int]]:
    """머리줄 → {열쇠: [열 번호…]}. 같은 이름이 가로로 이어진 칸(병합)은 한 열쇠의 여러 칸이 된다.
    예산은 「합계」로 시작하는 칸 가운데 맨 왼쪽(「상반기 합계」·이월금 포함 합계보다 그해 합계가 먼저 온다)."""
    col: dict[str, list[int]] = {}
    taken: set[int] = set()
    for key, words in HEAD:
        for j, c in enumerate(head):
            f = _flat(c)
            if j in taken or not f or not any(w in f for w in words):
                continue
            if key in col and not (col[key][-1] == j - 1 and _flat(head[j - 1]) == f):
                continue                                    # 첫 칸(과 가로로 이어진 같은 이름 칸)만
            col.setdefault(key, []).append(j)
            taken.add(j)
    tot = [j for j, c in enumerate(head) if _flat(c).startswith("합계") and j not in taken]
    if tot:
        col["budget"] = [tot[0]]
        taken.add(tot[0])
    return col


def _grid(ws) -> tuple[list[list], set[tuple[int, int]], dict[tuple[int, int], int]]:
    """시트 값 격자 — 병합 칸은 왼쪽 위 값으로 채운다. (격자, 채운 칸 좌표, 칸 → 가로 병합 너비)."""
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    width = 0
    for r in rows:                                          # 뒤쪽 빈 열은 잘라 낸다(서식만 있는 수천 열)
        k = len(r)
        while k and r[k - 1] in (None, ""):
            k -= 1
        width = max(width, k)
    rows = [(r + [None] * width)[:width] for r in rows]
    filled: set[tuple[int, int]] = set()
    span: dict[tuple[int, int], int] = {}
    for m in getattr(ws, "merged_cells", None) and ws.merged_cells.ranges or []:
        r0, c0, r1, c1 = m.min_row - 1, m.min_col - 1, m.max_row - 1, m.max_col - 1
        if r0 >= len(rows) or c0 >= width:
            continue
        v = rows[r0][c0]
        for i in range(r0, min(r1, len(rows) - 1) + 1):
            for j in range(c0, min(c1, width - 1) + 1):
                span[(i, j)] = c1 - c0 + 1
                if (i, j) != (r0, c0):
                    rows[i][j] = v
                    filled.add((i, j))
    return rows, filled, span


def read_book(path, hint: str = "") -> list[dict]:
    """엑셀 한 권 → 연도별 줄 [{year, name, sub, ministry, agency, period, year_period, manager, officer, dept, account_dept,
    majors, status, budget, budget_won, national_won, extra, file, sheet, line}]. hint 는 연도를 못 찾을 때 볼 문서 이름."""
    import openpyxl
    wb = openpyxl.load_workbook(Path(path), data_only=True)
    out = []
    for ws in wb.worksheets:
        rows, filled, span = _grid(ws)
        hi = next((i for i, r in enumerate(rows[:15]) if any("사업명" in _flat(c) for c in r)), None)
        if hi is None:
            continue
        top = " ".join(_text(c) for r in rows[:hi] for c in r if c not in (None, ""))
        m = _YEAR_TITLE.search(top) or _YEAR_ANY.search(ws.title or "") or _YEAR_ANY.search(hint or "") \
            or _YEAR_ANY.search(Path(path).stem)
        if not m:
            continue
        year = int(m.group(1))
        head = [_text(c) for c in rows[hi]]
        col = _columns(head)
        if "name" not in col:
            continue
        nj = col["name"][0]
        used = {j for js in col.values() for j in js}
        for i in range(hi + 1, len(rows)):
            r = rows[i]
            name = clean_name(_text(r[nj]))
            if not name or len(name) < 3 or name.startswith(_NOTE) or _TOTAL.match(name):
                continue
            if span.get((i, nj), 1) >= 3:                   # 여러 칸을 가로로 합친 줄 — 합계·메모
                continue
            if any(_TOTAL.match(_text(r[j])) for j in range(0, nj)):
                continue

            def get(key, r=r, i=i):
                js = col.get(key) or []
                if key in MONEY:
                    js = [j for j in js if (i, j) not in filled]
                return [r[j] for j in js]

            item = {"year": year, "name": name, "file": Path(path).name if not hint else hint, "sheet": ws.title, "line": i + 1}
            for key in ("sub", "ministry", "agency", "manager", "officer", "dept", "account_dept", "majors", "status"):
                item[key] = ", ".join(dict.fromkeys(t for t in (_text(v) for v in get(key)) if t))
            item["period"] = fmt_period(get("period"))
            item["year_period"] = fmt_period(get("year_period"))
            b = next((n for n in (money(v) for v in get("budget")) if n is not None), None)
            item["budget_won"], item["budget"] = b, won(b)
            item["national_won"] = next((n for n in (money(v) for v in get("national")) if n is not None), None)
            item["national"] = won(item["national_won"])
            item["extra"] = {head[j] or f"{j + 1}열": _text(r[j]) for j in range(len(head))
                             if j not in used and _text(r[j]) and head[j] and head[j] not in ("순", "번호")}
            out.append(item)
    return out


# ---------------------------------------------------------------- 잇기

class Matcher:
    """사업 관리 줄 찾기 — 이름 열쇠가 같으면 그 줄, 괄호 속 이름이 같으면 그 줄, 아니면 가장 긴 별칭 겹침(같은 길이로 여럿이면 모호)."""

    MIN = 4

    def __init__(self, rows: list[dict], not_program: set[str] | None = None):
        self.exact: dict[str, set[str]] = {}
        self.keys: list[tuple[str, str]] = []
        self.name = {r["key"]: r.get("name") or r["key"] for r in rows}
        self.not_program = not_program or set()
        for r in rows:
            for t in [r.get("name") or "", *(r.get("terms") or [])]:
                k = group_key(t)
                if k:
                    self.exact.setdefault(k, set()).add(r["key"])
                    if len(k) >= self.MIN:
                        self.keys.append((k, r["key"]))

    def find(self, name: str) -> tuple[str, list[str]]:
        """→ (갈래, 줄 열쇠들). 갈래: exact · inner · contains · ambiguous · new · not_program."""
        k = group_key(name)
        if not k:
            return "new", []
        hit = self.exact.get(k) or set()
        if len(hit) == 1:
            return "exact", sorted(hit)
        if len(hit) > 1:
            return "ambiguous", sorted(hit)
        if k in self.not_program:
            return "not_program", []
        for part in re.findall(r"\(([^)]{4,})\)", name or ""):
            h = self.exact.get(group_key(part)) or set()
            if len(h) == 1:
                return "inner", sorted(h)
        best, got = 0, set()
        for t, key in self.keys:
            if t in k or (len(k) >= 6 and k in t):
                n = min(len(t), len(k))
                if n > best:
                    best, got = n, {key}
                elif n == best:
                    got.add(key)
        if len(got) == 1:
            return "contains", sorted(got)
        if len(got) > 1:
            return "ambiguous", sorted(got)
        return "new", []


def _join(a: str, b: str) -> str:
    parts = [p for p in re.split(r"\s*,\s*", a or "") if p] + [p for p in re.split(r"\s*,\s*", b or "") if p]
    return ", ".join(dict.fromkeys(parts))


def _fold(items: list[dict]) -> dict:
    """같은 사업·같은 해의 여러 줄(세부 과제) → 한 칸 묶음. 금액은 더하고 글은 겹치지 않게 잇는다."""
    out: dict = {"lines": len(items)}
    for f in YEAR_FIELDS:
        if f in MONEY:
            ns = [it.get(f + "_won") for it in items if it.get(f + "_won") is not None]
            if ns:
                out[f + "_won"] = sum(ns)
                out[f] = won(out[f + "_won"])
            continue
        v = ""
        for it in items:
            v = _join(v, it.get(f) or "")
        if v:
            out[f] = v
    extra: dict = {}
    for it in items:
        for h, v in (it.get("extra") or {}).items():
            extra[h] = _join(extra.get(h, ""), v)
    if extra:
        out["extra"] = extra
    out["from"] = ", ".join(dict.fromkeys(str(it.get("file") or "") for it in items if it.get("file")))
    return out


def latest(years: dict) -> dict:
    """연도별 값 가운데 가장 최근 해의 값."""
    ys = [y for y in (years or {}) if str(y).isdigit()]
    return dict(years[max(ys, key=int)]) if ys else {}


def plan(items: list[dict], rows: list[dict], store: dict, overwrite: bool = False, not_program: set[str] | None = None,
         by: str = "") -> tuple[dict, dict]:
    """엑셀 줄 → (요약, 바뀐 저장값). store 는 고치지 않고 사본을 돌려준다.
    요약: matched·new·ambiguous·not_program(이름 목록), cells{fill, change, same, kept}, blank_top(빈 칸이던 표의 칸 가운데 채워지는 수)."""
    store = copy.deepcopy(store)
    store.setdefault("rows", {})
    store.setdefault("added", {})
    m = Matcher(rows, not_program)
    by_key = {r["key"]: r for r in rows}
    groups: dict[str, dict] = {}                       # 줄 열쇠 → {year: [items]}
    new_names: dict[str, list[dict]] = {}              # 이름 열쇠 → items (장부에 없는 사업)
    matched: dict[str, set[str]] = {}
    ambiguous: dict[str, list[str]] = {}
    skipped: set[str] = set()
    for it in items:
        how, keys = m.find(it["name"])
        if how in ("exact", "inner", "contains"):
            groups.setdefault(keys[0], {}).setdefault(str(it["year"]), []).append(it)
            matched.setdefault(keys[0], set()).add(it["name"])
        elif how == "ambiguous":
            ambiguous[it["name"]] = [m.name.get(k, k) for k in keys]
        elif how == "not_program":
            skipped.add(it["name"])
        else:
            new_names.setdefault(group_key(it["name"]) or it["name"], []).append(it)
    new_rows = []
    for gk, its in new_names.items():
        its.sort(key=lambda x: x["year"])
        last = its[-1]
        key = slug(last["name"])
        n = 2
        while key in by_key or key in store["added"]:
            key, n = f"{slug(last['name'])}-{n}", n + 1
        store["added"][key] = {"source": "엑셀", "name": last["name"],
                               "terms": list(dict.fromkeys(x["name"] for x in its)),
                               "dept": last.get("dept") or "", "period": last.get("period") or "",
                               "added_by": by, "added_at": time.strftime("%Y-%m-%d")}
        for it in its:
            groups.setdefault(key, {}).setdefault(str(it["year"]), []).append(it)
        new_rows.append({"key": key, "name": last["name"], "years": sorted({x["year"] for x in its})})
    cells = {"fill": 0, "change": 0, "same": 0, "kept": 0}
    blank_top = 0
    for key, years in groups.items():
        o = store["rows"].setdefault(key, {})
        ys = o.setdefault("years", {})
        for y, its in years.items():
            got = _fold(its)
            cur = ys.setdefault(y, {})
            hand = set(cur.get("by_hand") or [])
            for f in YEAR_FIELDS:
                v = got.get(f)
                if not v:
                    continue
                old = cur.get(f) or ""
                if old == v:
                    cells["same"] += 1
                    continue
                if f in hand and not overwrite:
                    cells["kept"] += 1
                    continue
                cells["fill" if not old else "change"] += 1
                cur[f] = v
                if f in MONEY:
                    cur[f + "_won"] = got.get(f + "_won")
                if f in hand:
                    hand.discard(f)
            if got.get("extra"):
                cur["extra"] = {**(cur.get("extra") or {}), **got["extra"]}
            cur["lines"] = got["lines"]
            cur["from"] = got.get("from") or ""
            cur["at"] = time.strftime("%Y-%m-%d")
            if hand:
                cur["by_hand"] = sorted(hand)
            else:
                cur.pop("by_hand", None)
        r = by_key.get(key)
        if r is not None:                               # 표에서 빈칸이던 칸 가운데 연도 값으로 채워 보이는 칸
            before, after = latest(r.get("years") or {}), latest(ys)
            blank_top += sum(1 for f in ("dept", "period", "budget", "manager") if not (r.get(f) or before.get(f)) and after.get(f))
    summary = {
        "lines": len(items),
        "years": sorted({it["year"] for it in items}),
        "matched": sorted(((by_key[k]["name"] if k in by_key else k), sorted(v)) for k, v in matched.items()),
        "new": new_rows,
        "ambiguous": sorted(ambiguous.items()),
        "not_program": sorted(skipped),
        "cells": cells,
        "blank_top": blank_top,
        "overwrite": overwrite,
    }
    return summary, store
